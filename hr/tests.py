import ast
import hashlib
import threading
import time
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from PIL import Image
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import DatabaseError, IntegrityError, connection, transaction
from django.test import RequestFactory, TestCase, TransactionTestCase, override_settings
from hr.audit import SENSITIVE_AUDIT_FIELDS, get_client_ip, is_ip_trusted, log_audit_event, verify_chain
from hr.forms import (
    DocumentEditForm,
    DocumentForm,
    DocumentReplaceForm,
    EndEngagementForm,
    EngagementForm,
    PersonForm,
)
from hr.models import (
    Applicant,
    AuditLog,
    CalendarEvent,
    CertificateRequest,
    Delegate,
    Document,
    EmployeeSequence,
    Engagement,
    LeaveCredit,
    LeaveRequest,
    LeaveType,
    Meeting,
    MeetingSeries,
    Memo,
    Office,
    StructurePosition,
    Person,
    Project,
    SalaryAction,
    SalaryCheckRun,
    ServiceInterruption,
    UserProfile,
    Vacancy,
)
from hr.sequences import get_next_employee_no, sync_employee_sequence
from hr.permissions import (
    PROTECTED_ROLES,
    ROLE_AUDITOR,
    ROLE_EXECUTIVE,
    ROLE_HRMO,
    ROLE_OFFICE_HEAD,
    ROLE_OFFICE_STAFF,
    ROLE_SYSADMIN,
    is_active_delegate,
    is_executive,
)



class OfficeAndSeedOfficesTestCase(TestCase):
    def test_seed_offices_creates_25_offices_and_is_idempotent(self):
        call_command("seed_offices")
        self.assertEqual(Office.objects.count(), 25)

        # Check legislative offices
        ossb = Office.objects.get(acronym="OSSB")
        self.assertEqual(ossb.branch, Office.Branch.LEGISLATIVE)
        self.assertIsNotNone(ossb.parent)
        self.assertEqual(ossb.parent.name, "Office of the Sangguniang Bayan")

        # Check institution
        college = Office.objects.get(name="San Pascual Polytechnic College")
        self.assertEqual(college.type, Office.Type.INSTITUTION)

        # Idempotence: calling again updates without duplicate
        call_command("seed_offices")
        self.assertEqual(Office.objects.count(), 25)


class PersonModelTestCase(TestCase):
    def test_fake_id_constraint(self):
        # Valid fake ID starting with 00
        p = Person.objects.create(
            employee_no="SP-00001",
            last_name="Dela Cruz",
            first_name="Juan",
            birthdate=date(1990, 1, 1),
            tin="000-123-456-000",
        )
        self.assertEqual(str(p), "Dela Cruz, Juan")

        # Real/non-zero-prefixed ID should fail DB constraint
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Person.objects.create(
                    employee_no="SP-00002",
                    last_name="Santos",
                    first_name="Maria",
                    birthdate=date(1992, 5, 10),
                    tin="123-456-789-000",  # Invalid for prototype
                )

    def test_duplicate_person_constraint(self):
        Person.objects.create(
            employee_no="SP-00001",
            last_name="Reyes",
            first_name="Pedro",
            birthdate=date(1985, 3, 15),
        )
        # Duplicate name + birthdate should raise IntegrityError
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Person.objects.create(
                    employee_no="SP-00002",
                    last_name="reyes",
                    first_name="PEDRO",
                    birthdate=date(1985, 3, 15),
                )


class EngagementModelTestCase(TestCase):
    def setUp(self):
        self.office = Office.objects.create(
            name="HRMO",
            branch=Office.Branch.EXECUTIVE,
            type=Office.Type.DEPARTMENT,
        )
        self.person = Person.objects.create(
            employee_no="SP-00010",
            last_name="Aquino",
            first_name="Corazon",
            birthdate=date(1980, 2, 2),
        )

    def test_plantilla_requires_sg_step_item(self):
        # Missing salary grade / step / item number on plantilla fails constraint
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Engagement.objects.create(
                    person=self.person,
                    office=self.office,
                    type=Engagement.Type.PLANTILLA,
                    position_title="Administrative Officer V",
                    start_date=date(2020, 1, 1),
                )

        # Valid plantilla engagement succeeds
        eng = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.PLANTILLA,
            position_title="Administrative Officer V",
            salary_grade=18,
            step=1,
            item_number="SP-HRMO-001",
            rate=Decimal("46725.00"),
            rate_basis=Engagement.RateBasis.MONTHLY,
            start_date=date(2020, 1, 1),
        )
        self.assertIsNotNone(eng.pk)

    def test_non_plantilla_forbids_salary_grade(self):
        # Job order with salary grade fails constraint
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Engagement.objects.create(
                    person=self.person,
                    office=self.office,
                    type=Engagement.Type.JOB_ORDER,
                    position_title="Office Clerk",
                    salary_grade=5,
                    step=1,
                    item_number="SP-JO-01",
                    start_date=date(2022, 1, 1),
                )

    def test_project_based_requires_project(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Engagement.objects.create(
                    person=self.person,
                    office=self.office,
                    type=Engagement.Type.PROJECT_BASED,
                    position_title="Site Inspector",
                    start_date=date(2023, 1, 1),
                    project=None,
                )

        project = Project.objects.create(
            name="Municipal Hall Renovation",
            office=self.office,
            start_date=date(2023, 1, 1),
        )
        eng = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.PROJECT_BASED,
            position_title="Site Inspector",
            start_date=date(2023, 1, 1),
            project=project,
        )
        self.assertIsNotNone(eng.pk)


class DocumentAndLeaveTestCase(TestCase):
    def setUp(self):
        self.office = Office.objects.create(
            name="Office of the Mayor",
            branch=Office.Branch.EXECUTIVE,
            type=Office.Type.DEPARTMENT,
        )
        self.person = Person.objects.create(
            employee_no="SP-00020",
            last_name="Bautista",
            first_name="Herbert",
            birthdate=date(1975, 4, 12),
        )

    def test_document_xor_person_engagement(self):
        eng = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.CASUAL,
            position_title="Staff",
            start_date=date(2021, 1, 1),
        )
        # Cannot attach to both person and engagement
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Document.objects.create(
                    person=self.person,
                    engagement=eng,
                    doc_type=Document.Type.PDS,
                )

        # Can attach to person alone
        doc = Document.objects.create(
            person=self.person,
            doc_type=Document.Type.PDS,
            physical_location="Cabinet 1, Folder A",
        )
        self.assertIsNotNone(doc.pk)

    def test_leave_credit_validation(self):
        # Person without active plantilla or casual engagement cannot have leave credit
        credit = LeaveCredit(
            person=self.person,
            leave_type=LeaveType.VACATION,
            days=Decimal("15.0"),
            as_of=date(2026, 1, 1),
        )
        with self.assertRaises(ValidationError):
            credit.clean()

        # Add leave-eligible casual engagement -> clean() passes
        Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.CASUAL,
            position_title="Utility Worker",
            start_date=date(2026, 1, 1),
            status=Engagement.Status.ACTIVE,
            leave_eligible=True,
        )
        credit.clean()  # should not raise


class SeedFakeCommandTestCase(TestCase):
    def test_seed_fake_command(self):
        call_command("seed_fake", "--clear")

        # Scaled headcount: 420 people matching survey composition
        people_count = Person.objects.count()
        self.assertEqual(people_count, 420)

        # Active engagement breakdown matches survey findings
        active_eng = Engagement.objects.filter(status=Engagement.Status.ACTIVE)
        self.assertEqual(active_eng.filter(type=Engagement.Type.PLANTILLA).count(), 94)
        self.assertEqual(active_eng.filter(type=Engagement.Type.CASUAL).count(), 28)
        self.assertEqual(active_eng.filter(type=Engagement.Type.ELECTED).count(), 12)
        jo_cos_count = active_eng.filter(
            type__in=[Engagement.Type.JOB_ORDER, Engagement.Type.CONTRACT_OF_SERVICE, Engagement.Type.PROJECT_BASED]
        ).count()
        self.assertGreaterEqual(jo_cos_count, 280)

        # All fake ID numbers start with "00"
        for p in Person.objects.all():
            if p.tin:
                self.assertTrue(p.tin.startswith("00"), f"TIN {p.tin} does not start with 00")
            if p.gsis_no:
                self.assertTrue(p.gsis_no.startswith("00"), f"GSIS {p.gsis_no} does not start with 00")

        # Career progression check: some people have multiple engagements (JO -> Casual -> Plantilla)
        people_with_multi_engagements = Person.objects.filter(engagements__status=Engagement.Status.ENDED).distinct()
        self.assertGreaterEqual(people_with_multi_engagements.count(), 10)

        # Projects and Documents exist
        self.assertEqual(Project.objects.count(), 6)
        self.assertGreaterEqual(Document.objects.count(), 60)

        # Leave requests exist
        self.assertGreaterEqual(LeaveRequest.objects.count(), 20)


class ViewsTestCase(TestCase):
    def setUp(self):
        User = get_user_model()
        call_command("seed_offices")
        call_command("seed_fake")
        call_command("setup_roles")

        # HRMO staff user (sees all offices, all people, and pay rates)
        self.user = User.objects.create_user(username="hrmouser", password="securepassword123")
        self.user.groups.add(Group.objects.get(name=ROLE_HRMO))

        # Regular employee user (self-service only)
        self.regular_person = Person.objects.first()
        self.regular_user = User.objects.create_user(username="reguser", password="securepassword123")
        UserProfile.objects.create(user=self.regular_user, person=self.regular_person)

    def test_views_require_login(self):
        # Unauthenticated users redirected to login
        resp = self.client.get("/offices/")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/accounts/login/", resp.url)

        resp = self.client.get("/people/")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/accounts/login/", resp.url)

    def test_office_views(self):
        self.client.force_login(self.user)
        resp = self.client.get("/offices/")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Municipal Offices Directory")
        self.assertContains(resp, "Office of the Municipal Mayor")

        # Filter by branch
        resp_leg = self.client.get("/offices/?branch=legislative")
        self.assertEqual(resp_leg.status_code, 200)
        self.assertContains(resp_leg, "Office of the Sangguniang Bayan")
        self.assertNotContains(resp_leg, "Office of the Municipal Mayor")

        # Office detail
        office = Office.objects.first()
        resp_det = self.client.get(f"/offices/{office.pk}/")
        self.assertEqual(resp_det.status_code, 200)
        self.assertContains(resp_det, office.name)

    def test_person_views(self):
        self.client.force_login(self.user)
        resp = self.client.get("/people/")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Personnel Directory")
        self.assertEqual(len(resp.context["page_obj"]), 25)

        person = Person.objects.first()
        resp_det = self.client.get(f"/people/{person.pk}/")
        self.assertEqual(resp_det.status_code, 200)
        self.assertContains(resp_det, person.last_name)
        self.assertContains(resp_det, person.employee_no)
        # HRMO user can see rates
        self.assertTrue(resp_det.context["can_view_rates"])

    def test_rate_visibility_for_non_staff(self):
        self.client.force_login(self.regular_user)
        resp_det = self.client.get(f"/people/{self.regular_person.pk}/")
        self.assertEqual(resp_det.status_code, 200)
        # Regular employee user does not see rates
        self.assertFalse(resp_det.context["can_view_rates"])
        self.assertContains(resp_det, "Pay rates are restricted")



class Step4AModelAndAdminTestCase(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        self.superuser = User.objects.create_superuser("superadmin", "super@test.com", "superpass123")
        self.staff_sysadmin = User.objects.create_user("sysadmin", "sys@test.com", "syspass123", is_staff=True)
        self.mayor = User.objects.create_user("mayor", "mayor@test.com", "mayorpass123")
        self.delegate_user = User.objects.create_user("delegate", "delegate@test.com", "delpass123")

        self.office = Office.objects.create(name="Office of the Mayor", branch=Office.Branch.EXECUTIVE, type=Office.Type.DEPARTMENT)
        self.person = Person.objects.create(
            employee_no="SP-00099",
            last_name="Santiago",
            first_name="Miriam",
            birthdate=date(1970, 6, 15),
        )

    def test_engagement_leave_eligible_defaults_and_toggle(self):
        # The type default (True for Plantilla/Casual) is a suggestion for entry screens;
        # the model stores exactly what it is given (DECISIONS L3, LeaveEligibleDefaultTestCase).
        self.assertTrue(Engagement.default_leave_eligible(Engagement.Type.PLANTILLA))
        self.assertTrue(Engagement.default_leave_eligible(Engagement.Type.CASUAL))
        self.assertFalse(Engagement.default_leave_eligible(Engagement.Type.JOB_ORDER))

        plantilla = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.PLANTILLA,
            position_title="Attorney IV",
            salary_grade=23,
            step=1,
            item_number="SP-MO-099",
            start_date=date(2025, 1, 1),
            leave_eligible=True,
        )
        self.assertTrue(plantilla.leave_eligible)
        self.assertTrue(self.person.earns_leave_credits())

        # Casual saved without the flag is not forced to True
        casual = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.CASUAL,
            position_title="Legal Aide",
            start_date=date(2025, 1, 1),
        )
        self.assertFalse(casual.leave_eligible)

        # Job order defaults leave_eligible=False
        jo = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.JOB_ORDER,
            position_title="Office Encoder",
            start_date=date(2025, 1, 1),
        )
        self.assertFalse(jo.leave_eligible)

        # HRMO can explicitly toggle leave_eligible on JO
        jo.leave_eligible = True
        jo.save()
        self.assertTrue(jo.leave_eligible)

    def test_document_original_on_file(self):
        doc = Document.objects.create(
            person=self.person,
            doc_type=Document.Type.PDS,
            physical_location="Cabinet 1, Drawer 2",
        )
        self.assertTrue(doc.original_on_file)

    def test_delegate_model_and_active_window(self):
        d = Delegate.objects.create(
            principal=self.mayor,
            delegate=self.delegate_user,
            valid_from=date(2026, 10, 1),
            valid_to=date(2026, 10, 15),
            notes="Covering Mayor leave",
        )
        self.assertTrue(d.is_active(date(2026, 10, 5)))
        self.assertFalse(d.is_active(date(2026, 10, 20)))  # Expired
        self.assertFalse(d.is_active(date(2026, 9, 28)))   # Prior

        # valid_to < valid_from constraint
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Delegate.objects.create(
                    principal=self.mayor,
                    delegate=self.delegate_user,
                    valid_from=date(2026, 10, 20),
                    valid_to=date(2026, 10, 10),
                )

    def test_admin_hr_models_forbidden_for_non_superusers(self):
        # Non-superuser (System Admin) accessing HR models in /admin/ gets 403
        self.client.force_login(self.staff_sysadmin)
        resp = self.client.get("/admin/hr/person/")
        self.assertEqual(resp.status_code, 403)

        # Superuser can access /admin/hr/person/
        self.client.force_login(self.superuser)
        resp_super = self.client.get("/admin/hr/person/")
        self.assertEqual(resp_super.status_code, 200)


class Step4CSecurityAndRolesTestCase(TestCase):
    def setUp(self):
        User = get_user_model()
        call_command("seed_offices")
        call_command("setup_roles")

        self.superuser = User.objects.create_superuser("superadmin", "super@test.com", "superpass123")
        self.sysadmin = User.objects.create_user("sysadmin", "sys@test.com", "syspass123", is_staff=True)
        self.sysadmin.groups.add(Group.objects.get(name=ROLE_SYSADMIN))

        self.hrmo_user = User.objects.create_user("hrmo", "hrmo@test.com", "hrmopass123")
        self.hrmo_user.groups.add(Group.objects.get(name=ROLE_HRMO))

        self.auditor_user = User.objects.create_user("auditor", "aud@test.com", "audpass123")
        self.auditor_user.groups.add(Group.objects.get(name=ROLE_AUDITOR))

        self.mayor = User.objects.create_user("mayor", "mayor@test.com", "mayorpass123")
        self.mayor.groups.add(Group.objects.get(name=ROLE_EXECUTIVE))

        self.delegate_user = User.objects.create_user("deluser", "del@test.com", "delpass123")

        # Two offices
        self.office_a = Office.objects.get(name="Office of the Municipal Mayor")
        self.office_b = Office.objects.get(name="Office of the Municipal Treasurer")

        # Office Head for Office A
        self.head_a = User.objects.create_user("head_a", "head_a@test.com", "headpass123")
        self.head_a.groups.add(Group.objects.get(name=ROLE_OFFICE_HEAD))
        UserProfile.objects.create(user=self.head_a, office=self.office_a)

        # Office Staff for Office A
        self.staff_a = User.objects.create_user("staff_a", "staff_a@test.com", "staffpass123")
        self.staff_a.groups.add(Group.objects.get(name=ROLE_OFFICE_STAFF))
        UserProfile.objects.create(user=self.staff_a, office=self.office_a)

        # Person A in Office A
        self.person_a = Person.objects.create(
            employee_no="SP-00101",
            last_name="Rizal",
            first_name="Jose",
            birthdate=date(1980, 6, 19),
            tin="000-111-222-000",
        )
        self.eng_a = Engagement.objects.create(
            person=self.person_a,
            office=self.office_a,
            type=Engagement.Type.PLANTILLA,
            position_title="Executive Assistant IV",
            salary_grade=22,
            step=1,
            item_number="SP-MO-001",
            rate=Decimal("50000.00"),
            rate_basis=Engagement.RateBasis.MONTHLY,
            start_date=date(2025, 1, 1),
            status=Engagement.Status.ACTIVE,
        )
        self.doc_a = Document.objects.create(
            person=self.person_a,
            doc_type=Document.Type.APPOINTMENT,
            physical_location="Cabinet A, Drawer 1",
            original_on_file=True,
        )

        # Person B in Office B
        self.person_b = Person.objects.create(
            employee_no="SP-00202",
            last_name="Bonifacio",
            first_name="Andres",
            birthdate=date(1982, 11, 30),
            tin="000-333-444-000",
        )
        self.eng_b = Engagement.objects.create(
            person=self.person_b,
            office=self.office_b,
            type=Engagement.Type.CASUAL,
            position_title="Revenue Collector I",
            rate=Decimal("15000.00"),
            rate_basis=Engagement.RateBasis.MONTHLY,
            start_date=date(2025, 1, 1),
            status=Engagement.Status.ACTIVE,
        )

        # Employee self-service user for Person A
        self.emp_a_user = User.objects.create_user("emp_a", "emp_a@test.com", "emppass123")
        UserProfile.objects.create(user=self.emp_a_user, person=self.person_a, office=self.office_a)

    def test_sysadmin_cannot_access_hr_front_views(self):
        self.client.force_login(self.sysadmin)
        resp_people = self.client.get("/people/")
        self.assertEqual(resp_people.status_code, 302)
        self.assertIn("/admin/", resp_people.url)

        resp_offices = self.client.get("/offices/")
        self.assertEqual(resp_offices.status_code, 302)
        self.assertIn("/admin/", resp_offices.url)

    def test_sysadmin_cannot_see_hr_models_in_admin(self):
        self.client.force_login(self.sysadmin)
        # Direct URL access to HR models is 403 Forbidden
        for model_path in ["person", "engagement", "document", "office", "leaverequest", "delegate"]:
            resp = self.client.get(f"/admin/hr/{model_path}/")
            self.assertEqual(resp.status_code, 403, f"Expected 403 for /admin/hr/{model_path}/")

        # Admin index does not show HR app
        resp_index = self.client.get("/admin/")
        self.assertEqual(resp_index.status_code, 200)
        self.assertNotContains(resp_index, "hr/person")
        self.assertNotContains(resp_index, "hr/engagement")

    def test_sysadmin_cannot_assign_protected_roles(self):
        from django.contrib.admin.sites import site
        from hr.admin import CustomUserAdmin
        User = get_user_model()
        user_admin = CustomUserAdmin(User, site)

        rf = RequestFactory()
        req = rf.get("/admin/auth/user/add/")
        req.user = self.sysadmin

        formfield = user_admin.formfield_for_manytomany(User._meta.get_field("groups"), req)
        group_choices = list(formfield.queryset.values_list("name", flat=True))
        for protected in PROTECTED_ROLES:
            self.assertNotIn(protected, group_choices, f"{protected} must not be selectable by SysAdmin")

        # Even with forged POST request, save_related blocks it
        new_user = User.objects.create_user("target_user", "target@test.com", "pass123")
        hrmo_group = Group.objects.get(name=ROLE_HRMO)
        new_user.groups.add(hrmo_group)

        class DummyForm:
            instance = new_user

            def save_m2m(self):
                pass

        with self.assertRaises(PermissionDenied):
            user_admin.save_related(req, DummyForm(), [], False)


        # Confirm target_user was stripped of HRMO Staff
        self.assertFalse(new_user.groups.filter(name=ROLE_HRMO).exists())

    def test_sysadmin_cannot_edit_superuser_or_protected_members(self):
        from django.contrib.admin.sites import site
        from hr.admin import CustomUserAdmin
        User = get_user_model()
        user_admin = CustomUserAdmin(User, site)

        rf = RequestFactory()
        req = rf.get(f"/admin/auth/user/{self.superuser.pk}/change/")
        req.user = self.sysadmin

        # Cannot edit emergency superuser
        self.assertFalse(user_admin.has_change_permission(req, self.superuser))
        self.assertFalse(user_admin.has_delete_permission(req, self.superuser))

        # Cannot edit HRMO staff member
        self.assertFalse(user_admin.has_change_permission(req, self.hrmo_user))
        self.assertFalse(user_admin.has_delete_permission(req, self.hrmo_user))

    def test_sysadmin_cannot_modify_groups_in_admin(self):
        from django.contrib.admin.sites import site
        from hr.admin import CustomGroupAdmin
        group_admin = CustomGroupAdmin(Group, site)

        rf = RequestFactory()
        req = rf.get("/admin/auth/group/")
        req.user = self.sysadmin

        self.assertFalse(group_admin.has_add_permission(req))
        self.assertFalse(group_admin.has_change_permission(req, Group.objects.first()))
        self.assertFalse(group_admin.has_delete_permission(req, Group.objects.first()))

    def test_office_head_scoping_and_minimization(self):
        self.client.force_login(self.head_a)

        # Scoped to own office
        resp_offices = self.client.get("/offices/")
        self.assertEqual(resp_offices.status_code, 200)
        self.assertContains(resp_offices, self.office_a.name)
        self.assertNotContains(resp_offices, self.office_b.name)

        resp_off_a = self.client.get(f"/offices/{self.office_a.pk}/")
        self.assertEqual(resp_off_a.status_code, 200)

        resp_off_b = self.client.get(f"/offices/{self.office_b.pk}/")
        self.assertEqual(resp_off_b.status_code, 404)

        # Scoped to active personnel in own office
        resp_people = self.client.get("/people/")
        self.assertEqual(resp_people.status_code, 200)
        self.assertContains(resp_people, self.person_a.last_name)
        self.assertNotContains(resp_people, self.person_b.last_name)

        # Person detail in own office: 200 OK, but NO pay and NO fake IDs
        resp_person_a = self.client.get(f"/people/{self.person_a.pk}/")
        self.assertEqual(resp_person_a.status_code, 200)
        self.assertFalse(resp_person_a.context["can_view_rates"])
        self.assertFalse(resp_person_a.context["can_view_id_numbers"])
        self.assertFalse(resp_person_a.context["can_view_doc_content"])
        self.assertNotContains(resp_person_a, "50000")
        self.assertNotContains(resp_person_a, "000-111-222-000")
        self.assertContains(resp_person_a, "Pay rates are restricted")
        self.assertContains(resp_person_a, "Government ID numbers are restricted")
        self.assertContains(resp_person_a, "Cabinet A, Drawer 1")

        # Person detail in other office: 404 Not Found
        resp_person_b = self.client.get(f"/people/{self.person_b.pk}/")
        self.assertEqual(resp_person_b.status_code, 404)

    def test_office_head_and_staff_cannot_download_documents(self):
        self.client.force_login(self.head_a)
        resp = self.client.get(f"/documents/{self.doc_a.pk}/download/")
        self.assertEqual(resp.status_code, 403)

        self.client.force_login(self.staff_a)
        resp_staff = self.client.get(f"/documents/{self.doc_a.pk}/download/")
        self.assertEqual(resp_staff.status_code, 403)

    def test_executive_sees_across_offices_without_pay_or_ids(self):
        self.client.force_login(self.mayor)

        # Sees across offices
        resp_a = self.client.get(f"/people/{self.person_a.pk}/")
        self.assertEqual(resp_a.status_code, 200)
        self.assertFalse(resp_a.context["can_view_rates"])
        self.assertFalse(resp_a.context["can_view_id_numbers"])
        self.assertNotContains(resp_a, "50000")
        self.assertNotContains(resp_a, "000-111-222-000")

        resp_b = self.client.get(f"/people/{self.person_b.pk}/")
        self.assertEqual(resp_b.status_code, 200)
        self.assertFalse(resp_b.context["can_view_rates"])
        self.assertFalse(resp_b.context["can_view_id_numbers"])
        self.assertNotContains(resp_b, "15000")
        self.assertNotContains(resp_b, "000-333-444-000")

    def test_executive_delegate_active_and_expired(self):
        from django.utils import timezone

        # Relative to today: is_executive() checks the delegation against the real date.
        today = timezone.localdate()
        d = Delegate.objects.create(
            principal=self.mayor,
            delegate=self.delegate_user,
            valid_from=today - timedelta(days=4),
            valid_to=today + timedelta(days=10),
            notes="Covering Mayor leave",
        )
        self.assertTrue(is_active_delegate(self.delegate_user, today))
        self.assertTrue(is_executive(self.delegate_user))

        # Active delegate can view across offices
        self.client.force_login(self.delegate_user)
        resp = self.client.get(f"/people/{self.person_a.pk}/")
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(resp.context["can_view_rates"])

        # Expired delegation
        d.valid_from = today - timedelta(days=34)
        d.valid_to = today - timedelta(days=20)
        d.save()
        self.assertFalse(is_active_delegate(self.delegate_user))
        self.assertFalse(is_executive(self.delegate_user))

        # Expired delegate has no role access
        resp_exp = self.client.get(f"/people/{self.person_a.pk}/")
        self.assertEqual(resp_exp.status_code, 404)

    def test_auditor_full_read_with_pay_and_ids(self):
        self.client.force_login(self.auditor_user)

        resp = self.client.get(f"/people/{self.person_a.pk}/")
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.context["can_view_rates"])
        self.assertTrue(resp.context["can_view_id_numbers"])
        self.assertTrue(resp.context["can_view_doc_content"])
        self.assertContains(resp, "000-111-222-000")
        self.assertContains(resp, "50000.00")


    def test_employee_self_service_privacy_and_access(self):
        self.client.force_login(self.emp_a_user)

        # Can view own record
        resp_own = self.client.get(f"/people/{self.person_a.pk}/")
        self.assertEqual(resp_own.status_code, 200)
        self.assertFalse(resp_own.context["can_view_rates"])
        self.assertTrue(resp_own.context["can_view_id_numbers"])
        self.assertContains(resp_own, "000-111-222-000")

        # Cannot view other personnel
        resp_other = self.client.get(f"/people/{self.person_b.pk}/")
        self.assertEqual(resp_other.status_code, 404)


class Step5AuditLogTestCase(TestCase):
    def setUp(self):
        User = get_user_model()
        call_command("seed_offices")
        call_command("setup_roles")

        self.superuser = User.objects.create_superuser("superadmin", "super@test.com", "superpass123")
        self.auditor = User.objects.create_user("auditor_user", "auditor@test.com", "audpass123")
        self.auditor.groups.add(Group.objects.get(name=ROLE_AUDITOR))

        self.hrmo = User.objects.create_user("hrmo_user", "hrmo@test.com", "hrmopass123")
        self.hrmo.groups.add(Group.objects.get(name=ROLE_HRMO))

        self.office_head = User.objects.create_user("head_user", "head@test.com", "headpass123")
        self.office_head.groups.add(Group.objects.get(name=ROLE_OFFICE_HEAD))

        self.office = Office.objects.get(name="Office of the Municipal Mayor")
        UserProfile.objects.create(user=self.office_head, office=self.office)

        self.person = Person.objects.create(
            employee_no="SP-00888",
            last_name="Mabini",
            first_name="Apolinario",
            birthdate=date(1985, 7, 23),
            tin="000-888-999-000",
        )

    def test_database_trigger_blocks_update_and_delete(self):
        # Create an audit log entry
        entry = log_audit_event(
            action=AuditLog.Action.CREATE,
            model_name="Person",
            object_id=self.person.pk,
            object_repr=str(self.person),
            changes={"first_name": {"old": None, "new": "Apolinario"}},
        )
        self.assertIsNotNone(entry)
        self.assertIsNotNone(entry.pk)

        # Attempting DELETE via SQL/ORM must fail at the PostgreSQL trigger level
        with self.assertRaises(DatabaseError):
            with transaction.atomic():
                AuditLog.objects.filter(pk=entry.pk).delete()

        # Attempting UPDATE via SQL/ORM must fail at the PostgreSQL trigger level
        with self.assertRaises(DatabaseError):
            with transaction.atomic():
                AuditLog.objects.filter(pk=entry.pk).update(object_repr="Tampered description")

    def test_sensitive_fields_redaction_in_changes_diff(self):
        # 1. Update Person with new fake TIN
        self.person.tin = "000-777-666-000"
        self.person.save()

        person_update_log = AuditLog.objects.filter(
            model_name="Person", object_id=str(self.person.pk), action=AuditLog.Action.UPDATE
        ).first()
        self.assertIsNotNone(person_update_log)
        # tin field name is logged, but its value is redacted
        self.assertIn("tin", person_update_log.changes)
        self.assertEqual(person_update_log.changes["tin"], "[REDACTED]")
        self.assertNotIn("000-777-666-000", str(person_update_log.changes))

        # 2. Create Engagement with pay rate
        eng = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.CASUAL,
            position_title="Legal Researcher",
            rate=Decimal("28000.00"),
            rate_basis=Engagement.RateBasis.MONTHLY,
            start_date=date(2025, 1, 1),
        )
        eng_log = AuditLog.objects.filter(
            model_name="Engagement", object_id=str(eng.pk), action=AuditLog.Action.CREATE
        ).first()
        self.assertIsNotNone(eng_log)
        self.assertIn("rate", eng_log.changes)
        self.assertEqual(eng_log.changes["rate"], "[REDACTED]")
        self.assertNotIn("28000.00", str(eng_log.changes))

        # 3. Password reset on User
        target_user = get_user_model().objects.create_user("pwd_target", "pwd@test.com", "initialpass123")
        target_user.set_password("newsecretpass456")
        target_user.save()

        pwd_log = AuditLog.objects.filter(
            model_name="User", object_id=str(target_user.pk), action=AuditLog.Action.PASSWORD_RESET
        ).first()
        self.assertIsNotNone(pwd_log)
        self.assertIn("password", pwd_log.changes)
        self.assertEqual(pwd_log.changes["password"], "[REDACTED]")
        self.assertNotIn("pbkdf2_sha256", str(pwd_log.changes))
        self.assertNotIn("newsecretpass456", str(pwd_log.changes))

    def test_user_deletion_does_not_affect_audit_log_rows(self):
        User = get_user_model()
        temp_user = User.objects.create_user("temp_user", "temp@test.com", "pass123")
        temp_uid = temp_user.pk

        entry = log_audit_event(
            action=AuditLog.Action.UPDATE,
            model_name="Person",
            object_id=self.person.pk,
            object_repr=str(self.person),
            user=temp_user,
        )

        self.assertEqual(entry.user_id, temp_uid)
        self.assertEqual(entry.username, "temp_user")

        # Delete user
        temp_user.delete()

        # AuditLog entry still exists intact
        reloaded = AuditLog.objects.get(pk=entry.pk)
        self.assertEqual(reloaded.user_id, temp_uid)
        self.assertEqual(reloaded.username, "temp_user")

    def test_transaction_binding_failure_rolls_back_action(self):
        # If audit log creation fails, the outer action transaction must roll back
        from unittest.mock import patch

        with self.assertRaises(RuntimeError):
            with transaction.atomic():
                p = Person(
                    employee_no="SP-00999",
                    last_name="Silang",
                    first_name="Gabriela",
                    birthdate=date(1991, 3, 19),
                )
                p.save()
                with patch("hr.models.AuditLog.objects.create", side_effect=RuntimeError("Audit DB down")):
                    log_audit_event(
                        action=AuditLog.Action.UPDATE,
                        model_name="Person",
                        object_id=p.pk,
                        object_repr=str(p),
                    )

        # Confirm person was rolled back and never persisted
        self.assertFalse(Person.objects.filter(employee_no="SP-00999").exists())

    def test_trusted_proxy_client_ip_resolution(self):
        rf = RequestFactory()

        # Untrusted remote address: X-Forwarded-For must be IGNORED
        req_untrusted = rf.get("/people/", REMOTE_ADDR="198.51.100.42", HTTP_X_FORWARDED_FOR="1.1.1.1")
        self.assertEqual(get_client_ip(req_untrusted), "198.51.100.42")

        # Trusted proxy: the rightmost hop that is not a trusted proxy is the client.
        # "203.0.113.195" was supplied by the client itself and must not be believed.
        req_trusted = rf.get("/people/", REMOTE_ADDR="127.0.0.1", HTTP_X_FORWARDED_FOR="203.0.113.195, 10.0.0.1")
        self.assertEqual(get_client_ip(req_trusted), "10.0.0.1")

        # When the next hop is itself a trusted proxy, keep walking left.
        with override_settings(TRUSTED_PROXIES=["127.0.0.1", "10.0.0.0/8"]):
            self.assertEqual(get_client_ip(req_trusted), "203.0.113.195")

        # A malformed hop stops the walk at the last trustworthy address.
        req_garbage = rf.get("/people/", REMOTE_ADDR="127.0.0.1", HTTP_X_FORWARDED_FOR="<script>")
        self.assertEqual(get_client_ip(req_garbage), "127.0.0.1")

        # IPv6 is stored in PostgreSQL's canonical form, so the row hash survives a round trip.
        req_v6 = rf.get("/people/", REMOTE_ADDR="127.0.0.1", HTTP_X_FORWARDED_FOR="2001:DB8:0:0:0:0:0:1")
        self.assertEqual(get_client_ip(req_v6), "2001:db8::1")
        log_audit_event(action=AuditLog.Action.VIEW, model_name="Ip", request=req_v6)
        self.assertTrue(verify_chain().is_valid)

    def test_selective_view_logging(self):
        self.client.force_login(self.hrmo)

        # 1. Directory listing (/people/) must NOT log VIEW
        init_view_count = AuditLog.objects.filter(action=AuditLog.Action.VIEW).count()
        self.client.get("/people/")
        self.assertEqual(AuditLog.objects.filter(action=AuditLog.Action.VIEW).count(), init_view_count)

        # 2. Offices directory (/offices/) must NOT log VIEW
        self.client.get("/offices/")
        self.assertEqual(AuditLog.objects.filter(action=AuditLog.Action.VIEW).count(), init_view_count)

        # 3. Person detail (/people/<id>/) MUST log VIEW
        resp_detail = self.client.get(f"/people/{self.person.pk}/")
        self.assertEqual(resp_detail.status_code, 200)
        self.assertEqual(
            AuditLog.objects.filter(
                action=AuditLog.Action.VIEW, model_name="Person", object_id=str(self.person.pk)
            ).count(),
            1,
        )

    def test_auditor_interface_access_and_export(self):
        # Non-auditor (Office Head) gets 403 Forbidden on /audit/
        self.client.force_login(self.office_head)
        resp_denied = self.client.get("/audit/")
        self.assertEqual(resp_denied.status_code, 403)

        # 403 denial was logged as PERMISSION_DENIED
        denial_log = AuditLog.objects.filter(action=AuditLog.Action.PERMISSION_DENIED).first()
        self.assertIsNotNone(denial_log)
        self.assertEqual(denial_log.changes.get("path"), "/audit/")

        # Auditor gets 200 OK
        self.client.force_login(self.auditor)
        resp_audit = self.client.get("/audit/")
        self.assertEqual(resp_audit.status_code, 200)
        self.assertContains(resp_audit, "System Audit Trail")

        # Auditor exports CSV
        resp_export = self.client.get("/audit/export/csv/")
        self.assertEqual(resp_export.status_code, 200)
        self.assertEqual(resp_export["Content-Type"], "text/csv")
        self.assertIn("Timestamp,Action,User ID", resp_export.content.decode("utf-8"))

        # EXPORT action was recorded in AuditLog
        self.assertTrue(
            AuditLog.objects.filter(action=AuditLog.Action.EXPORT, model_name="AuditLog").exists()
        )

    def test_failed_login_records_username_only_for_existing_account(self):
        # 1. Failed login for existing account: records the actual username
        self.client.post("/accounts/login/", {"username": "auditor_user", "password": "wrongpassword"})
        log_existing = AuditLog.objects.filter(
            action=AuditLog.Action.FAILED_LOGIN, changes__username_attempted="auditor_user"
        ).first()
        self.assertIsNotNone(log_existing)
        self.assertEqual(log_existing.username, "auditor_user")
        self.assertIn("auditor_user", log_existing.object_repr)

        # 2. Failed login for non-existent account: records "[unknown]" to prevent credential leaks
        self.client.post("/accounts/login/", {"username": "nonexistent_stranger", "password": "somepassword"})
        log_unknown = AuditLog.objects.filter(
            action=AuditLog.Action.FAILED_LOGIN, changes__username_attempted="[unknown]"
        ).first()
        self.assertIsNotNone(log_unknown)
        self.assertEqual(log_unknown.username, "[unknown]")
        self.assertIn("[unknown]", log_unknown.object_repr)
        self.assertNotIn("nonexistent_stranger", str(log_unknown.changes))

    def test_permission_denied_deduplication(self):
        # Office head accesses protected /audit/ repeatedly
        self.client.force_login(self.office_head)

        # 1. First request creates PERMISSION_DENIED entry
        resp1 = self.client.get("/audit/")
        self.assertEqual(resp1.status_code, 403)
        count_after_first = AuditLog.objects.filter(
            action=AuditLog.Action.PERMISSION_DENIED,
            user_id=self.office_head.pk,
            changes__path="/audit/",
        ).count()
        self.assertEqual(count_after_first, 1)

        # 2. Second request within the same minute does NOT create a duplicate entry
        resp2 = self.client.get("/audit/")
        self.assertEqual(resp2.status_code, 403)
        count_after_second = AuditLog.objects.filter(
            action=AuditLog.Action.PERMISSION_DENIED,
            user_id=self.office_head.pk,
            changes__path="/audit/",
        ).count()
        self.assertEqual(count_after_second, 1)

        # 3. Request to a different forbidden path DOES create a new entry
        doc = Document.objects.create(
            person=self.person,
            doc_type=Document.Type.APPOINTMENT,
            physical_location="Cabinet A",
        )
        resp3 = self.client.get(f"/documents/{doc.pk}/download/")
        self.assertEqual(resp3.status_code, 403)
        count_doc_denied = AuditLog.objects.filter(
            action=AuditLog.Action.PERMISSION_DENIED,
            user_id=self.office_head.pk,
            changes__path=f"/documents/{doc.pk}/download/",
        ).count()
        self.assertEqual(count_doc_denied, 1)

        # 4. A different user accessing /audit/ creates a separate entry for that user
        self.client.force_login(self.hrmo)
        resp4 = self.client.get("/audit/")
        self.assertEqual(resp4.status_code, 403)
        count_hrmo_denied = AuditLog.objects.filter(
            action=AuditLog.Action.PERMISSION_DENIED,
            user_id=self.hrmo.pk,
            changes__path="/audit/",
        ).count()
        self.assertEqual(count_hrmo_denied, 1)

    def test_trusted_proxies_environment_setting(self):
        from django.test import override_settings

        with override_settings(TRUSTED_PROXIES=["127.0.0.1", "::1", "172.18.0.0/16"]):
            self.assertTrue(is_ip_trusted("127.0.0.1"))
            self.assertTrue(is_ip_trusted("::1"))
            self.assertTrue(is_ip_trusted("172.18.0.42"))
            self.assertFalse(is_ip_trusted("192.168.1.1"))

    def test_codebase_inspection_forbids_unlogged_queryset_mutations(self):
        # Scans hr/ app views and domain code to ensure no .update(), .bulk_create(), or raw SQL are used
        from pathlib import Path

        hr_dir = Path(__file__).resolve().parent
        forbidden_calls = {"bulk_create", "bulk_update", "raw"}

        for py_file in hr_dir.glob("*.py"):
            if py_file.name in ["tests.py", "audit.py", "signals.py"]:
                continue

            with open(py_file, "r", encoding="utf-8") as f:
                tree = ast.parse(f.read(), filename=str(py_file))

            for node in ast.walk(tree):
                if isinstance(node, ast.Call):
                    if isinstance(node.func, ast.Attribute):
                        call_name = node.func.attr
                        # Disallow bulk_create(), bulk_update(), raw() in domain code
                        if call_name in forbidden_calls:
                            self.fail(
                                f"Forbidden call '{call_name}' in {py_file.name} line {node.lineno}: "
                                "Domain operations must use model save() or delete() to guarantee audit trail."
                            )


class Step6APersonEntryTestCase(TestCase):
    def setUp(self):
        User = get_user_model()
        call_command("seed_offices")
        call_command("setup_roles")

        self.hrmo_user = User.objects.create_user("hrmo_editor", "hrmo@test.com", "pass123")
        self.hrmo_user.groups.add(Group.objects.get(name=ROLE_HRMO))

        self.sysadmin = User.objects.create_user("sys_admin", "sys@test.com", "pass123", is_staff=True)
        self.sysadmin.groups.add(Group.objects.get(name=ROLE_SYSADMIN))

        self.office_head = User.objects.create_user("head_user", "head@test.com", "pass123")
        self.office_head.groups.add(Group.objects.get(name=ROLE_OFFICE_HEAD))

        self.auditor = User.objects.create_user("audit_user", "audit@test.com", "pass123")
        self.auditor.groups.add(Group.objects.get(name=ROLE_AUDITOR))

        self.executive = User.objects.create_user("exec_user", "exec@test.com", "pass123")
        self.executive.groups.add(Group.objects.get(name=ROLE_EXECUTIVE))

        self.employee = User.objects.create_user("reg_emp", "emp@test.com", "pass123")

        self.existing_person = Person.objects.create(
            employee_no="SP-00001",
            last_name="Del Rosario",
            first_name="Carmen",
            middle_name="Santos",
            suffix="",
            sex=Person.Sex.FEMALE,
            birthdate=date(1990, 5, 20),
            tin="000-111-222-000",
            version=1,
        )

    def test_non_hrmo_roles_cannot_access_person_create_or_edit(self):
        non_hrmo_users = [
            self.sysadmin,
            self.office_head,
            self.auditor,
            self.executive,
            self.employee,
        ]

        for user in non_hrmo_users:
            self.client.force_login(user)

            # GET create URL -> 403 Forbidden
            resp_create_get = self.client.get("/people/new/")
            self.assertEqual(
                resp_create_get.status_code,
                403,
                f"Role {user.username} must get 403 on GET /people/new/",
            )

            # POST create URL -> 403 Forbidden
            resp_create_post = self.client.post("/people/new/", {"last_name": "Test"})
            self.assertEqual(
                resp_create_post.status_code,
                403,
                f"Role {user.username} must get 403 on POST /people/new/",
            )

            # GET edit URL -> 403 Forbidden
            resp_edit_get = self.client.get(f"/people/{self.existing_person.pk}/edit/")
            self.assertEqual(
                resp_edit_get.status_code,
                403,
                f"Role {user.username} must get 403 on GET /people/{self.existing_person.pk}/edit/",
            )

            # POST edit URL -> 403 Forbidden
            resp_edit_post = self.client.post(
                f"/people/{self.existing_person.pk}/edit/", {"last_name": "Hacked"}
            )
            self.assertEqual(
                resp_edit_post.status_code,
                403,
                f"Role {user.username} must get 403 on POST /people/{self.existing_person.pk}/edit/",
            )

    def test_hrmo_can_create_person_with_race_safe_employee_no_and_audit_log(self):
        self.client.force_login(self.hrmo_user)

        # GET form
        resp_get = self.client.get("/people/new/")
        self.assertEqual(resp_get.status_code, 200)
        self.assertContains(resp_get, "PROTOTYPE - FAKE DATA ONLY")
        self.assertContains(resp_get, "Add New Personnel Record")

        # POST form
        payload = {
            "last_name": "  Luna  ",
            "first_name": " Antonio ",
            "middle_name": " Novicio ",
            "suffix": "",
            "sex": "M",
            "birthdate": "1988-10-29",
            "mobile": " 0917-888-9999 ",
            "email": "antonio.luna@sanpascual.gov.ph",
            "address": "San Pascual, Batangas",
            "tin": "000-999-888-000",
            "gsis_no": "0012345678",
            "philhealth_no": "00-123456789-0",
            "pagibig_no": "0012-3456-7890",
            "sss_no": "00-1234567-8",
        }
        resp_post = self.client.post("/people/new/", payload)
        self.assertEqual(resp_post.status_code, 302)

        new_person = Person.objects.get(last_name="Luna")
        self.assertEqual(new_person.first_name, "Antonio")
        self.assertEqual(new_person.middle_name, "Novicio")
        self.assertEqual(new_person.mobile, "0917-888-9999")
        self.assertTrue(new_person.employee_no.startswith("SP-"))
        self.assertEqual(new_person.version, 1)

        # Audit log verifies CREATE event with sensitive IDs redacted
        audit_entry = AuditLog.objects.filter(
            model_name="Person", object_id=str(new_person.pk), action=AuditLog.Action.CREATE
        ).first()
        self.assertIsNotNone(audit_entry)
        self.assertIn("tin", audit_entry.changes)
        self.assertEqual(audit_entry.changes["tin"], "[REDACTED]")
        self.assertNotIn("000-999-888-000", str(audit_entry.changes))

    def test_hrmo_can_edit_person_with_version_increment(self):
        self.client.force_login(self.hrmo_user)

        # Edit person
        payload = {
            "last_name": "Del Rosario",
            "first_name": "Carmen",
            "middle_name": "Santos",
            "suffix": "",
            "sex": "F",
            "birthdate": "1990-05-20",
            "mobile": "0918-111-2222",
            "email": "carmen.delrosario@sanpascual.gov.ph",
            "address": "Poblacion, San Pascual",
            "tin": "000-111-222-000",
            "gsis_no": "0011223344",
            "philhealth_no": "00-223344556-0",
            "pagibig_no": "0011-2233-4455",
            "sss_no": "00-9988776-5",
            "version": self.existing_person.version,
        }
        resp = self.client.post(f"/people/{self.existing_person.pk}/edit/", payload)
        self.assertEqual(resp.status_code, 302)

        self.existing_person.refresh_from_db()
        self.assertEqual(self.existing_person.mobile, "0918-111-2222")
        self.assertEqual(self.existing_person.version, 2)

    def test_lost_update_protection_rejects_stale_submission(self):
        self.client.force_login(self.hrmo_user)

        # Stale submission: submitted version is 1, but DB version is already 2
        self.existing_person.version = 2
        self.existing_person.save()

        payload = {
            "last_name": "Del Rosario",
            "first_name": "Carmen",
            "middle_name": "Santos",
            "suffix": "",
            "sex": "F",
            "birthdate": "1990-05-20",
            "mobile": "0918-999-9999",
            "version": 1,  # Stale version!
        }
        resp = self.client.post(f"/people/{self.existing_person.pk}/edit/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "modified by another user")

        # Database record was NOT modified
        self.existing_person.refresh_from_db()
        self.assertNotEqual(self.existing_person.mobile, "0918-999-9999")
        self.assertEqual(self.existing_person.version, 2)

    def test_duplicate_person_validation_considers_middle_name_and_suffix(self):
        self.client.force_login(self.hrmo_user)

        # Attempt to create duplicate of existing_person
        payload = {
            "last_name": "del rosario",
            "first_name": "CARMEN",
            "middle_name": "santos",
            "suffix": "",
            "sex": "F",
            "birthdate": "1990-05-20",
        }
        resp = self.client.post("/people/new/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "already exists")
        self.assertContains(resp, self.existing_person.employee_no)

        # Distinct middle name succeeds
        payload["middle_name"] = "Garcia"
        resp_distinct = self.client.post("/people/new/", payload)
        self.assertEqual(resp_distinct.status_code, 302)

    def test_birthdate_age_validation(self):
        self.client.force_login(self.hrmo_user)

        # 1. Future birthdate
        resp_future = self.client.post("/people/new/", {
            "last_name": "Future",
            "first_name": "Baby",
            "birthdate": "2030-01-01",
        })
        self.assertContains(resp_future, "Birthdate cannot be in the future")

        # 2. Too young (< 18)
        resp_young = self.client.post("/people/new/", {
            "last_name": "Young",
            "first_name": "Child",
            "birthdate": "2015-01-01",
        })
        self.assertContains(resp_young, "at least 18 years old")

        # 3. Too old (> 90)
        resp_old = self.client.post("/people/new/", {
            "last_name": "Elder",
            "first_name": "Senior",
            "birthdate": "1920-01-01",
        })
        self.assertContains(resp_old, "cannot exceed 90 years")

    def test_prototype_fake_id_rule(self):
        from django.test import override_settings

        self.client.force_login(self.hrmo_user)

        # With PROTOTYPE_FAKE_DATA_ONLY=True, non-00 ID is rejected
        with override_settings(PROTOTYPE_FAKE_DATA_ONLY=True):
            resp = self.client.post("/people/new/", {
                "last_name": "RealID",
                "first_name": "Tester",
                "birthdate": "1990-01-01",
                "tin": "123-456-789-000",  # Invalid (doesn't start with 00)
            })
            self.assertContains(resp, "must start with 00")

    def test_field_level_gating_removes_id_fields_without_view_ids(self):
        # User without view_ids has ID fields completely absent from PersonForm
        form = PersonForm(user=self.employee)
        for f in ["tin", "gsis_no", "philhealth_no", "pagibig_no", "sss_no"]:
            self.assertNotIn(f, form.fields)

        # User with view_ids (HRMO) has ID fields present
        hrmo_form = PersonForm(user=self.hrmo_user)
        for f in ["tin", "gsis_no", "philhealth_no", "pagibig_no", "sss_no"]:
            self.assertIn(f, hrmo_form.fields)

    def test_employee_sequence_generation_and_resync(self):
        from hr.sequences import get_next_employee_no, sync_employee_sequence

        # Sequential generation
        no1 = get_next_employee_no()
        no2 = get_next_employee_no()
        self.assertTrue(no1.startswith("SP-"))
        self.assertTrue(no2.startswith("SP-"))
        num1 = int(no1.split("-")[1])
        num2 = int(no2.split("-")[1])
        self.assertEqual(num2, num1 + 1)

        # Create person using no2 to verify sync_employee_sequence finds it
        Person.objects.create(
            employee_no=no2,
            last_name="Sequence",
            first_name="Test",
            birthdate=date(1995, 1, 1),
        )
        synced = sync_employee_sequence()
        self.assertGreaterEqual(synced, num2)


class Step6BEngagementEntryTestCase(TestCase):
    def setUp(self):
        User = get_user_model()
        call_command("seed_offices")
        call_command("setup_roles")

        self.hrmo_user = User.objects.create_user("hrmo_editor", "hrmo@test.com", "pass123")
        self.hrmo_user.groups.add(Group.objects.get(name=ROLE_HRMO))

        self.sysadmin = User.objects.create_user("sys_admin", "sys@test.com", "pass123", is_staff=True)
        self.sysadmin.groups.add(Group.objects.get(name=ROLE_SYSADMIN))

        self.office_head = User.objects.create_user("head_user", "head@test.com", "pass123")
        self.office_head.groups.add(Group.objects.get(name=ROLE_OFFICE_HEAD))

        self.auditor = User.objects.create_user("audit_user", "audit@test.com", "pass123")
        self.auditor.groups.add(Group.objects.get(name=ROLE_AUDITOR))

        self.executive = User.objects.create_user("exec_user", "exec@test.com", "pass123")
        self.executive.groups.add(Group.objects.get(name=ROLE_EXECUTIVE))

        self.employee = User.objects.create_user("reg_emp", "emp@test.com", "pass123")

        self.office = Office.objects.get(name="Office of the Municipal Mayor")
        self.person = Person.objects.create(
            employee_no="SP-00050",
            last_name="Silang",
            first_name="Diego",
            birthdate=date(1985, 12, 16),
            tin="000-555-666-000",
            version=1,
        )

        self.engagement = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.CASUAL,
            position_title="Administrative Aide IV",
            rate=Decimal("18000.00"),
            rate_basis=Engagement.RateBasis.MONTHLY,
            start_date=date(2025, 1, 1),
            end_date=date(2025, 12, 31),
            status=Engagement.Status.ACTIVE,
            leave_eligible=True,
            version=1,
        )

        self.project = Project.objects.create(
            name="San Pascual Digital Portal",
            office=self.office,
            start_date=date(2025, 1, 1),
            funding_reference="FAKE-CAF-2025-900",  # required to hire against it (DECISIONS P2)
        )

    def test_non_hrmo_roles_cannot_access_engagement_actions(self):
        non_hrmo_users = [
            self.sysadmin,
            self.office_head,
            self.auditor,
            self.executive,
            self.employee,
        ]

        for user in non_hrmo_users:
            self.client.force_login(user)

            # Create URL
            resp_create = self.client.get(f"/people/{self.person.pk}/engagements/new/")
            self.assertEqual(resp_create.status_code, 403, f"{user.username} must get 403 on create")

            # Edit URL
            resp_edit = self.client.get(f"/engagements/{self.engagement.pk}/edit/")
            self.assertEqual(resp_edit.status_code, 403, f"{user.username} must get 403 on edit")

            # End URL
            resp_end = self.client.get(f"/engagements/{self.engagement.pk}/end/")
            self.assertEqual(resp_end.status_code, 403, f"{user.username} must get 403 on end")

            # Renew URL
            resp_renew = self.client.get(f"/engagements/{self.engagement.pk}/renew/")
            self.assertEqual(resp_renew.status_code, 403, f"{user.username} must get 403 on renew")

    def test_hrmo_can_create_plantilla_engagement_with_audit_redaction(self):
        self.client.force_login(self.hrmo_user)

        payload = {
            "office": self.office.pk,
            "type": Engagement.Type.PLANTILLA,
            "position_title": "Administrative Officer V",
            "rate": "46725.00",
            "rate_basis": Engagement.RateBasis.MONTHLY,
            "start_date": "2026-01-01",
            "end_date": "",
            "status": Engagement.Status.ACTIVE,
            "leave_eligible": True,
            "salary_grade": 18,
            "step": 1,
            "item_number": "SP-MO-018",
        }
        resp = self.client.post(f"/people/{self.person.pk}/engagements/new/", payload)
        self.assertEqual(resp.status_code, 302)

        new_eng = Engagement.objects.filter(item_number="SP-MO-018").first()
        self.assertIsNotNone(new_eng)
        self.assertEqual(new_eng.salary_grade, 18)
        self.assertEqual(new_eng.version, 1)

        # Audit log verifies CREATE event with rate redacted
        audit = AuditLog.objects.filter(
            model_name="Engagement", object_id=str(new_eng.pk), action=AuditLog.Action.CREATE
        ).first()
        self.assertIsNotNone(audit)
        self.assertEqual(audit.changes.get("rate"), "[REDACTED]")
        self.assertNotIn("46725.00", str(audit.changes))

    def test_plantilla_validation_requires_sg_step_and_item(self):
        self.client.force_login(self.hrmo_user)

        # Missing plantilla fields fails validation
        payload = {
            "office": self.office.pk,
            "type": Engagement.Type.PLANTILLA,
            "position_title": "Officer",
            "start_date": "2026-01-01",
            "status": Engagement.Status.ACTIVE,
        }
        resp = self.client.post(f"/people/{self.person.pk}/engagements/new/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Salary grade, step, and item number are required")

    def test_non_plantilla_forbids_salary_grade_and_item(self):
        self.client.force_login(self.hrmo_user)

        payload = {
            "office": self.office.pk,
            "type": Engagement.Type.JOB_ORDER,
            "position_title": "Encoder",
            "start_date": "2026-01-01",
            "status": Engagement.Status.ACTIVE,
            "salary_grade": 5,
            "step": 1,
            "item_number": "SP-JO-05",
        }
        resp = self.client.post(f"/people/{self.person.pk}/engagements/new/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "must be empty for non-Plantilla")

    def test_project_based_validation(self):
        self.client.force_login(self.hrmo_user)

        # Project-based without project fails
        payload = {
            "office": self.office.pk,
            "type": Engagement.Type.PROJECT_BASED,
            "position_title": "Civil Engineer",
            "start_date": "2026-01-01",
            "status": Engagement.Status.ACTIVE,
        }
        resp = self.client.post(f"/people/{self.person.pk}/engagements/new/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "require an associated project")

        # With project succeeds
        payload["project"] = self.project.pk
        resp_valid = self.client.post(f"/people/{self.person.pk}/engagements/new/", payload)
        self.assertEqual(resp_valid.status_code, 302)

    def test_end_date_before_start_date_validation(self):
        self.client.force_login(self.hrmo_user)

        payload = {
            "office": self.office.pk,
            "type": Engagement.Type.JOB_ORDER,
            "position_title": "Clerk",
            "start_date": "2026-06-01",
            "end_date": "2026-05-01",
            "status": Engagement.Status.ACTIVE,
        }
        resp = self.client.post(f"/people/{self.person.pk}/engagements/new/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "cannot be earlier than start date")

    def test_overlap_warning_requires_confirmation_tick(self):
        self.client.force_login(self.hrmo_user)

        # self.engagement is CASUAL in self.office from 2025-01-01 to 2025-12-31
        # Attempting to add another overlapping CASUAL engagement without confirmation fails with warning
        payload = {
            "office": self.office.pk,
            "type": Engagement.Type.CASUAL,
            "position_title": "Overlapping Aide",
            "start_date": "2025-06-01",
            "end_date": "2025-11-30",
            "status": Engagement.Status.ACTIVE,
        }
        resp = self.client.post(f"/people/{self.person.pk}/engagements/new/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Warning:")
        self.assertContains(resp, "Check &#x27;I understand, save anyway&#x27; to proceed")

        # Ticking confirm_overlap allows save
        payload["confirm_overlap"] = True
        resp_confirm = self.client.post(f"/people/{self.person.pk}/engagements/new/", payload)
        self.assertEqual(resp_confirm.status_code, 302)
        self.assertTrue(Engagement.objects.filter(position_title="Overlapping Aide").exists())

    def test_end_engagement_action(self):
        self.client.force_login(self.hrmo_user)

        # GET confirm screen
        resp_get = self.client.get(f"/engagements/{self.engagement.pk}/end/")
        self.assertEqual(resp_get.status_code, 200)
        self.assertContains(resp_get, "End Appointment / Contract")
        self.assertContains(resp_get, self.engagement.position_title)

        # POST confirm
        resp_post = self.client.post(
            f"/engagements/{self.engagement.pk}/end/",
            {
                "end_date": "2025-12-15",
                "notes": "Contract wrapped up early",
                "version": self.engagement.version,
            },
        )
        self.assertEqual(resp_post.status_code, 302)

        self.engagement.refresh_from_db()
        self.assertEqual(self.engagement.status, Engagement.Status.ENDED)
        self.assertEqual(self.engagement.end_date, date(2025, 12, 15))
        self.assertIn("Contract wrapped up early", self.engagement.notes)
        self.assertEqual(self.engagement.version, 2)

    def test_renew_creates_new_engagement_preserving_old(self):
        self.client.force_login(self.hrmo_user)

        # GET renew form (pre-filled from self.engagement)
        resp_get = self.client.get(f"/engagements/{self.engagement.pk}/renew/")
        self.assertEqual(resp_get.status_code, 200)
        self.assertContains(resp_get, "Renew / New Appointment")
        # start_date pre-filled to day after old end_date (2025-12-31 -> 2026-01-01)
        self.assertContains(resp_get, 'value="2026-01-01"')

        # POST creates new appointment
        payload = {
            "office": self.office.pk,
            "type": Engagement.Type.PLANTILLA,  # Career progression to Plantilla!
            "position_title": "Administrative Aide IV (Permanent)",
            "rate": "18000.00",
            "rate_basis": Engagement.RateBasis.MONTHLY,
            "start_date": "2026-01-01",
            "status": Engagement.Status.ACTIVE,
            "leave_eligible": True,
            "salary_grade": 4,
            "step": 1,
            "item_number": "SP-MO-004",
        }
        resp_post = self.client.post(f"/engagements/{self.engagement.pk}/renew/", payload)
        self.assertEqual(resp_post.status_code, 302)

        # Old engagement remains unchanged
        self.engagement.refresh_from_db()
        self.assertEqual(self.engagement.type, Engagement.Type.CASUAL)
        self.assertEqual(self.engagement.position_title, "Administrative Aide IV")

        # New engagement created
        new_eng = Engagement.objects.get(item_number="SP-MO-004")
        self.assertEqual(new_eng.type, Engagement.Type.PLANTILLA)
        self.assertEqual(new_eng.person, self.person)

    def test_lost_update_protection_on_engagement(self):
        self.client.force_login(self.hrmo_user)

        self.engagement.version = 2
        self.engagement.save()

        # Submit with stale version 1
        payload = {
            "office": self.office.pk,
            "type": self.engagement.type,
            "position_title": "Hacked Title",
            "start_date": self.engagement.start_date.isoformat(),
            "status": Engagement.Status.ACTIVE,
            "version": 1,
        }
        resp = self.client.post(f"/engagements/{self.engagement.pk}/edit/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "modified by another user")

        self.engagement.refresh_from_db()
        self.assertNotEqual(self.engagement.position_title, "Hacked Title")

    def test_pay_rate_fields_stripped_without_view_pay(self):
        # Regular employee has no view_pay permission
        form = EngagementForm(user=self.employee, person=self.person)
        self.assertNotIn("rate", form.fields)
        self.assertNotIn("rate_basis", form.fields)

        # HRMO staff has view_pay permission
        hrmo_form = EngagementForm(user=self.hrmo_user, person=self.person)
        self.assertIn("rate", hrmo_form.fields)
        self.assertIn("rate_basis", hrmo_form.fields)


class Step6CDocumentEntryTestCase(TestCase):
    """Test suite for Document data-entry, file validation, storage, and replacement (Step 6C)."""

    def setUp(self):
        call_command("seed_offices")
        call_command("setup_roles")

        User = get_user_model()
        self.hrmo_user = User.objects.create_user(username="hrmo_doc_user", password="Password123!")
        self.hrmo_user.groups.add(Group.objects.get(name=ROLE_HRMO))

        self.sys_admin = User.objects.create_user(username="sys_admin_doc", password="Password123!")
        self.sys_admin.groups.add(Group.objects.get(name=ROLE_SYSADMIN))

        self.office_head = User.objects.create_user(username="head_doc", password="Password123!")
        self.office_head.groups.add(Group.objects.get(name=ROLE_OFFICE_HEAD))

        self.executive = User.objects.create_user(username="exec_doc", password="Password123!")
        self.executive.groups.add(Group.objects.get(name=ROLE_EXECUTIVE))

        self.auditor = User.objects.create_user(username="audit_doc", password="Password123!")
        self.auditor.groups.add(Group.objects.get(name=ROLE_AUDITOR))

        self.employee = User.objects.create_user(username="emp_doc", password="Password123!")

        self.office = Office.objects.get(acronym="MTO")
        UserProfile.objects.create(user=self.office_head, office=self.office)

        self.person = Person.objects.create(
            employee_no="SP-00010",
            last_name="Bautista",
            first_name="Ramon",
            middle_name="Santos",
            suffix="",
            sex=Person.Sex.MALE,
            birthdate=date(1988, 7, 20),
            version=1,
        )
        self.engagement = Engagement.objects.create(
            person=self.person,
            office=self.office,
            type=Engagement.Type.CASUAL,
            position_title="Revenue Collector I",
            start_date=date(2025, 1, 1),
            status=Engagement.Status.ACTIVE,
            version=1,
        )
        self.sample_doc = Document.objects.create(
            person=self.person,
            doc_type=Document.Type.APPOINTMENT,
            status=Document.Status.RECEIVED,
            original_on_file=True,
            physical_location="Cab 1, D 2",
            version=1,
        )

    def _make_image_bytes(self, img_format="JPEG"):
        buf = BytesIO()
        img = Image.new("RGB", (30, 30), color="teal")
        img.save(buf, format=img_format)
        buf.seek(0)
        return buf.getvalue()

    def test_non_hrmo_roles_cannot_access_document_actions(self):
        non_hrmo_users = [
            self.sys_admin,
            self.office_head,
            self.executive,
            self.auditor,
            self.employee,
        ]
        for user in non_hrmo_users:
            self.client.force_login(user)

            # Create for person
            resp = self.client.get(f"/people/{self.person.pk}/documents/new/")
            self.assertEqual(resp.status_code, 403, f"{user.username} must get 403 on person doc create")

            # Create for engagement
            resp = self.client.get(f"/engagements/{self.engagement.pk}/documents/new/")
            self.assertEqual(resp.status_code, 403, f"{user.username} must get 403 on engagement doc create")

            # Edit doc
            resp = self.client.get(f"/documents/{self.sample_doc.pk}/edit/")
            self.assertEqual(resp.status_code, 403, f"{user.username} must get 403 on doc edit")

            # Replace doc
            resp = self.client.get(f"/documents/{self.sample_doc.pk}/replace/")
            self.assertEqual(resp.status_code, 403, f"{user.username} must get 403 on doc replace")

    def test_hrmo_can_create_paper_only_document(self):
        self.client.force_login(self.hrmo_user)

        payload = {
            "doc_type": Document.Type.SERVICE_RECORD,
            "status": Document.Status.FILED,
            "original_on_file": True,
            "physical_location": "201 Cabinet A, Folder Bautista",
            "notes": "Original wet-signed service record archived.",
        }
        resp = self.client.post(f"/people/{self.person.pk}/documents/new/", payload)
        self.assertEqual(resp.status_code, 302)

        doc = Document.objects.filter(doc_type=Document.Type.SERVICE_RECORD, person=self.person).first()
        self.assertIsNotNone(doc)
        self.assertEqual(doc.version, 1)
        self.assertTrue(doc.original_on_file)
        self.assertEqual(doc.physical_location, "201 Cabinet A, Folder Bautista")
        self.assertFalse(bool(doc.file))

    def test_hrmo_can_upload_pdf_with_uuid_storage_and_sanitized_name(self):
        self.client.force_login(self.hrmo_user)

        pdf_bytes = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF"
        uploaded_file = SimpleUploadedFile("Bautista Ramon PDS (Draft 2026).pdf", pdf_bytes, content_type="application/pdf")

        payload = {
            "doc_type": Document.Type.PDS,
            "status": Document.Status.APPROVED,
            "original_on_file": True,
            "physical_location": "Cabinet 2, Drawer 1",
            "file": uploaded_file,
        }
        resp = self.client.post(f"/people/{self.person.pk}/documents/new/", payload)
        self.assertEqual(resp.status_code, 302)

        doc = Document.objects.filter(doc_type=Document.Type.PDS, person=self.person).first()
        self.assertIsNotNone(doc)
        self.assertEqual(doc.display_name, "Bautista Ramon PDS (Draft 2026).pdf")
        # Ensure stored file path is random UUID, not the original filename
        self.assertNotIn("Bautista", doc.file.name)
        self.assertTrue(doc.file.name.endswith(".pdf"))
        self.assertIn("documents/", doc.file.name)

    def test_upload_validation_rejects_unsupported_extension(self):
        self.client.force_login(self.hrmo_user)

        bad_file = SimpleUploadedFile("script.sh", b"#!/bin/bash\necho bad\n", content_type="application/x-sh")
        payload = {
            "doc_type": Document.Type.OTHER,
            "status": Document.Status.RECEIVED,
            "file": bad_file,
        }
        resp = self.client.post(f"/people/{self.person.pk}/documents/new/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Unsupported file format")

    def test_upload_validation_rejects_oversized_file(self):
        from django.test import override_settings

        self.client.force_login(self.hrmo_user)

        # Set MAX_UPLOAD_SIZE to 50 bytes for test
        with override_settings(MAX_UPLOAD_SIZE=50):
            pdf_bytes = b"%PDF-1.4 " + (b"A" * 100)
            uploaded_file = SimpleUploadedFile("oversized.pdf", pdf_bytes, content_type="application/pdf")
            payload = {
                "doc_type": Document.Type.PDS,
                "status": Document.Status.RECEIVED,
                "file": uploaded_file,
            }
            resp = self.client.post(f"/people/{self.person.pk}/documents/new/", payload)
            self.assertEqual(resp.status_code, 200)
            self.assertContains(resp, "File exceeds maximum upload size")

    def test_magic_bytes_check_rejects_spoofed_pdf(self):
        self.client.force_login(self.hrmo_user)

        # File named .pdf but containing plain text
        spoofed = SimpleUploadedFile("fake.pdf", b"This is plain text pretending to be a PDF file.", content_type="application/pdf")
        payload = {
            "doc_type": Document.Type.PDS,
            "status": Document.Status.RECEIVED,
            "file": spoofed,
        }
        resp = self.client.post(f"/people/{self.person.pk}/documents/new/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "invalid magic bytes signature")

    def test_image_validation_with_pillow_for_jpeg_and_png(self):
        self.client.force_login(self.hrmo_user)

        # 1. Valid JPEG
        jpg_bytes = self._make_image_bytes(img_format="JPEG")
        jpg_file = SimpleUploadedFile("scan.jpg", jpg_bytes, content_type="image/jpeg")
        resp_jpg = self.client.post(
            f"/people/{self.person.pk}/documents/new/",
            {"doc_type": Document.Type.CERTIFICATE, "status": Document.Status.RECEIVED, "file": jpg_file},
        )
        self.assertEqual(resp_jpg.status_code, 302)

        # 2. Valid PNG
        png_bytes = self._make_image_bytes(img_format="PNG")
        png_file = SimpleUploadedFile("scan.png", png_bytes, content_type="image/png")
        resp_png = self.client.post(
            f"/people/{self.person.pk}/documents/new/",
            {"doc_type": Document.Type.CLEARANCE, "status": Document.Status.RECEIVED, "file": png_file},
        )
        self.assertEqual(resp_png.status_code, 302)

    def test_image_validation_rejects_corrupted_image(self):
        self.client.force_login(self.hrmo_user)

        # Header has JPEG magic bytes but payload is corrupted garbage
        corrupted_bytes = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00corrupted_payload_data"
        bad_img = SimpleUploadedFile("corrupt.jpg", corrupted_bytes, content_type="image/jpeg")
        resp = self.client.post(
            f"/people/{self.person.pk}/documents/new/",
            {"doc_type": Document.Type.CERTIFICATE, "status": Document.Status.RECEIVED, "file": bad_img},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "corrupted or invalid")

    def test_document_download_safe_headers_and_attachment(self):
        pdf_bytes = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF"
        uploaded_file = SimpleUploadedFile("Appointment_Letter.pdf", pdf_bytes, content_type="application/pdf")
        doc = Document.objects.create(
            person=self.person,
            doc_type=Document.Type.APPOINTMENT,
            status=Document.Status.APPROVED,
            file=uploaded_file,
            display_name="Appointment_Letter.pdf",
            version=1,
        )

        self.client.force_login(self.hrmo_user)
        resp = self.client.get(f"/documents/{doc.pk}/download/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment; filename="Appointment_Letter.pdf"', resp["Content-Disposition"])
        self.assertEqual(resp["X-Content-Type-Options"], "nosniff")

        # Verify audit log recorded VIEW event
        audit = AuditLog.objects.filter(
            action=AuditLog.Action.VIEW, model_name="Document", object_id=str(doc.pk)
        ).first()
        self.assertIsNotNone(audit)

    def test_document_metadata_edit_and_lost_update_protection(self):
        self.client.force_login(self.hrmo_user)

        # Valid edit
        payload = {
            "doc_type": self.sample_doc.doc_type,
            "status": Document.Status.APPROVED,
            "original_on_file": False,
            "physical_location": "Cabinet 9, Drawer 9",
            "notes": "Updated filing location",
            "version": 1,
        }
        resp = self.client.post(f"/documents/{self.sample_doc.pk}/edit/", payload)
        self.assertEqual(resp.status_code, 302)

        self.sample_doc.refresh_from_db()
        self.assertEqual(self.sample_doc.version, 2)
        self.assertEqual(self.sample_doc.physical_location, "Cabinet 9, Drawer 9")
        self.assertFalse(self.sample_doc.original_on_file)

        # Stale version edit rejection
        stale_payload = {
            "doc_type": self.sample_doc.doc_type,
            "status": Document.Status.APPROVED,
            "version": 1,  # Stale version!
        }
        resp_stale = self.client.post(f"/documents/{self.sample_doc.pk}/edit/", stale_payload)
        self.assertEqual(resp_stale.status_code, 200)
        self.assertContains(resp_stale, "modified by another user")

    def test_document_replace_preserves_old_scan_and_creates_new_record(self):
        self.client.force_login(self.hrmo_user)

        # Original scan
        pdf1 = b"%PDF-1.4 original scan content"
        f1 = SimpleUploadedFile("First_Scan.pdf", pdf1, content_type="application/pdf")
        doc1 = Document.objects.create(
            person=self.person,
            doc_type=Document.Type.APPOINTMENT,
            status=Document.Status.APPROVED,
            file=f1,
            display_name="First_Scan.pdf",
            version=1,
        )

        # Replacement scan
        pdf2 = b"%PDF-1.4 new clearer scan content"
        f2 = SimpleUploadedFile("Second_Scan.pdf", pdf2, content_type="application/pdf")
        payload = {
            "file": f2,
            "notes": "Re-scanned with municipal official stamp",
            "original_on_file": True,
            "physical_location": "Cabinet 1, Archive",
            "version": 1,
        }
        resp = self.client.post(f"/documents/{doc1.pk}/replace/", payload)
        self.assertEqual(resp.status_code, 302)

        # Old document is marked superseded and pointed to new scan
        doc1.refresh_from_db()
        self.assertTrue(doc1.is_replaced)
        self.assertIsNotNone(doc1.replaced_by)
        self.assertEqual(doc1.version, 2)

        # New document exists with clean version 1
        new_doc = doc1.replaced_by
        self.assertEqual(new_doc.version, 1)
        self.assertEqual(new_doc.display_name, "Second_Scan.pdf")
        self.assertEqual(new_doc.person, self.person)
        self.assertFalse(new_doc.is_replaced)
        self.assertNotEqual(doc1.file.name, new_doc.file.name)


class Step6DDemoUsersTestCase(TestCase):
    """Test suite for create_demo_users command and demo accounts setup (Step 6D)."""

    def setUp(self):
        call_command("seed_offices")
        call_command("setup_roles")
        self.person = Person.objects.create(
            employee_no="SP-00020",
            last_name="Flores",
            first_name="Maria",
            middle_name="Santos",
            suffix="",
            sex=Person.Sex.FEMALE,
            birthdate=date(1992, 3, 15),
            version=1,
        )

    def test_create_demo_users_aborts_when_debug_is_false(self):
        from django.test import override_settings

        with override_settings(DEBUG=False):
            with self.assertRaises(CommandError) as ctx:
                call_command("create_demo_users")
            self.assertIn("DJANGO_DEBUG=1", str(ctx.exception))

    def test_create_demo_users_creates_all_roles_with_must_change_password(self):
        from io import StringIO
        from django.test import override_settings

        out = StringIO()
        with override_settings(DEBUG=True):
            call_command("create_demo_users", stdout=out)

        output_str = out.getvalue()
        self.assertIn("demo_hrmo", output_str)
        self.assertIn("demo_head", output_str)
        self.assertIn("demo_exec", output_str)
        self.assertIn("demo_auditor", output_str)
        self.assertIn("demo_admin", output_str)
        self.assertIn("demo_employee", output_str)

        User = get_user_model()

        # 1. demo_hrmo
        u_hrmo = User.objects.get(username="demo_hrmo")
        self.assertTrue(u_hrmo.groups.filter(name=ROLE_HRMO).exists())
        self.assertTrue(u_hrmo.profile.must_change_password)
        self.assertIsNotNone(u_hrmo.profile.office)
        self.assertEqual(u_hrmo.profile.office.acronym, "HRMO")

        # 2. demo_head
        u_head = User.objects.get(username="demo_head")
        self.assertTrue(u_head.groups.filter(name=ROLE_OFFICE_HEAD).exists())
        self.assertTrue(u_head.profile.must_change_password)
        self.assertIsNotNone(u_head.profile.office)
        self.assertEqual(u_head.profile.office.acronym, "MTO")

        # 3. demo_exec
        u_exec = User.objects.get(username="demo_exec")
        self.assertTrue(u_exec.groups.filter(name=ROLE_EXECUTIVE).exists())
        self.assertTrue(u_exec.profile.must_change_password)
        self.assertIsNotNone(u_head.profile.office)
        self.assertEqual(u_exec.profile.office.acronym, "MO")

        # 4. demo_auditor
        u_auditor = User.objects.get(username="demo_auditor")
        self.assertTrue(u_auditor.groups.filter(name=ROLE_AUDITOR).exists())
        self.assertTrue(u_auditor.profile.must_change_password)

        # 5. demo_admin
        u_admin = User.objects.get(username="demo_admin")
        self.assertTrue(u_admin.groups.filter(name=ROLE_SYSADMIN).exists())
        self.assertTrue(u_admin.is_staff)
        self.assertTrue(u_admin.profile.must_change_password)

        # 6. demo_employee
        u_emp = User.objects.get(username="demo_employee")
        self.assertEqual(u_emp.groups.count(), 0)
        self.assertTrue(u_emp.profile.must_change_password)
        self.assertIsNotNone(u_emp.profile.person)
        self.assertEqual(u_emp.profile.person, self.person)


class AuditHardeningTestCase(TestCase):
    """Test suite for audit hash chain, verification command, request tracing, and last_login suppression."""

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        call_command("seed_offices")
        call_command("setup_roles")
        User = get_user_model()
        self.auditor_user = User.objects.create_user(username="auditor_test", password="Password123!")
        self.auditor_user.groups.add(Group.objects.get(name=ROLE_AUDITOR))

    def test_user_update_ignored_when_only_last_login_changes(self):
        from django.utils import timezone

        User = get_user_model()
        user = User.objects.create_user(username="login_user", password="Password123!")

        # Update ONLY last_login
        user.last_login = timezone.now()
        user.save()

        # No UPDATE audit event should be logged for User
        after_login_count = AuditLog.objects.filter(
            model_name="User", object_id=str(user.pk), action=AuditLog.Action.UPDATE
        ).count()
        self.assertEqual(after_login_count, 0)

        # Updating another field (e.g., first_name) must log an UPDATE event
        user.first_name = "Maria"
        user.save()

        update_log = AuditLog.objects.filter(
            model_name="User", object_id=str(user.pk), action=AuditLog.Action.UPDATE
        ).first()
        self.assertIsNotNone(update_log)
        self.assertIn("first_name", update_log.changes)

    def test_audit_hash_chain_creation_and_verify_chain(self):
        entry1 = log_audit_event(
            action=AuditLog.Action.CREATE,
            model_name="TestModel",
            object_id="1",
            object_repr="Test Obj 1",
        )
        self.assertIsNotNone(entry1.row_hash)
        self.assertEqual(len(entry1.row_hash), 64)
        self.assertEqual(entry1.row_hash, entry1.calculate_row_hash(prev_hash=entry1.prev_hash))

        entry2 = log_audit_event(
            action=AuditLog.Action.UPDATE,
            model_name="TestModel",
            object_id="1",
            object_repr="Test Obj 1 Updated",
        )
        self.assertEqual(entry2.prev_hash, entry1.row_hash)
        self.assertEqual(entry2.row_hash, entry2.calculate_row_hash(prev_hash=entry1.row_hash))

        # verify_chain() must report valid
        is_valid, count, error_msg, broken_id = verify_chain()
        self.assertTrue(is_valid)
        self.assertGreaterEqual(count, 2)
        self.assertEqual(error_msg, "")
        self.assertIsNone(broken_id)

    def test_verify_audit_chain_management_command(self):
        from io import StringIO

        out = StringIO()
        call_command("verify_audit_chain", stdout=out)
        self.assertIn("Audit hash chain verified", out.getvalue())
        self.assertIn("record(s) intact", out.getvalue().lower())

    def test_broken_prev_hash_detection(self):
        from django.utils import timezone

        bad_entry = AuditLog(
            timestamp=timezone.now(),
            action=AuditLog.Action.CREATE,
            model_name="TamperedModel",
            object_id="999",
            object_repr="Tampered Obj",
            prev_hash="corrupted_prev_hash_12345",
            row_hash="corrupted_row_hash_12345",
        )
        bad_entry.save()

        is_valid, count, error_msg, broken_id = verify_chain()
        self.assertFalse(is_valid)
        self.assertIn("prev_hash mismatch", error_msg)
        self.assertEqual(broken_id, bad_entry.pk)

        # verify_audit_chain command must fail with CommandError
        with self.assertRaises(CommandError) as ctx:
            call_command("verify_audit_chain")
        self.assertIn(f"Broken audit chain detected at record #{bad_entry.pk}", str(ctx.exception))

    def test_tampered_row_hash_detection(self):
        from django.utils import timezone

        last_entry = AuditLog.objects.order_by("-id").first()
        prev_hash = last_entry.row_hash if last_entry else ("0" * 64)
        tampered_entry = AuditLog(
            timestamp=timezone.now(),
            action=AuditLog.Action.CREATE,
            model_name="TamperedModel",
            object_id="1000",
            object_repr="Tampered Row Hash",
            prev_hash=prev_hash,
            row_hash="tampered_hash_that_does_not_match_canonical_hash",
        )
        tampered_entry.save()

        is_valid, count, error_msg, broken_id = verify_chain()
        self.assertFalse(is_valid)
        self.assertIn("row_hash mismatch", error_msg)
        self.assertEqual(broken_id, tampered_entry.pk)

    def test_request_id_and_session_key_hash_captured(self):
        from django.contrib.sessions.backends.db import SessionStore

        rf = RequestFactory()
        req = rf.get("/people/", HTTP_X_REQUEST_ID="trace-req-456")
        req.user = self.auditor_user
        req.request_id = "trace-req-456"
        req.session = SessionStore()
        req.session.create()
        raw_key = req.session.session_key
        expected_hash = hashlib.sha256(raw_key.encode("utf-8")).hexdigest()

        entry = log_audit_event(
            action=AuditLog.Action.VIEW,
            model_name="TestTracing",
            object_id="123",
            request=req,
        )
        self.assertEqual(entry.request_id, "trace-req-456")
        self.assertEqual(entry.session_key_hash, expected_hash)
        self.assertNotEqual(entry.session_key_hash, raw_key)

    def test_audit_list_chain_consistent_badge_without_anchor(self):
        self.client.force_login(self.auditor_user)
        self.client.post("/audit/reverify/")
        resp = self.client.get("/audit/")
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.context["verification"].is_valid)
        self.assertEqual(resp.context["verification"].anchor_status, "none")
        # Badge must display "Chain consistent" and NOT "tamper-evident"
        self.assertContains(resp, "Chain consistent")
        self.assertNotContains(resp, "tamper-evident")
        # Must include the "Re-verify all" form
        self.assertContains(resp, "Re-verify all")
        self.assertContains(resp, 'action="/audit/reverify/"')

    def test_audit_list_tamper_evident_badge_with_anchor(self):
        latest = AuditLog.objects.order_by("-id").first()
        valid_anchor = f"{latest.id}:{latest.row_hash}"

        self.client.force_login(self.auditor_user)
        self.client.post("/audit/reverify/", {"anchor": valid_anchor})
        resp = self.client.get("/audit/")
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.context["verification"].is_valid)
        self.assertEqual(resp.context["verification"].anchor_status, "verified")
        # With anchor verified, it displays "tamper-evident"
        self.assertContains(resp, "Anchor verified")
        self.assertContains(resp, "tamper-evident")

    def test_audit_list_unconfirmed_when_chain_broken(self):
        from django.utils import timezone

        bad_entry = AuditLog(
            timestamp=timezone.now(),
            action=AuditLog.Action.CREATE,
            model_name="BrokenModel",
            object_id="777",
            prev_hash="invalid_prev",
            row_hash="invalid_row",
        )
        bad_entry.save()

        self.client.force_login(self.auditor_user)
        self.client.post("/audit/reverify/")
        resp = self.client.get("/audit/")
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(resp.context["verification"].is_valid)
        self.assertContains(resp, "Chain verification unconfirmed")
        self.assertNotContains(resp, "tamper-evident")

    def test_audit_row_lock_select_for_update_supported(self):
        # Verify select_for_update can acquire row locks on AuditLog without permission issues
        entry = log_audit_event(
            action=AuditLog.Action.CREATE,
            model_name="LockTest",
            object_id="1",
        )
        with transaction.atomic():
            locked = AuditLog.objects.select_for_update().order_by("-id").first()
            self.assertEqual(locked.pk, entry.pk)

    def test_x_request_id_trusted_proxy_and_format_validation(self):
        # 1. Untrusted remote IP with valid format header: rejected, generates random UUID
        resp_untrusted = self.client.get(
            "/people/",
            REMOTE_ADDR="198.51.100.42",
            HTTP_X_REQUEST_ID="req-custom-trace-12345",
        )
        self.assertNotEqual(resp_untrusted["X-Request-ID"], "req-custom-trace-12345")
        self.assertEqual(len(resp_untrusted["X-Request-ID"]), 32)  # UUID4 hex

        # 2. Trusted proxy with invalid format header (<8 chars, special chars, >64 chars): rejected
        resp_invalid_chars = self.client.get(
            "/people/",
            REMOTE_ADDR="127.0.0.1",
            HTTP_X_REQUEST_ID="invalid;sql--injection",
        )
        self.assertNotEqual(resp_invalid_chars["X-Request-ID"], "invalid;sql--injection")

        resp_too_short = self.client.get(
            "/people/",
            REMOTE_ADDR="127.0.0.1",
            HTTP_X_REQUEST_ID="short",
        )
        self.assertNotEqual(resp_too_short["X-Request-ID"], "short")

        # 3. Trusted proxy with valid format header: accepted!
        resp_valid = self.client.get(
            "/people/",
            REMOTE_ADDR="127.0.0.1",
            HTTP_X_REQUEST_ID="gateway-trace-valid-0099",
        )
        self.assertEqual(resp_valid["X-Request-ID"], "gateway-trace-valid-0099")

    def test_audit_anchor_management_command(self):
        from io import StringIO

        out_normal = StringIO()
        call_command("audit_anchor", stdout=out_normal)
        content = out_normal.getvalue()
        self.assertIn("Latest Audit Log Anchor:", content)
        self.assertIn("Audit ID:", content)
        self.assertIn("Row Hash:", content)
        self.assertIn("Timestamp:", content)
        self.assertIn("Anchor:", content)

        out_raw = StringIO()
        call_command("audit_anchor", "--raw", stdout=out_raw)
        raw_val = out_raw.getvalue().strip()
        self.assertIn(":", raw_val)
        latest = AuditLog.objects.order_by("-id").first()
        self.assertEqual(raw_val, f"{latest.id}:{latest.row_hash}")

    def test_verify_audit_chain_with_anchor_options(self):
        from io import StringIO

        latest = AuditLog.objects.order_by("-id").first()
        valid_anchor = f"{latest.id}:{latest.row_hash}"

        # 1. Valid anchor succeeds and reports tamper-evident
        out = StringIO()
        call_command("verify_audit_chain", f"--anchor={valid_anchor}", stdout=out)
        self.assertIn("Audit hash chain verified", out.getvalue())
        self.assertIn("tamper-evident", out.getvalue())

        # 2. Missing row anchor fails
        with self.assertRaises(CommandError) as ctx_missing:
            call_command("verify_audit_chain", "--anchor=999999:abc1234567890def")
        self.assertIn("not found in audit log", str(ctx_missing.exception))

        # 3. Mismatched hash anchor fails
        wrong_anchor = f"{latest.id}:0000000000000000000000000000000000000000000000000000000000000000"
        with self.assertRaises(CommandError) as ctx_mismatch:
            call_command("verify_audit_chain", f"--anchor={wrong_anchor}")
        self.assertIn("Anchor hash mismatch", str(ctx_mismatch.exception))


class MustChangePasswordTestCase(TestCase):
    """B4: UserProfile.must_change_password is enforced by middleware (DECISIONS A1)."""

    NEW_PASSWORD = "Fresh-Pass-2026-xyz"

    def setUp(self):
        call_command("seed_offices")
        call_command("setup_roles")
        User = get_user_model()
        self.user = User.objects.create_user(username="otp_user", password="Password123!")
        self.user.groups.add(Group.objects.get(name=ROLE_HRMO))
        self.profile = UserProfile.objects.create(user=self.user, must_change_password=True)

    def _assert_redirects_to_password_change(self, resp):
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp["Location"], "/accounts/password/")

    def test_flagged_user_is_redirected_from_every_page(self):
        self.client.force_login(self.user)
        for url in ["/people/", "/offices/", "/audit/", "/admin/", "/does-not-exist/"]:
            with self.subTest(url=url):
                self._assert_redirects_to_password_change(self.client.get(url))
        # POSTs are blocked too, not just page views
        self._assert_redirects_to_password_change(self.client.post("/people/new/", {}))

    def test_flagged_user_can_reach_password_change_logout_and_static(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get("/accounts/password/").status_code, 200)
        static_resp = self.client.get("/static/does-not-matter.css")
        self.assertNotEqual(static_resp.get("Location"), "/accounts/password/")
        logout_resp = self.client.post("/accounts/logout/")
        self.assertEqual(logout_resp.status_code, 302)
        self.assertEqual(logout_resp["Location"], "/accounts/login/")

    def test_successful_change_clears_flag_and_unblocks(self):
        self.client.force_login(self.user)
        resp = self.client.post(
            "/accounts/password/",
            {
                "old_password": "Password123!",
                "new_password1": self.NEW_PASSWORD,
                "new_password2": self.NEW_PASSWORD,
            },
        )
        self.assertEqual(resp.status_code, 302)
        self.profile.refresh_from_db()
        self.assertFalse(self.profile.must_change_password)
        self.assertEqual(self.client.get("/people/").status_code, 200)

    def test_failed_change_keeps_flag(self):
        self.client.force_login(self.user)
        resp = self.client.post(
            "/accounts/password/",
            {"old_password": "wrong-password", "new_password1": self.NEW_PASSWORD, "new_password2": self.NEW_PASSWORD},
        )
        self.assertEqual(resp.status_code, 200)
        self.profile.refresh_from_db()
        self.assertTrue(self.profile.must_change_password)
        self._assert_redirects_to_password_change(self.client.get("/people/"))

    def test_unflagged_users_and_users_without_profile_are_not_redirected(self):
        self.profile.must_change_password = False
        self.profile.save()
        self.client.force_login(self.user)
        self.assertEqual(self.client.get("/people/").status_code, 200)

        User = get_user_model()
        no_profile = User.objects.create_user(username="no_profile_user", password="Password123!")
        no_profile.groups.add(Group.objects.get(name=ROLE_HRMO))
        self.client.force_login(no_profile)
        self.assertEqual(self.client.get("/people/").status_code, 200)

    def test_anonymous_users_go_to_login_not_password_change(self):
        resp = self.client.get("/people/")
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(resp["Location"].startswith("/accounts/login/"))


class FakeIdGuardTestCase(TestCase):
    """B3: PROTOTYPE_FAKE_DATA_ONLY=0 only skips the form check; the DB constraint still rejects real IDs."""

    def setUp(self):
        call_command("setup_roles")
        User = get_user_model()
        self.hrmo_user = User.objects.create_user(username="fake_id_hrmo", password="Password123!")
        self.hrmo_user.groups.add(Group.objects.get(name=ROLE_HRMO))

    def test_every_id_field_has_a_fake_only_db_constraint(self):
        from hr.models import FAKE_ID_FIELDS

        names = {c.name for c in Person._meta.constraints}
        for f in FAKE_ID_FIELDS:
            self.assertIn(f"person_{f}_fake_only", names)

    def test_flag_off_still_rejects_real_ids_in_form_and_database(self):
        from django.test import override_settings

        data = {"last_name": "Real", "first_name": "Id", "birthdate": "1990-01-01", "tin": "123-456-789-000"}
        with override_settings(PROTOTYPE_FAKE_DATA_ONLY=False):
            form = PersonForm(data=data, user=self.hrmo_user)
            self.assertFalse(form.is_valid())
            self.assertIn("must start with 00", str(form.errors))

        with self.assertRaises(IntegrityError), transaction.atomic():
            Person.objects.create(
                employee_no="SP-99999", last_name="Real", first_name="Id",
                birthdate=date(1990, 1, 1), tin="123-456-789-000",
            )


class AuditVerificationStatusTestCase(TestCase):
    """B1 + B7: /audit/ shows the stored result of the last full scan and never verifies on page load."""

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        call_command("seed_offices")
        call_command("setup_roles")
        User = get_user_model()
        self.auditor = User.objects.create_user(username="verify_auditor", password="Password123!")
        self.auditor.groups.add(Group.objects.get(name=ROLE_AUDITOR))
        self.hrmo = User.objects.create_user(username="verify_hrmo", password="Password123!")
        self.hrmo.groups.add(Group.objects.get(name=ROLE_HRMO))
        log_audit_event(action=AuditLog.Action.CREATE, model_name="VerifyTest", object_id="1")

    def _status(self):
        from hr.models import AuditChainStatus

        return AuditChainStatus.objects.filter(pk=1).first()

    def _verify_log_count(self):
        return AuditLog.objects.filter(action="verify_chain").count()

    def _age_status(self, seconds=61):
        from datetime import timedelta
        from hr.models import AuditChainStatus

        AuditChainStatus.objects.filter(pk=1).update(
            verified_at=self._status().verified_at - timedelta(seconds=seconds)
        )

    def test_page_load_never_runs_the_hash_scan(self):
        from unittest import mock

        self.client.force_login(self.auditor)
        real = AuditLog.calculate_row_hash
        with mock.patch.object(AuditLog, "calculate_row_hash", autospec=True, side_effect=real) as spy:
            self.assertEqual(self.client.get("/audit/").status_code, 200)
            self.assertEqual(self.client.get("/audit/?reverify=1").status_code, 200)
        self.assertEqual(spy.call_count, 0)

    def test_stale_cached_checkpoint_does_not_crash(self):
        """Old code raised UnboundLocalError when the cached checkpoint row no longer matched."""
        from django.core.cache import cache

        cache.set(
            "audit_chain_verification_state",
            {"is_valid": True, "last_verified_id": 999999, "last_verified_hash": "f" * 64, "count": 3},
        )
        is_valid, count, error_msg, broken_id = verify_chain()
        self.assertTrue(is_valid, error_msg)
        self.assertEqual(count, AuditLog.objects.count())
        self.client.force_login(self.auditor)
        self.assertEqual(self.client.get("/audit/").status_code, 200)

    def test_missing_result_shows_not_verified_yet(self):
        self.client.force_login(self.auditor)
        resp = self.client.get("/audit/")
        self.assertEqual(resp.status_code, 200)
        self.assertIsNone(resp.context["verification"])
        self.assertContains(resp, "Not verified yet")
        self.assertNotContains(resp, "tamper-evident")

    def test_stale_result_shows_timestamp_and_unverified_newer_entries(self):
        from datetime import timedelta
        from django.utils import timezone
        from hr.models import AuditChainStatus

        AuditChainStatus.objects.create(
            pk=1,
            verified_at=timezone.now() - timedelta(days=3),
            is_valid=True,
            record_count=1,
            last_record_id=1,
            source="command",
        )
        self.client.force_login(self.auditor)
        resp = self.client.get("/audit/")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Last verified")
        self.assertGreater(resp.context["entries_since"], 0)

        # A result pointing past the end of the table (e.g. after a restore) must not crash either.
        AuditChainStatus.objects.filter(pk=1).update(last_record_id=10**9)
        resp = self.client.get("/audit/")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["entries_since"], 0)

    def test_reverify_runs_full_scan_stores_result_and_is_logged(self):
        self.client.force_login(self.auditor)
        rows_before = AuditLog.objects.count()
        resp = self.client.post("/audit/reverify/")
        self.assertRedirects(resp, "/audit/", fetch_redirect_response=False)

        status = self._status()
        self.assertTrue(status.is_valid)
        self.assertEqual(status.record_count, rows_before)
        self.assertEqual(status.source, "web")
        self.assertEqual(status.verified_by, "verify_auditor")
        self.assertEqual(status.anchor_status, "none")
        log = AuditLog.objects.get(action="verify_chain")
        self.assertEqual(log.username, "verify_auditor")
        self.assertEqual(log.changes["is_valid"], True)

        page = self.client.get("/audit/")
        self.assertContains(page, "Chain consistent")
        self.assertNotContains(page, "tamper-evident")

    def test_reverify_limited_to_once_per_minute(self):
        self.client.force_login(self.auditor)
        self.client.post("/audit/reverify/")
        resp = self.client.post("/audit/reverify/", follow=True)
        self.assertContains(resp, "once per minute")
        self.assertEqual(self._verify_log_count(), 1)

        self._age_status()
        self.client.post("/audit/reverify/")
        self.assertEqual(self._verify_log_count(), 2)

    def test_reverify_with_valid_and_wrong_anchor(self):
        latest = AuditLog.objects.order_by("-id").first()
        self.client.force_login(self.auditor)
        self.client.post("/audit/reverify/", {"anchor": f"{latest.id}:{latest.row_hash}"})
        self.assertEqual(self._status().anchor_status, "verified")
        self.assertContains(self.client.get("/audit/"), "tamper-evident")

        self._age_status()
        self.client.post("/audit/reverify/", {"anchor": f"{latest.id}:" + "0" * 64})
        status = self._status()
        self.assertFalse(status.is_valid)
        self.assertEqual(status.anchor_status, "not_verified")
        page = self.client.get("/audit/")
        self.assertContains(page, "Chain verification unconfirmed")
        self.assertNotContains(page, "tamper-evident")

    def test_broken_chain_is_stored_as_broken(self):
        from django.utils import timezone

        bad = AuditLog(
            timestamp=timezone.now(),
            action=AuditLog.Action.CREATE,
            model_name="Broken",
            prev_hash="bad_prev",
            row_hash="bad_row",
        )
        bad.save()
        self.client.force_login(self.auditor)
        self.client.post("/audit/reverify/")
        status = self._status()
        self.assertFalse(status.is_valid)
        self.assertEqual(status.broken_id, bad.pk)
        self.assertContains(self.client.get("/audit/"), "Chain verification unconfirmed")

    def test_reverify_requires_auditor_and_post(self):
        self.client.force_login(self.hrmo)
        self.assertEqual(self.client.post("/audit/reverify/").status_code, 403)
        self.assertIsNone(self._status())
        self.client.force_login(self.auditor)
        self.assertEqual(self.client.get("/audit/reverify/").status_code, 405)
        self.assertIsNone(self._status())

    def test_management_command_stores_result(self):
        from io import StringIO
        from django.utils import timezone

        call_command("verify_audit_chain", stdout=StringIO())
        status = self._status()
        self.assertTrue(status.is_valid)
        self.assertEqual(status.source, "command")
        self.assertEqual(status.record_count, AuditLog.objects.count())

        AuditLog(timestamp=timezone.now(), action=AuditLog.Action.CREATE, prev_hash="x", row_hash="y").save()
        with self.assertRaises(CommandError):
            call_command("verify_audit_chain", stdout=StringIO(), stderr=StringIO())
        self.assertFalse(self._status().is_valid)


class DocumentMissingVersionTestCase(TestCase):
    """B6: a missing version on document edit/replace is a conflict, like Person and Engagement."""

    def setUp(self):
        call_command("setup_roles")
        User = get_user_model()
        self.hrmo_user = User.objects.create_user(username="doc_version_hrmo", password="Password123!")
        self.hrmo_user.groups.add(Group.objects.get(name=ROLE_HRMO))
        self.person = Person.objects.create(
            employee_no="SP-00077", last_name="Version", first_name="Check", birthdate=date(1985, 3, 3)
        )
        self.doc = Document.objects.create(
            person=self.person,
            doc_type=Document.Type.APPOINTMENT,
            physical_location="Cab 1",
            version=3,
        )
        self.client.force_login(self.hrmo_user)

    def test_edit_without_version_is_rejected(self):
        resp = self.client.post(
            f"/documents/{self.doc.pk}/edit/",
            {"doc_type": Document.Type.CLEARANCE, "status": Document.Status.FILED, "physical_location": "Cab 9"},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "modified by another user")
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.version, 3)
        self.assertEqual(self.doc.doc_type, Document.Type.APPOINTMENT)
        self.assertEqual(self.doc.physical_location, "Cab 1")

    def test_replace_without_version_is_rejected(self):
        new_file = SimpleUploadedFile("New.pdf", b"%PDF-1.4 replacement", content_type="application/pdf")
        resp = self.client.post(f"/documents/{self.doc.pk}/replace/", {"file": new_file})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "modified by another user")
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.version, 3)
        self.assertFalse(self.doc.is_replaced)
        self.assertIsNone(self.doc.replaced_by)
        self.assertEqual(Document.objects.count(), 1)

    def test_edit_with_current_version_still_saves(self):
        resp = self.client.post(
            f"/documents/{self.doc.pk}/edit/",
            {"doc_type": Document.Type.CLEARANCE, "status": Document.Status.FILED, "physical_location": "Cab 9", "version": 3},
        )
        self.assertEqual(resp.status_code, 302)
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.version, 4)


class LeaveEligibleDefaultTestCase(TestCase):
    """B2: leave_eligible is editable by HRMO for every type; the type default is only a form initial value."""

    def setUp(self):
        call_command("seed_offices")
        call_command("setup_roles")
        User = get_user_model()
        self.hrmo_user = User.objects.create_user(username="leave_hrmo", password="Password123!")
        self.hrmo_user.groups.add(Group.objects.get(name=ROLE_HRMO))
        self.office = Office.objects.get(acronym="MTO")
        self.person = Person.objects.create(
            employee_no="SP-00088", last_name="Leave", first_name="Flag", birthdate=date(1980, 2, 2)
        )

    def _payload(self, eng_type, **extra):
        payload = {
            "office": self.office.pk,
            "type": eng_type,
            "position_title": "Clerk",
            "start_date": "2026-01-01",
            "end_date": "",
            "status": Engagement.Status.ACTIVE,
            # leave_eligible omitted = checkbox left unticked
        }
        if eng_type == Engagement.Type.PLANTILLA:
            payload.update({"salary_grade": 8, "step": 1, "item_number": "SP-MTO-008"})
        payload.update(extra)
        return payload

    def test_model_save_keeps_explicit_false_for_plantilla_and_casual(self):
        plantilla = Engagement.objects.create(
            person=self.person, office=self.office, type=Engagement.Type.PLANTILLA, position_title="Clerk",
            salary_grade=8, step=1, item_number="SP-MTO-001", start_date=date(2026, 1, 1), leave_eligible=False,
        )
        casual = Engagement.objects.create(
            person=self.person, office=self.office, type=Engagement.Type.CASUAL, position_title="Aide",
            start_date=date(2026, 1, 1), leave_eligible=False,
        )
        plantilla.refresh_from_db()
        casual.refresh_from_db()
        self.assertFalse(plantilla.leave_eligible)
        self.assertFalse(casual.leave_eligible)

    def test_hrmo_can_create_plantilla_and_casual_without_leave_via_form(self):
        self.client.force_login(self.hrmo_user)
        for eng_type in (Engagement.Type.PLANTILLA, Engagement.Type.CASUAL):
            with self.subTest(type=eng_type):
                resp = self.client.post(f"/people/{self.person.pk}/engagements/new/", self._payload(eng_type))
                self.assertEqual(resp.status_code, 302)
                eng = Engagement.objects.filter(person=self.person, type=eng_type).latest("pk")
                self.assertFalse(eng.leave_eligible)

    def test_hrmo_can_untick_leave_on_existing_plantilla(self):
        eng = Engagement.objects.create(
            person=self.person, office=self.office, type=Engagement.Type.PLANTILLA, position_title="Clerk",
            salary_grade=8, step=1, item_number="SP-MTO-008", start_date=date(2026, 1, 1), leave_eligible=True,
        )
        self.client.force_login(self.hrmo_user)
        resp = self.client.post(
            f"/engagements/{eng.pk}/edit/", self._payload(Engagement.Type.PLANTILLA, version=eng.version)
        )
        self.assertEqual(resp.status_code, 302)
        eng.refresh_from_db()
        self.assertFalse(eng.leave_eligible)

    def test_type_default_is_applied_as_form_initial_value(self):
        for eng_type, expected in [
            (Engagement.Type.PLANTILLA, True),
            (Engagement.Type.CASUAL, True),
            (Engagement.Type.JOB_ORDER, False),
            (Engagement.Type.CONTRACT_OF_SERVICE, False),
        ]:
            with self.subTest(type=eng_type):
                self.assertEqual(Engagement.default_leave_eligible(eng_type), expected)
                form = EngagementForm(initial={"type": eng_type}, user=self.hrmo_user, person=self.person)
                self.assertEqual(bool(form["leave_eligible"].initial), expected)


class SetupRolesAuditTestCase(TestCase):
    """B5: setup_roles runs entirely inside disable_audit(), including group creation."""

    def test_setup_roles_writes_no_audit_rows(self):
        self.assertFalse(Group.objects.exists())
        before = AuditLog.objects.count()
        call_command("setup_roles")
        self.assertTrue(Group.objects.filter(name=ROLE_HRMO).exists())
        self.assertEqual(AuditLog.objects.count(), before)

        # Re-running (idempotent) also writes nothing.
        call_command("setup_roles")
        self.assertEqual(AuditLog.objects.count(), before)


class AuditChainConcurrencyTestCase(TransactionTestCase):
    """Concurrent audit writes must append to one linear chain, never fork it."""

    def _fixture_teardown(self):
        # TransactionTestCase flushes with TRUNCATE, which the append-only trigger blocks.
        # The test database runs as the owner role, so lift that trigger for the flush only.
        with connection.cursor() as cursor:
            cursor.execute("ALTER TABLE hr_auditlog DISABLE TRIGGER trg_hr_auditlog_no_truncate")
        try:
            super()._fixture_teardown()
        finally:
            with connection.cursor() as cursor:
                cursor.execute("ALTER TABLE hr_auditlog ENABLE TRIGGER trg_hr_auditlog_no_truncate")

    def _writer(self, label, errors, logged=None, release=None):
        def run():
            try:
                with transaction.atomic():
                    log_audit_event(action=AuditLog.Action.VIEW, model_name="Race", object_repr=label)
                    if logged:
                        logged.set()
                    if release:
                        release.wait(5)
            except Exception as exc:  # surfaced by the assertion in _race()
                errors.append(exc)
            finally:
                connection.close()

        return threading.Thread(target=run)

    def _race(self):
        """Writer A appends and holds its transaction open while writer B tries to append."""
        errors = []
        a_logged, release_a = threading.Event(), threading.Event()
        writer_a = self._writer("A", errors, logged=a_logged, release=release_a)
        writer_b = self._writer("B", errors)
        writer_a.start()
        self.assertTrue(a_logged.wait(5))
        writer_b.start()
        time.sleep(0.5)  # B is now waiting on A (or, without serialization, already done)
        release_a.set()
        writer_a.join(10)
        writer_b.join(10)
        self.assertEqual(errors, [])

    def test_concurrent_first_writes_share_one_chain(self):
        self._race()
        self.assertEqual(AuditLog.objects.count(), 2)
        result = verify_chain()
        self.assertTrue(result.is_valid, result.error_msg)

    def test_concurrent_writes_after_existing_rows_share_one_chain(self):
        log_audit_event(action=AuditLog.Action.VIEW, model_name="Seed", object_repr="seed")
        self._race()
        self.assertEqual(AuditLog.objects.count(), 3)
        result = verify_chain()
        self.assertTrue(result.is_valid, result.error_msg)



class DataMinimizationTestCase(TestCase):
    """R4/R5/R6: each role sees only the person fields its row in DECISIONS allows."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    BIRTHDATE = "June 19, 1980"

    def setUp(self):
        self.setUp_roles()
        self.person_a.mobile = "0917-555-0101"
        self.person_a.address = "12 Mabini St, Poblacion"
        self.person_a.save()
        self.eng_a.notes = "Confidential HR note"
        self.eng_a.save()
        self.doc_a.display_name = "Rizal_Medical_Certificate.pdf"
        self.doc_a.notes = "Signed copy pending"
        self.doc_a.save()
        # Person B used to work in Office A (ended), now in Office B.
        Engagement.objects.create(
            person=self.person_b, office=self.office_a, type=Engagement.Type.JOB_ORDER,
            position_title="Former Clerk", start_date=date(2023, 1, 1), end_date=date(2024, 12, 31),
            status=Engagement.Status.ENDED,
        )

    def _person_page(self, user):
        self.client.force_login(user)
        resp = self.client.get(f"/people/{self.person_a.pk}/")
        self.assertEqual(resp.status_code, 200)
        return resp

    def test_office_staff_sees_contacts_but_no_personal_details_documents_or_pay(self):
        resp = self._person_page(self.staff_a)
        self.assertContains(resp, "0917-555-0101")
        for hidden in (self.BIRTHDATE, "12 Mabini St", "SG 22", "Cabinet A, Drawer 1",
                       "Rizal_Medical_Certificate.pdf", "Signed copy pending", "Confidential HR note"):
            self.assertNotContains(resp, hidden)
        self.assertContains(self.client.get("/people/"), "0917-555-0101")

    def test_office_head_sees_document_status_but_not_file_names_or_notes(self):
        resp = self._person_page(self.head_a)
        self.assertContains(resp, "0917-555-0101")
        self.assertContains(resp, "Cabinet A, Drawer 1")
        for hidden in (self.BIRTHDATE, "12 Mabini St", "SG 22", "Rizal_Medical_Certificate.pdf",
                       "Signed copy pending", "Confidential HR note"):
            self.assertNotContains(resp, hidden)

    def test_executive_sees_name_position_office_type_only(self):
        resp = self._person_page(self.mayor)
        self.assertContains(resp, "Executive Assistant IV")
        for hidden in (self.BIRTHDATE, "12 Mabini St", "0917-555-0101", "SG 22", "Cabinet A, Drawer 1",
                       "Rizal_Medical_Certificate.pdf", "Confidential HR note"):
            self.assertNotContains(resp, hidden)
        self.assertNotContains(self.client.get("/people/"), "0917-555-0101")

    def test_hrmo_and_auditor_see_everything(self):
        for user in (self.hrmo_user, self.auditor_user):
            resp = self._person_page(user)
            for shown in (self.BIRTHDATE, "12 Mabini St", "0917-555-0101", "SG 22",
                          "Rizal_Medical_Certificate.pdf", "Signed copy pending", "Confidential HR note"):
                self.assertContains(resp, shown)

    def test_employee_sees_own_details_and_documents_but_not_salary_grade(self):
        resp = self._person_page(self.emp_a_user)
        for shown in (self.BIRTHDATE, "12 Mabini St", "0917-555-0101", "Rizal_Medical_Certificate.pdf"):
            self.assertContains(resp, shown)
        self.assertNotContains(resp, "SG 22")

    def test_employee_can_open_documents_attached_to_own_engagement(self):
        from hr.permissions import can_view_document_content

        eng_doc = Document.objects.create(engagement=self.eng_a, doc_type=Document.Type.APPOINTMENT)
        self.assertTrue(can_view_document_content(self.emp_a_user, eng_doc))
        self.assertFalse(can_view_document_content(self.head_a, eng_doc))

    def test_office_page_hides_salary_grade_and_past_staff_from_office_roles(self):
        for user in (self.head_a, self.staff_a):
            self.client.force_login(user)
            resp = self.client.get(f"/offices/{self.office_a.pk}/")
            self.assertEqual(resp.status_code, 200)
            self.assertNotContains(resp, "SG 22")
            self.assertNotContains(resp, "Bonifacio")  # transferred out (R5)
        self.client.force_login(self.hrmo_user)
        resp = self.client.get(f"/offices/{self.office_a.pk}/")
        self.assertContains(resp, "SG 22")
        self.assertContains(resp, "Bonifacio")


class SysadminEscalationTestCase(TestCase):
    """R3/A1: System Admin cannot reach office roles, and admin-set passwords are one-time."""

    setUp = Step4CSecurityAndRolesTestCase.setUp

    def test_office_roles_are_protected(self):
        self.assertIn(ROLE_OFFICE_HEAD, PROTECTED_ROLES)
        self.assertIn(ROLE_OFFICE_STAFF, PROTECTED_ROLES)
        self.client.force_login(self.sysadmin)
        resp = self.client.get(f"/admin/auth/user/{self.head_a.pk}/change/")
        self.assertNotEqual(resp.status_code, 200)
        resp = self.client.get(f"/admin/auth/user/{self.head_a.pk}/password/")
        self.assertNotEqual(resp.status_code, 200)

    def test_admin_password_reset_forces_change_at_next_login(self):
        # An account with no HR access (here: an office link but no role or person).
        clerk = get_user_model().objects.create_user("it_clerk", password="Password123!")
        UserProfile.objects.create(user=clerk, office=self.office_a)
        self.client.force_login(self.sysadmin)
        new_pw = "Rq7!fresh-one-time-pass"
        resp = self.client.post(
            f"/admin/auth/user/{clerk.pk}/password/",
            {"password1": new_pw, "password2": new_pw, "usable_password": "true"},
        )
        self.assertEqual(resp.status_code, 302)
        clerk.profile.refresh_from_db()
        self.assertTrue(clerk.profile.must_change_password)

    def _assert_locked_for_sysadmin(self, target):
        """System Admin can neither open, list, reset nor delete the account (R3)."""
        self.client.force_login(self.sysadmin)
        old_hash = get_user_model().objects.get(pk=target.pk).password
        for url in (f"/admin/auth/user/{target.pk}/change/", f"/admin/auth/user/{target.pk}/password/",
                    f"/admin/auth/user/{target.pk}/delete/"):
            self.assertNotEqual(self.client.get(url).status_code, 200, url)
        new_pw = "Zt8!taken-over-account"
        self.client.post(f"/admin/auth/user/{target.pk}/password/",
                         {"password1": new_pw, "password2": new_pw, "usable_password": "true"})
        self.assertEqual(get_user_model().objects.get(pk=target.pk).password, old_hash)
        listing = self.client.get("/admin/auth/user/")
        self.assertEqual(listing.status_code, 200)
        self.assertNotContains(listing, f"/admin/auth/user/{target.pk}/change/")

    def test_sysadmin_cannot_take_over_an_employee_account(self):
        # No role, but linked to a person: a reset would open that person's PDS, IDs and leave.
        self._assert_locked_for_sysadmin(self.emp_a_user)

    def test_sysadmin_cannot_take_over_a_delegate_account(self):
        today = date.today()
        Delegate.objects.create(principal=self.mayor, delegate=self.delegate_user,
                                valid_from=today, valid_to=today + timedelta(days=30))
        self._assert_locked_for_sysadmin(self.delegate_user)

    def test_future_delegate_is_locked_but_an_expired_one_is_not(self):
        today = date.today()
        delegation = Delegate.objects.create(principal=self.mayor, delegate=self.delegate_user,
                                             valid_from=today + timedelta(days=5), valid_to=today + timedelta(days=9))
        self._assert_locked_for_sysadmin(self.delegate_user)
        delegation.valid_from, delegation.valid_to = today - timedelta(days=9), today - timedelta(days=5)
        delegation.save()
        self.client.force_login(self.sysadmin)
        self.assertEqual(self.client.get(f"/admin/auth/user/{self.delegate_user.pk}/change/").status_code, 200)

    def test_sysadmin_cannot_take_over_an_account_with_direct_hr_permissions(self):
        from django.contrib.auth.models import Permission

        reader = get_user_model().objects.create_user("pay_reader", password="Password123!")
        reader.user_permissions.add(Permission.objects.get(content_type__app_label="hr", codename="view_pay"))
        self._assert_locked_for_sysadmin(reader)

    def test_sysadmin_still_manages_accounts_without_hr_access_and_superuser_resets_employees(self):
        plain = get_user_model().objects.create_user("plain_it", password="Password123!")
        self.client.force_login(self.sysadmin)
        self.assertEqual(self.client.get(f"/admin/auth/user/{plain.pk}/change/").status_code, 200)
        self.assertContains(self.client.get("/admin/auth/user/"), f"/admin/auth/user/{plain.pk}/change/")
        self.client.force_login(self.superuser)
        new_pw = "Wm3!superuser-reset-pass"
        resp = self.client.post(f"/admin/auth/user/{self.emp_a_user.pk}/password/",
                                {"password1": new_pw, "password2": new_pw, "usable_password": "true"})
        self.assertEqual(resp.status_code, 302)

    def test_admin_created_account_must_change_password(self):
        self.client.force_login(self.sysadmin)
        pw = "Kp4!brand-new-account"
        resp = self.client.post(
            "/admin/auth/user/add/",
            {"username": "new_clerk", "usable_password": "true", "password1": pw, "password2": pw, "_save": "Save"},
        )
        self.assertEqual(resp.status_code, 302)
        new_user = get_user_model().objects.get(username="new_clerk")
        self.assertTrue(UserProfile.objects.get(user=new_user).must_change_password)


class LoginLockoutAuditTestCase(TestCase):
    """U5/S9: lockouts are per real client behind a proxy and never store unknown usernames."""

    def setUp(self):
        self.user = get_user_model().objects.create_user("clerk", password="Correct-Horse-91")

    def _fail(self, username, client_ip, times=5):
        for _ in range(times):
            self.client.post(
                "/accounts/login/", {"username": username, "password": "wrong-password"},
                HTTP_X_FORWARDED_FOR=client_ip,
            )

    def test_lockout_behind_proxy_is_per_client_ip(self):
        self._fail("clerk", "198.51.100.1")
        locked = self.client.post(
            "/accounts/login/", {"username": "clerk", "password": "Correct-Horse-91"},
            HTTP_X_FORWARDED_FOR="198.51.100.1",
        )
        self.assertNotEqual(locked.status_code, 302)
        other_client = self.client.post(
            "/accounts/login/", {"username": "clerk", "password": "Correct-Horse-91"},
            HTTP_X_FORWARDED_FOR="198.51.100.2",
        )
        self.assertEqual(other_client.status_code, 302)

    def test_lockout_does_not_store_unknown_username(self):
        secret = "MyActualPassword!2026"
        self._fail(secret, "198.51.100.9")
        lockouts = AuditLog.objects.filter(action=AuditLog.Action.ACCOUNT_LOCKOUT)
        self.assertTrue(lockouts.exists())
        self.assertEqual(lockouts.first().username, "[unknown]")
        self.assertEqual(lockouts.first().ip_address, "198.51.100.9")
        for row in AuditLog.objects.all():
            self.assertNotIn(secret, f"{row.username} {row.object_repr} {row.changes}")

    def test_lockout_keeps_real_username(self):
        self._fail("clerk", "198.51.100.3")
        self.assertEqual(
            AuditLog.objects.filter(action=AuditLog.Action.ACCOUNT_LOCKOUT).first().username, "clerk"
        )


class AuditExportAndInputHardeningTestCase(TestCase):
    """CSV formula escaping, malformed filters, superseded scans, and app-role privileges."""

    def setUp(self):
        call_command("seed_offices")
        call_command("setup_roles")
        User = get_user_model()
        self.auditor = User.objects.create_user("aud_csv", password="Password123!")
        self.auditor.groups.add(Group.objects.get(name=ROLE_AUDITOR))
        self.hrmo = User.objects.create_user("hrmo_csv", password="Password123!")
        self.hrmo.groups.add(Group.objects.get(name=ROLE_HRMO))

    def test_csv_safe_escapes_formula_prefixes(self):
        from hr.views import csv_safe

        for value in ("=1+1", "+1", "-1", "@SUM(A1)", "\t=x", "\r=x"):
            self.assertTrue(csv_safe(value).startswith("'"), value)
        self.assertEqual(csv_safe("Rizal, Jose"), "Rizal, Jose")
        self.assertEqual(csv_safe(None), "")

    def test_audit_export_neutralizes_formula_cells(self):
        log_audit_event(action=AuditLog.Action.VIEW, model_name="Person",
                        object_repr='=HYPERLINK("http://evil.example","x")')
        self.client.force_login(self.auditor)
        content = self.client.get("/audit/export/csv/").content.decode()
        self.assertIn("'=HYPERLINK", content)
        self.assertNotIn(",=HYPERLINK", content)
        self.assertNotIn('"=HYPERLINK', content)

    def test_malformed_filters_do_not_500(self):
        self.client.force_login(self.hrmo)
        self.assertEqual(self.client.get("/people/?office=abc&type=nope&status=zzz").status_code, 200)
        self.client.force_login(self.auditor)
        self.assertEqual(self.client.get("/audit/?start_date=junk&end_date=2026-02-31").status_code, 200)

    def test_superseded_document_cannot_be_replaced_again(self):
        person = Person.objects.create(employee_no="SP-00555", last_name="Luna", first_name="Juan",
                                       birthdate=date(1985, 1, 1))
        current = Document.objects.create(person=person, doc_type=Document.Type.PDS)
        old = Document.objects.create(person=person, doc_type=Document.Type.PDS,
                                      is_replaced=True, replaced_by=current)
        self.client.force_login(self.hrmo)
        self.assertEqual(self.client.get(f"/documents/{old.pk}/replace/").status_code, 302)
        pdf = SimpleUploadedFile("scan.pdf", b"%PDF-1.4 test", content_type="application/pdf")
        before = Document.objects.count()
        resp = self.client.post(f"/documents/{old.pk}/replace/", {"version": old.version, "file": pdf})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Document.objects.count(), before)
        old.refresh_from_db()
        self.assertEqual(old.replaced_by_id, current.pk)

    def test_privilege_migration_uses_configured_role(self):
        import importlib
        import os
        from unittest import mock

        migration = importlib.import_module("hr.migrations.0011_reapply_app_role_audit_privileges")
        with mock.patch.dict(os.environ, {"APP_DB_USER": "no_such_role_xyz"}):
            with connection.schema_editor() as editor, self.assertRaises(RuntimeError):
                migration.restrict_app_role(None, editor)

        role = os.environ.get("APP_DB_USER") or "hr_app"
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1 FROM pg_roles WHERE rolname = %s", [role])
            if cursor.fetchone() is None:
                self.skipTest(f"role {role} not present in this cluster")
        with connection.schema_editor() as editor:
            migration.restrict_app_role(None, editor)
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT has_table_privilege(%s, 'hr_auditlog', 'UPDATE'),"
                " has_table_privilege(%s, 'hr_auditlog', 'DELETE'),"
                " has_column_privilege(%s, 'hr_auditlog', 'id', 'UPDATE')",
                [role, role, role],
            )
            self.assertEqual(cursor.fetchone(), (False, False, False))


class ProjectManagementTestCase(TestCase):
    """P1-P3: Project Maintainer role, funding reference, and contract-within-project rules."""

    def setUp(self):
        from hr.permissions import ROLE_PROJECT_MAINTAINER

        call_command("seed_offices")
        call_command("setup_roles")
        User = get_user_model()
        self.office = Office.objects.get(name="Office of the Municipal Engineer")

        self.maintainer = User.objects.create_user("planning_officer", password="Password123!")
        self.maintainer.groups.add(Group.objects.get(name=ROLE_PROJECT_MAINTAINER))
        self.hrmo = User.objects.create_user("hrmo_proj", password="Password123!")
        self.hrmo.groups.add(Group.objects.get(name=ROLE_HRMO))
        self.sysadmin = User.objects.create_user("sysadmin_proj", password="Password123!", is_staff=True)
        self.sysadmin.groups.add(Group.objects.get(name=ROLE_SYSADMIN))

        self.project = Project.objects.create(
            name="Coastal Road Concreting", office=self.office, status=Project.Status.ONGOING,
            start_date=date(2026, 1, 1), end_date=date(2026, 12, 31), funding_reference="FAKE-CAF-2026-001",
        )
        self.person = Person.objects.create(
            employee_no="SP-00901", last_name="Dagohoy", first_name="Francisco", birthdate=date(1990, 4, 4)
        )

    def _project_payload(self, **overrides):
        payload = {
            "name": "Evacuation Center Phase 3", "office": self.office.pk, "status": Project.Status.PLANNED,
            "start_date": "2026-03-01", "end_date": "2027-02-28", "funding_reference": "FAKE-CAF-2026-077",
            "notes": "",
        }
        payload.update(overrides)
        return payload

    def _contract_payload(self, **overrides):
        payload = {
            "office": self.office.pk, "type": Engagement.Type.PROJECT_BASED, "position_title": "Site Inspector",
            "start_date": "2026-02-01", "end_date": "2026-11-30", "status": Engagement.Status.ACTIVE,
            "project": self.project.pk,
        }
        payload.update(overrides)
        return payload

    # --- Role and access --------------------------------------------------------

    def test_setup_roles_creates_project_maintainer_group(self):
        from hr.permissions import ROLE_PROJECT_MAINTAINER

        perms = set(Group.objects.get(name=ROLE_PROJECT_MAINTAINER).permissions.values_list("codename", flat=True))
        self.assertIn("manage_projects", perms)
        self.assertFalse({"view_person", "view_hr_records", "manage_hr_records"} & perms)
        hrmo_perms = set(Group.objects.get(name=ROLE_HRMO).permissions.values_list("codename", flat=True))
        self.assertIn("view_project", hrmo_perms)
        self.assertNotIn("manage_projects", hrmo_perms)

    def test_maintainer_creates_project_and_it_is_audited(self):
        self.client.force_login(self.maintainer)
        resp = self.client.post("/projects/new/", self._project_payload())
        self.assertEqual(resp.status_code, 302)
        created = Project.objects.get(name="Evacuation Center Phase 3")
        self.assertEqual(created.funding_reference, "FAKE-CAF-2026-077")
        self.assertTrue(AuditLog.objects.filter(
            action=AuditLog.Action.CREATE, model_name="Project", object_id=str(created.pk)
        ).exists())

    def test_funding_reference_is_required(self):
        self.client.force_login(self.maintainer)
        resp = self.client.post("/projects/new/", self._project_payload(funding_reference="  "))
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(Project.objects.filter(name="Evacuation Center Phase 3").exists())

    def test_maintainer_has_no_personnel_access(self):
        self.client.force_login(self.maintainer)
        self.assertRedirects(self.client.get("/"), "/projects/")
        self.assertEqual(self.client.get("/people/").status_code, 403)
        self.assertEqual(self.client.get(f"/people/{self.person.pk}/").status_code, 404)
        self.assertEqual(self.client.get(f"/people/{self.person.pk}/engagements/new/").status_code, 403)

    def test_maintainer_sees_headcount_but_not_names(self):
        Engagement.objects.create(
            person=self.person, office=self.office, type=Engagement.Type.PROJECT_BASED,
            position_title="Site Inspector", start_date=date(2026, 2, 1), end_date=date(2026, 11, 30),
            project=self.project,
        )
        self.client.force_login(self.maintainer)
        resp = self.client.get(f"/projects/{self.project.pk}/")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["active_count"], 1)
        self.assertNotContains(resp, "Dagohoy")
        self.client.force_login(self.hrmo)
        self.assertContains(self.client.get(f"/projects/{self.project.pk}/"), "Dagohoy")

    def test_hrmo_cannot_create_or_edit_projects_without_the_role(self):
        from hr.permissions import ROLE_PROJECT_MAINTAINER

        self.client.force_login(self.hrmo)
        self.assertEqual(self.client.get("/projects/").status_code, 200)
        self.assertEqual(self.client.post("/projects/new/", self._project_payload()).status_code, 403)
        self.assertEqual(self.client.get(f"/projects/{self.project.pk}/edit/").status_code, 403)

        # Roles stack: an HRMO staffer added to Project Maintainer can do both (LGU decision, P1).
        self.hrmo.groups.add(Group.objects.get(name=ROLE_PROJECT_MAINTAINER))
        self.assertEqual(self.client.post("/projects/new/", self._project_payload()).status_code, 302)

    def test_sysadmin_cannot_assign_project_maintainer(self):
        from hr.permissions import ROLE_PROJECT_MAINTAINER

        self.assertIn(ROLE_PROJECT_MAINTAINER, PROTECTED_ROLES)

    def test_project_edit_rejects_stale_version(self):
        self.client.force_login(self.maintainer)
        payload = self._project_payload(
            name=self.project.name, status=Project.Status.ONGOING, start_date="2026-01-01",
            end_date="2026-12-31", version=self.project.version,
        )
        self.assertEqual(self.client.post(f"/projects/{self.project.pk}/edit/", payload).status_code, 302)
        resp = self.client.post(f"/projects/{self.project.pk}/edit/", payload)  # same, now stale version
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "modified by another user")
        self.project.refresh_from_db()
        self.assertEqual(self.project.version, 2)

    def test_project_end_cannot_cut_active_contracts(self):
        Engagement.objects.create(
            person=self.person, office=self.office, type=Engagement.Type.PROJECT_BASED,
            position_title="Site Inspector", start_date=date(2026, 2, 1), end_date=date(2026, 11, 30),
            project=self.project,
        )
        self.client.force_login(self.maintainer)
        payload = self._project_payload(
            name=self.project.name, status=Project.Status.ONGOING, start_date="2026-01-01",
            end_date="2026-06-30", version=self.project.version,
        )
        resp = self.client.post(f"/projects/{self.project.pk}/edit/", payload)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "run past this date")
        self.project.refresh_from_db()
        self.assertEqual(self.project.end_date, date(2026, 12, 31))

    # --- Hiring against projects --------------------------------------------------

    def test_hrmo_hires_onto_open_funded_project(self):
        self.client.force_login(self.hrmo)
        resp = self.client.post(f"/people/{self.person.pk}/engagements/new/", self._contract_payload())
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(Engagement.objects.filter(person=self.person, project=self.project).exists())

    def test_unfunded_or_closed_projects_are_not_offered(self):
        unfunded = Project.objects.create(
            name="Unfunded Idea", office=self.office, status=Project.Status.PLANNED, start_date=date(2026, 1, 1),
        )
        completed = Project.objects.create(
            name="Finished Bridge", office=self.office, status=Project.Status.COMPLETED,
            start_date=date(2025, 1, 1), end_date=date(2025, 12, 31), funding_reference="FAKE-CAF-2025-050",
        )
        form = EngagementForm(user=self.hrmo, person=self.person)
        offered = set(form.fields["project"].queryset)
        self.assertIn(self.project, offered)
        self.assertNotIn(unfunded, offered)
        self.assertNotIn(completed, offered)

        self.client.force_login(self.hrmo)
        resp = self.client.post(
            f"/people/{self.person.pk}/engagements/new/", self._contract_payload(project=unfunded.pk)
        )
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(Engagement.objects.filter(project=unfunded).exists())

    def test_existing_contract_keeps_its_closed_project_on_edit(self):
        eng = Engagement.objects.create(
            person=self.person, office=self.office, type=Engagement.Type.PROJECT_BASED,
            position_title="Site Inspector", start_date=date(2026, 2, 1), end_date=date(2026, 11, 30),
            project=self.project,
        )
        self.project.status = Project.Status.COMPLETED
        self.project.save()
        form = EngagementForm(instance=eng, user=self.hrmo, person=self.person)
        self.assertIn(self.project, set(form.fields["project"].queryset))

    def test_contract_must_fit_inside_project_dates(self):
        self.client.force_login(self.hrmo)
        url = f"/people/{self.person.pk}/engagements/new/"
        too_late = self.client.post(url, self._contract_payload(end_date="2027-03-31"))
        self.assertContains(too_late, "Contract must end by the project end date")
        open_ended = self.client.post(url, self._contract_payload(end_date=""))
        self.assertContains(open_ended, "Contract must end by the project end date")
        too_early = self.client.post(url, self._contract_payload(start_date="2025-12-01"))
        self.assertContains(too_early, "cannot start before the project")
        self.assertFalse(Engagement.objects.filter(person=self.person).exists())


class ProvisionAccountsTestCase(TestCase):
    """A1/A4/D9: bulk employee accounts with expiring one-time passwords and no roles."""

    def setUp(self):
        call_command("seed_offices")
        call_command("setup_roles")
        self.mto = Office.objects.get(acronym="MTO")
        self.hrmo_office = Office.objects.get(acronym="HRMO")
        self.active_mto = self._person("SP-00011", "Aquino", self.mto)
        self.active_hrmo = self._person("SP-00012", "Bautista", self.hrmo_office)
        self.ended_only = self._person("SP-00013", "Cruz", self.mto, status=Engagement.Status.ENDED)
        self.linked = self._person("SP-00014", "Diaz", self.mto)
        linked_user = get_user_model().objects.create_user("SP-00014", password="Existing-Pass-91")
        UserProfile.objects.create(user=linked_user, person=self.linked)
        self.conflict = self._person("SP-00015", "Estrada", self.mto)
        get_user_model().objects.create_user("SP-00015", password="Someone-Else-91")  # not linked

    def _person(self, emp_no, last_name, office, status=Engagement.Status.ACTIVE):
        person = Person.objects.create(
            employee_no=emp_no, last_name=last_name, first_name="Test", birthdate=date(1990, 1, 1)
        )
        Engagement.objects.create(
            person=person, office=office, type=Engagement.Type.JOB_ORDER, position_title="Clerk",
            start_date=date(2026, 1, 1), end_date=date(2026, 12, 31) if status != Engagement.Status.ACTIVE else None,
            status=status,
        )
        return person

    def _provision(self, *args):
        from io import StringIO

        out = StringIO()
        call_command("provision_accounts", *args, stdout=out)
        text = out.getvalue()
        import re

        passwords = dict(re.findall(r"^(SP-\d{5})\s+.*\s(\S+)$", text, flags=re.M))
        return text, passwords

    def test_dry_run_creates_nothing(self):
        text, _ = self._provision("--dry-run")
        self.assertIn("2 account(s) to create", text)
        self.assertIn("Skipped SP-00015", text)
        self.assertFalse(get_user_model().objects.filter(username__in=["SP-00011", "SP-00012"]).exists())

    def test_creates_accounts_only_for_active_unlinked_people(self):
        _, passwords = self._provision()
        self.assertEqual(set(passwords), {"SP-00011", "SP-00012"})
        User = get_user_model()
        self.assertFalse(User.objects.filter(username="SP-00013").exists())  # no active appointment
        self.assertFalse(UserProfile.objects.filter(person=self.conflict).exists())  # conflict untouched
        self.assertTrue(User.objects.get(username="SP-00014").check_password("Existing-Pass-91"))

    def test_office_filter(self):
        _, passwords = self._provision("--office", "mto")
        self.assertEqual(set(passwords), {"SP-00011"})

    def test_provisioned_account_is_a_plain_employee_account(self):
        from django.utils import timezone
        from hr.permissions import can_view_ids, get_visible_people

        _, passwords = self._provision()
        user = get_user_model().objects.get(username="SP-00011")
        self.assertTrue(user.check_password(passwords["SP-00011"]))
        self.assertFalse(user.is_superuser)
        self.assertFalse(user.is_staff)
        self.assertFalse(user.groups.exists())
        self.assertEqual(user.get_all_permissions(), set())
        self.assertEqual(list(get_visible_people(user)), [self.active_mto])
        self.assertTrue(can_view_ids(user, self.active_mto))
        self.assertFalse(can_view_ids(user, self.active_hrmo))
        profile = user.profile
        self.assertTrue(profile.must_change_password)
        remaining = profile.password_expires_at - timezone.now()
        self.assertGreater(remaining.days, 12)
        self.assertLessEqual(remaining.days, 14)

    def test_first_login_forces_password_change_and_clears_expiry(self):
        _, passwords = self._provision()
        resp = self.client.post("/accounts/login/", {"username": "SP-00011", "password": passwords["SP-00011"]})
        self.assertEqual(resp.status_code, 302)
        self.assertRedirects(self.client.get(f"/people/{self.active_mto.pk}/"), "/accounts/password/")
        new_pw = "My-Own-Strong-Pass-2026"
        resp = self.client.post("/accounts/password/", {
            "old_password": passwords["SP-00011"], "new_password1": new_pw, "new_password2": new_pw,
        })
        self.assertEqual(resp.status_code, 302)
        profile = UserProfile.objects.get(user__username="SP-00011")
        self.assertFalse(profile.must_change_password)
        self.assertIsNone(profile.password_expires_at)
        self.assertEqual(self.client.get(f"/people/{self.active_mto.pk}/").status_code, 200)

    def test_expired_one_time_password_is_refused(self):
        from datetime import timedelta
        from django.utils import timezone

        _, passwords = self._provision()
        UserProfile.objects.filter(user__username="SP-00011").update(
            password_expires_at=timezone.now() - timedelta(minutes=1)
        )
        resp = self.client.post("/accounts/login/", {"username": "SP-00011", "password": passwords["SP-00011"]})
        self.assertEqual(resp.status_code, 200)  # login form again, generic error
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_passwords_never_reach_the_audit_log(self):
        _, passwords = self._provision()
        created = AuditLog.objects.filter(action=AuditLog.Action.CREATE, model_name="User",
                                          object_repr__in=["SP-00011", "SP-00012"])
        self.assertEqual(created.count(), 2)
        for row in AuditLog.objects.all():
            for password in passwords.values():
                self.assertNotIn(password, f"{row.object_repr} {row.changes}")
        self.assertEqual(created.first().changes["password"], "[REDACTED]")
        self.assertTrue(verify_chain().is_valid)

    def test_slips_page_lists_each_account(self):
        import os
        import tempfile

        path = os.path.join(tempfile.mkdtemp(), "account_slips_test.html")
        _, passwords = self._provision("--slips", path)
        page = open(path, encoding="utf-8").read()
        for emp_no, password in passwords.items():
            self.assertIn(emp_no, page)
            self.assertIn(password, page)
        os.remove(path)

    def test_admin_reset_password_also_expires(self):
        from hr.accounts import require_password_change

        user = get_user_model().objects.create_user("reset_me", password="Old-Pass-2026x")
        profile = require_password_change(user)
        self.assertTrue(profile.must_change_password)
        self.assertIsNotNone(profile.password_expires_at)

    def test_one_time_passwords_avoid_look_alike_characters(self):
        from hr.accounts import generate_one_time_password

        for _ in range(200):
            pwd = generate_one_time_password()
            self.assertEqual(len(pwd), 12)
            self.assertFalse(set(pwd) & set("0O1lI"))


class TamperingTestCase(TestCase):
    """Inspect-element and forged-request attacks: the server must ignore or refuse them all.

    Each test sends what a user could produce by editing the page in the browser's developer
    tools (or replaying it with curl/Burp): extra fields, removed fields, other people's IDs,
    pages their role never shows, or a changed redirect.
    """

    setUp = Step4CSecurityAndRolesTestCase.setUp

    def _person_payload(self, **extra):
        payload = {
            "last_name": "Rizal", "first_name": "Jose", "birthdate": "1980-06-19",
            "tin": "000-111-222-000", "version": Person.objects.get(pk=self.person_a.pk).version,
        }
        payload.update(extra)
        return payload

    def _engagement_payload(self, **extra):
        payload = {
            "office": self.office_a.pk, "type": Engagement.Type.PLANTILLA, "position_title": "Executive Assistant IV",
            "rate": "50000.00", "rate_basis": Engagement.RateBasis.MONTHLY, "start_date": "2025-01-01",
            "status": Engagement.Status.ACTIVE, "salary_grade": 22, "step": 1, "item_number": "SP-MO-001",
            "version": Engagement.objects.get(pk=self.eng_a.pk).version,
        }
        payload.update(extra)
        return payload

    def test_extra_field_cannot_change_employee_number(self):
        self.client.force_login(self.hrmo_user)
        resp = self.client.post(f"/people/{self.person_a.pk}/edit/", self._person_payload(employee_no="SP-99999"))
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Person.objects.get(pk=self.person_a.pk).employee_no, "SP-00101")

    def test_extra_field_cannot_move_engagement_to_another_person(self):
        self.client.force_login(self.hrmo_user)
        resp = self.client.post(f"/engagements/{self.eng_a.pk}/edit/", self._engagement_payload(person=self.person_b.pk))
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Engagement.objects.get(pk=self.eng_a.pk).person_id, self.person_a.pk)

    def test_fields_hidden_from_a_role_are_ignored_when_posted_anyway(self):
        # A records clerk with manage_hr_records but no view_pay / view_ids: the rate and ID
        # inputs are not on their page, so they add them back with developer tools.
        from django.contrib.auth.models import Permission

        clerk = get_user_model().objects.create_user("records_clerk", password="Password123!")
        clerk.user_permissions.add(Permission.objects.get(codename="manage_hr_records"))
        self.client.force_login(clerk)

        payload = self._engagement_payload(rate="999999.00")
        payload.pop("rate_basis")
        self.assertEqual(self.client.post(f"/engagements/{self.eng_a.pk}/edit/", payload).status_code, 302)
        self.assertEqual(Engagement.objects.get(pk=self.eng_a.pk).rate, Decimal("50000.00"))

        resp = self.client.post(f"/people/{self.person_a.pk}/edit/", self._person_payload(tin="00-TAMPERED"))
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Person.objects.get(pk=self.person_a.pk).tin, "000-111-222-000")

    def test_guessing_other_ids_is_refused(self):
        doc_b = Document.objects.create(person=self.person_b, doc_type=Document.Type.PDS)
        self.client.force_login(self.emp_a_user)
        self.assertEqual(self.client.get(f"/people/{self.person_b.pk}/").status_code, 404)
        self.assertEqual(self.client.get(f"/documents/{doc_b.pk}/download/").status_code, 403)
        self.client.force_login(self.head_a)
        resp = self.client.get(f"/people/?office={self.office_b.pk}")
        self.assertNotContains(resp, self.person_b.last_name)
        self.assertEqual(self.client.get(f"/offices/{self.office_b.pk}/").status_code, 404)

    def test_pages_a_role_never_sees_are_refused(self):
        self.client.force_login(self.emp_a_user)
        for url in ("/people/new/", f"/people/{self.person_a.pk}/edit/", f"/engagements/{self.eng_a.pk}/end/",
                    "/projects/new/", "/audit/", "/audit/export/csv/"):
            self.assertEqual(self.client.post(url, {}).status_code, 403, url)
        self.client.force_login(self.staff_a)
        resp = self.client.post(f"/engagements/{self.eng_a.pk}/end/", {"end_date": "2026-01-31", "version": 1})
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Engagement.objects.get(pk=self.eng_a.pk).status, Engagement.Status.ACTIVE)

    def test_sysadmin_cannot_tick_is_superuser(self):
        target = get_user_model().objects.create_user("plain_clerk", password="Password123!")
        self.client.force_login(self.sysadmin)
        resp = self.client.post(f"/admin/auth/user/{target.pk}/change/", {
            "username": "plain_clerk", "first_name": "", "last_name": "", "email": "",
            "is_active": "on", "is_superuser": "on",
            "date_joined_0": "2026-10-11", "date_joined_1": "09:00:00",
            "last_login_0": "", "last_login_1": "", "_save": "Save",
        })
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(get_user_model().objects.get(pk=target.pk).is_superuser)

    def test_cross_site_post_without_csrf_token_is_refused(self):
        from django.test import Client

        client = Client(enforce_csrf_checks=True)
        client.force_login(self.hrmo_user)
        resp = client.post("/people/new/", {"last_name": "X", "first_name": "Y", "birthdate": "1990-01-01"})
        self.assertEqual(resp.status_code, 403)

    def test_login_redirect_cannot_point_off_site(self):
        resp = self.client.post("/accounts/login/?next=https://evil.example/", {
            "username": "hrmo", "password": "hrmopass123", "next": "https://evil.example/",
        })
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(resp.url.startswith("https://evil.example"))


class CalendarTestCase(TestCase):
    """DECISIONS C1-C2: one calendar, every source scoped by the existing role rules."""

    def setUp(self):
        User = get_user_model()
        call_command("seed_offices")
        call_command("setup_roles")
        self.office_a = Office.objects.get(name="Office of the Municipal Mayor")
        self.office_b = Office.objects.get(name="Office of the Municipal Treasurer")

        self.hrmo = User.objects.create_user("cal_hrmo", password="pass123")
        self.hrmo.groups.add(Group.objects.get(name=ROLE_HRMO))

        def person(no, last, office, start):
            p = Person.objects.create(employee_no=no, last_name=last, first_name="X", birthdate=date(1990, 1, 1))
            Engagement.objects.create(
                person=p, office=office, type=Engagement.Type.JOB_ORDER, position_title="Aide",
                start_date=start, end_date=date(2026, 12, 31),
            )
            LeaveRequest.objects.create(
                person=p, leave_type=LeaveType.SICK, start_date=date(2026, 10, 20),
                end_date=date(2026, 10, 21), days=Decimal("2"), status=LeaveRequest.Status.APPROVED,
            )
            return p

        self.coworker = person("SP-09001", "Coworker", self.office_a, date(2026, 10, 5))
        self.me = person("SP-09002", "Myself", self.office_a, date(2025, 1, 1))
        self.outsider = person("SP-09003", "Outsider", self.office_b, date(2026, 10, 6))

        # Office Staff in Office A, linked to their own person record.
        self.staff = User.objects.create_user("cal_staff", password="pass123")
        self.staff.groups.add(Group.objects.get(name=ROLE_OFFICE_STAFF))
        UserProfile.objects.create(user=self.staff, office=self.office_a, person=self.me)

        CalendarEvent.objects.create(title="Town fiesta", kind="holiday", start_date=date(2026, 10, 9))
        CalendarEvent.objects.create(title="Treasury interview", kind="interview",
                                     start_date=date(2026, 10, 14), office=self.office_b)

    def get_month(self, user):
        self.client.force_login(user)
        resp = self.client.get("/calendar/?month=2026-10")
        self.assertEqual(resp.status_code, 200)
        return resp.content.decode()

    def test_hrmo_sees_every_source(self):
        html = self.get_month(self.hrmo)
        for text in ["Town fiesta", "Treasury interview", "Starts: Coworker", "Starts: Outsider",
                     "Coworker, X · Sick leave", "Outsider, X · Sick leave"]:
            self.assertIn(text, html)

    def test_office_staff_sees_own_office_and_own_leave_only(self):
        html = self.get_month(self.staff)
        self.assertIn("Town fiesta", html)                 # municipality-wide event
        self.assertIn("Starts: Coworker", html)            # hire in own office
        self.assertIn("Myself, X · Sick leave", html)       # own leave
        self.assertNotIn("Coworker, X · Sick leave", html)  # coworker's leave type is not theirs to see
        self.assertNotIn("Treasury interview", html)       # other office's event
        self.assertNotIn("Outsider", html)                 # other office's people

    def test_only_hrmo_manages_events_and_changes_are_audited(self):
        data = {"title": "CSC report due", "kind": "deadline", "start_date": "2026-10-30"}
        self.client.force_login(self.staff)
        self.assertEqual(self.client.post("/calendar/events/new/", data).status_code, 403)
        self.assertFalse(CalendarEvent.objects.filter(title="CSC report due").exists())

        self.client.force_login(self.hrmo)
        resp = self.client.post("/calendar/events/new/", data)
        self.assertRedirects(resp, "/calendar/?month=2026-10")
        event = CalendarEvent.objects.get(title="CSC report due")
        self.assertEqual(event.created_by, self.hrmo)

        bad = self.client.post(f"/calendar/events/{event.pk}/", {**data, "end_date": "2026-10-01"})
        self.assertContains(bad, "End date cannot be before start date.")

        self.client.post(f"/calendar/events/{event.pk}/", {"delete": "1"})
        self.assertFalse(CalendarEvent.objects.filter(pk=event.pk).exists())
        actions = set(AuditLog.objects.filter(model_name="CalendarEvent").values_list("action", flat=True))
        self.assertEqual(actions, {"create", "delete"})

    def test_bad_month_falls_back_instead_of_500(self):
        self.client.force_login(self.hrmo)
        for month in ["junk", "9999-12", "2026-13"]:
            self.assertEqual(self.client.get(f"/calendar/?month={month}").status_code, 200)


class CalendarVisualsTestCase(TestCase):
    """DECISIONS C3-C4: busy days collapse, day page, filters, text prefixes, agenda, demo data."""

    def setUp(self):
        from hr.permissions import ROLE_PROJECT_MAINTAINER

        User = get_user_model()
        call_command("seed_offices")
        call_command("setup_roles")
        self.office_a = Office.objects.get(name="Office of the Municipal Mayor")
        self.office_b = Office.objects.get(name="Office of the Municipal Treasurer")
        self.year_end = date(2026, 12, 31)

        self.hrmo = User.objects.create_user("vis_hrmo", password="pass123")
        self.hrmo.groups.add(Group.objects.get(name=ROLE_HRMO))
        self.head_a = User.objects.create_user("vis_head", password="pass123")
        self.head_a.groups.add(Group.objects.get(name=ROLE_OFFICE_HEAD))
        UserProfile.objects.create(user=self.head_a, office=self.office_a)
        self.mayor = User.objects.create_user("vis_mayor", password="pass123")
        self.mayor.groups.add(Group.objects.get(name=ROLE_EXECUTIVE))
        self.maintainer = User.objects.create_user("vis_projects", password="pass123")
        self.maintainer.groups.add(Group.objects.get(name=ROLE_PROJECT_MAINTAINER))

        # Six contracts in Office A and one in Office B all end on Dec 31.
        self.people = []
        for i in range(7):
            office = self.office_a if i < 6 else self.office_b
            p = Person.objects.create(employee_no=f"SP-0950{i}", last_name=f"Endsoon{i}", first_name="Ana",
                                      birthdate=date(1990, 1, 1))
            Engagement.objects.create(person=p, office=office, type=Engagement.Type.JOB_ORDER,
                                      position_title="Utility Worker", start_date=date(2026, 1, 1),
                                      end_date=self.year_end)
            self.people.append(p)
        self.employee = User.objects.create_user("vis_employee", password="pass123")
        UserProfile.objects.create(user=self.employee, person=self.people[0])
        CalendarEvent.objects.create(title="Rizal Day", kind="holiday", start_date=date(2026, 12, 30))
        CalendarEvent.objects.create(title="MTO year-end meeting", kind="meeting",
                                     start_date=date(2026, 12, 31), office=self.office_b)
        Project.objects.create(name="Seawall Phase 1", office=self.office_a, status=Project.Status.ONGOING,
                               start_date=date(2026, 1, 1), end_date=self.year_end, funding_reference="FAKE-1")

    def page(self, user, url):
        self.client.force_login(user)
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        return resp.content.decode()

    def test_busy_day_collapses_into_one_summary_entry(self):
        html = self.page(self.hrmo, "/calendar/?month=2026-12")
        self.assertIn("7 contract ends", html)
        self.assertIn("/calendar/day/2026-12-31/#contract_end", html)
        self.assertNotIn("Endsoon3", html)  # individual names live on the day page, not the month

    def test_small_groups_show_short_labels_with_text_prefix(self):
        html = self.page(self.head_a, "/calendar/?month=2026-12")  # 6 ends in Office A: still collapsed
        self.assertIn("6 contract ends", html)
        Engagement.objects.filter(person__in=self.people[1:6]).update(end_date=date(2026, 12, 15))
        html = self.page(self.head_a, "/calendar/?month=2026-12")
        self.assertIn("<strong>End</strong> Endsoon0, A.", html)
        self.assertIn("Contract ends: Endsoon0, Ana, Utility Worker", html)  # full text in the title
        self.assertIn("<strong>Holiday</strong> Rizal Day", html)

    def test_day_page_lists_everything_with_the_same_scoping(self):
        html = self.page(self.hrmo, "/calendar/day/2026-12-31/")
        for p in self.people:
            self.assertIn(p.last_name, html)
        self.assertIn("MTO year-end meeting", html)
        self.assertIn("Project ends: Seawall Phase 1", html)

        html = self.page(self.head_a, "/calendar/day/2026-12-31/")
        self.assertIn("Endsoon0", html)
        self.assertNotIn("Endsoon6", html)              # Office B's person
        self.assertNotIn("MTO year-end meeting", html)  # Office B's event
        self.assertNotIn("Seawall", html)               # Office Heads do not see projects

        self.client.force_login(self.hrmo)
        for bad in ("2026-02-31", "junk", "9999-01-01"):
            self.assertEqual(self.client.get(f"/calendar/day/{bad}/").status_code, 404, bad)

    def test_filters_by_kind_and_office(self):
        html = self.page(self.hrmo, "/calendar/?month=2026-12&filtered=1&show=events")
        self.assertIn("Rizal Day", html)
        self.assertNotIn("contract ends", html)

        html = self.page(self.hrmo, f"/calendar/day/2026-12-31/?filtered=1&show=contract_end&office={self.office_b.pk}")
        self.assertIn("Endsoon6", html)
        self.assertNotIn("Endsoon0", html)

        # Month navigation keeps the filters.
        html = self.page(self.hrmo, "/calendar/?month=2026-12&filtered=1&show=leave")
        self.assertIn("month=2027-01&amp;filtered=1&amp;show=leave", html)

    def test_office_filter_cannot_reach_another_office(self):
        html = self.page(self.head_a, f"/calendar/day/2026-12-31/?office={self.office_b.pk}")
        self.assertNotIn("Endsoon6", html)
        self.assertIn("Endsoon0", html)  # tampered office ignored, own office still shown

    def test_agenda_list_for_small_screens(self):
        html = self.page(self.hrmo, "/calendar/?month=2026-12")
        self.assertIn("cal-agenda", html)
        self.assertIn("Wed, Dec 30", html)
        self.assertIn("<strong>Holiday</strong> Rizal Day", html)

    def test_other_roles(self):
        html = self.page(self.mayor, "/calendar/day/2026-12-31/")
        self.assertIn("Endsoon6", html)          # executives see across offices
        html = self.page(self.employee, "/calendar/day/2026-12-31/")
        self.assertIn("Endsoon0", html)          # own contract end
        self.assertNotIn("Endsoon1", html)
        self.assertNotIn("MTO year-end meeting", html)
        html = self.page(self.maintainer, "/calendar/day/2026-12-31/")
        self.assertIn("Seawall Phase 1", html)   # projects
        self.assertNotIn("Endsoon", html)        # no personnel

    def test_only_hrmo_can_open_or_delete_an_event(self):
        event = CalendarEvent.objects.get(title="Rizal Day")
        self.client.force_login(self.head_a)
        self.assertEqual(self.client.get(f"/calendar/events/{event.pk}/").status_code, 403)
        self.assertEqual(self.client.post(f"/calendar/events/{event.pk}/", {"delete": "1"}).status_code, 403)
        self.assertTrue(CalendarEvent.objects.filter(pk=event.pk).exists())

    def test_cell_label_drops_a_repeated_kind_word(self):
        CalendarEvent.objects.create(title="Interview: Clerk applicants", kind="interview", start_date=date(2026, 12, 10))
        html = self.page(self.hrmo, "/calendar/?month=2026-12")
        self.assertIn("<strong>Interview</strong> Clerk applicants", html)
        self.assertIn('title="Interview: Clerk applicants"', html)  # full title kept

    def test_seed_demo_calendar_fills_the_weeks_around_today(self):
        from io import StringIO
        from django.test import override_settings
        from django.utils import timezone

        with self.assertRaises(CommandError):
            call_command("seed_demo_calendar", stdout=StringIO())

        # Sample leave goes only to leave-eligible people.
        Engagement.objects.filter(person__in=self.people[:4]).update(leave_eligible=True)
        with override_settings(DEBUG=True):
            call_command("seed_demo_calendar", stdout=StringIO())
            first_counts = (CalendarEvent.objects.count(), LeaveRequest.objects.count())
            call_command("seed_demo_calendar", stdout=StringIO())  # replaces, never duplicates
        self.assertEqual((CalendarEvent.objects.count(), LeaveRequest.objects.count()), first_counts)

        today = timezone.localdate()
        demo_events = CalendarEvent.objects.filter(notes__startswith="[demo]")
        self.assertGreaterEqual(demo_events.count(), 9)
        for event in demo_events:
            self.assertLessEqual(abs((event.start_date - today).days), 60)
        self.assertEqual(LeaveRequest.objects.filter(reason__startswith="[demo]").count(), 4)
        self.assertEqual(CalendarEvent.objects.filter(title="Rizal Day").count(), 1)  # real events untouched


class HostedDeploymentTestCase(TestCase):
    """Settings for hosts such as Railway: health probe and client IP from X-Real-IP."""

    @override_settings(SECURE_SSL_REDIRECT=True, ALLOWED_HOSTS=["hr.example.test"])
    def test_health_probe_answers_before_host_check_and_https_redirect(self):
        # Railway probes over plain HTTP with its own Host header.
        resp = self.client.get("/healthz/", HTTP_HOST="healthcheck.railway.app")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.content, b"ok")
        self.assertFalse(AuditLog.objects.exists())
        # Every other page still gets the host check and the redirect.
        self.assertEqual(self.client.get("/", HTTP_HOST="healthcheck.railway.app").status_code, 400)
        self.assertEqual(self.client.get("/", HTTP_HOST="hr.example.test").status_code, 301)

    def test_health_probe_reports_database_failure(self):
        from unittest.mock import patch

        with patch("hr.middleware.connection") as conn:
            conn.cursor.side_effect = DatabaseError("down")
            resp = self.client.get("/healthz")
        self.assertEqual(resp.status_code, 503)

    def test_client_ip_header_from_trusted_proxy_only(self):
        rf = RequestFactory()
        headers = {"HTTP_X_REAL_IP": "203.0.113.7", "HTTP_X_FORWARDED_FOR": "203.0.113.7, 152.233.0.9"}
        with override_settings(TRUSTED_PROXIES=["100.0.0.0/8"], CLIENT_IP_HEADER="X-Real-IP"):
            # From Railway's edge: the header names the client, not the CDN hop at the right of XFF.
            self.assertEqual(get_client_ip(rf.get("/", REMOTE_ADDR="100.64.0.3", **headers)), "203.0.113.7")
            # Sent straight to the app by anyone else: the header is ignored.
            self.assertEqual(get_client_ip(rf.get("/", REMOTE_ADDR="198.51.100.4", **headers)), "198.51.100.4")
            # Missing or malformed header: fall back to the X-Forwarded-For walk.
            bad = rf.get("/", REMOTE_ADDR="100.64.0.3", HTTP_X_REAL_IP="nope", HTTP_X_FORWARDED_FOR="203.0.113.9")
            self.assertEqual(get_client_ip(bad), "203.0.113.9")
        # Not configured: X-Real-IP is never read.
        with override_settings(TRUSTED_PROXIES=["100.0.0.0/8"], CLIENT_IP_HEADER=""):
            self.assertEqual(get_client_ip(rf.get("/", REMOTE_ADDR="100.64.0.3", **headers)), "152.233.0.9")


class OrgChartTestCase(TestCase):
    """O1-O5: office tree, staff charts, reporting lines, My Office and exports."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp
    _seq = 500

    def setUp(self):
        self.setUp_roles()
        self.boss = self._eng("Dagohoy", "Francisco", Engagement.Type.ELECTED, "Municipal Mayor")
        self.eng_a.reports_to = self.boss
        self.eng_a.save()
        self.clerk = self._eng("Silang", "Gabriela", Engagement.Type.CASUAL, "Administrative Aide IV", reports_to=self.eng_a)
        self.jo = self._eng("Luna", "Antonio", Engagement.Type.JOB_ORDER, "Job Order Support Staff", reports_to=self.eng_a)
        self.jo_user = get_user_model().objects.create_user("emp_jo", "jo@test.com", "jopass12345")
        UserProfile.objects.create(user=self.jo_user, person=self.jo.person)

    def _eng(self, last, first, eng_type, title, office=None, reports_to=None):
        OrgChartTestCase._seq += 1
        person = Person.objects.create(
            employee_no=f"SP-0{OrgChartTestCase._seq}", last_name=last, first_name=first, birthdate=date(1985, 1, 1)
        )
        return Engagement.objects.create(
            person=person, office=office or self.office_a, type=eng_type, position_title=title,
            start_date=date(2025, 1, 1), status=Engagement.Status.ACTIVE, reports_to=reports_to,
        )

    def _chart(self, user, office=None):
        self.client.force_login(user)
        return self.client.get(f"/offices/{(office or self.office_a).pk}/chart/")

    # --- What is drawn ----------------------------------------------------------------

    def test_staff_chart_draws_chart_types_and_only_counts_job_orders(self):
        from hr.orgchart import build_staff_chart

        staff = build_staff_chart(self.office_a, show_types=True)
        self.assertEqual([n.title for n in staff.chart.nodes], ["Francisco Dagohoy", "Jose Rizal", "Gabriela Silang"])
        self.assertEqual(staff.head, self.boss)
        self.assertEqual(staff.not_drawn, {"Job Order": 1})
        self.assertEqual(len(staff.chart.edges), 2)

    def test_line_to_an_ended_appointment_starts_a_new_top_box(self):
        from hr.orgchart import build_staff_chart

        self.boss.status = Engagement.Status.ENDED
        self.boss.save()
        staff = build_staff_chart(self.office_a)
        self.assertEqual(staff.roots, [self.eng_a])
        self.assertEqual(staff.broken_lines, [self.eng_a])

    def test_layout_hangs_many_leaves_without_overlap(self):
        from hr.orgchart import ChartNode, layout

        root = ChartNode(title="Head", children=[ChartNode(title=f"Staff {i}", subtitle="Clerk") for i in range(9)])
        chart = layout([root])
        self.assertEqual({e.kind for e in chart.edges}, {"left", "right"})
        boxes = [(n.x, n.y, n.x + n.w, n.y + n.h) for n in chart.nodes]
        for i, a in enumerate(boxes):
            for b in boxes[i + 1:]:
                self.assertFalse(a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3], (a, b))

    def test_long_text_is_wrapped_and_shortened(self):
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from hr.orgchart import BOLD, wrap

        lines = wrap("Municipal Social Welfare and Development Office of the Municipality of San Pascual", BOLD, 11, 180, 2)
        self.assertEqual(len(lines), 2)
        self.assertTrue(lines[-1].endswith("…"))
        self.assertTrue(all(stringWidth(line, BOLD, 11) <= 180 for line in lines))

    def test_office_tree_follows_office_parents(self):
        from hr.orgchart import build_office_tree

        chart = build_office_tree(whole=True)
        self.assertEqual(len(chart.nodes), Office.objects.count() + 1)  # plus the municipality box
        self.assertEqual(len(chart.edges), Office.objects.count())

    # --- Who sees what (O3) -------------------------------------------------------------

    def test_employee_sees_own_office_chart_only_with_names_and_positions(self):
        resp = self._chart(self.emp_a_user)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Francisco Dagohoy")
        self.assertContains(resp, ">YOU<")
        self.assertNotContains(resp, ">Elected<")  # no appointment type for employees
        self.assertNotContains(resp, f"/people/{self.boss.person_id}/")  # no links to others' records
        self.assertEqual(self._chart(self.emp_a_user, self.office_b).status_code, 403)
        self.assertEqual(self.client.get(f"/offices/{self.office_b.pk}/chart/export/png/").status_code, 403)

    def test_new_job_order_hire_sees_supervisor_head_and_coworkers(self):
        self.client.force_login(self.jo_user)
        resp = self.client.get("/my-office/", follow=True)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["supervisor"], self.eng_a)
        self.assertEqual(resp.context["staff"].head, self.boss)
        self.assertContains(resp, "Your coworkers in this office (3)")
        self.assertContains(resp, "Gabriela Silang")

    def test_office_roles_see_types_and_links_hrmo_sees_every_office(self):
        resp = self._chart(self.staff_a)
        self.assertContains(resp, ">Elected<")
        self.assertContains(resp, f"/people/{self.boss.person_id}/")
        self.assertEqual(self._chart(self.staff_a, self.office_b).status_code, 403)
        self.assertEqual(self._chart(self.hrmo_user, self.office_b).status_code, 200)
        self.assertNotContains(self._chart(self.hrmo_user), "SG 22")  # never pay on charts

    def test_sysadmin_and_people_without_an_office_get_no_charts(self):
        self.assertEqual(self._chart(self.sysadmin).status_code, 403)
        self.assertEqual(self.client.get("/offices/chart/").status_code, 403)
        self.client.force_login(self.delegate_user)
        self.assertEqual(self.client.get("/offices/chart/").status_code, 403)
        self.assertRedirects(self.client.get("/my-office/"), "/", fetch_redirect_response=False)

    def test_office_tree_shows_heads_only_for_visible_offices(self):
        self.client.force_login(self.staff_a)
        resp = self.client.get("/offices/chart/")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Francisco Dagohoy")  # own office head
        self.assertContains(resp, "Office of the Municipal Treasurer")
        self.assertNotContains(resp, f"/offices/{self.office_b.pk}/chart/")

    # --- The clicked office as the head (O6) and the Mayor's Office structure (O7) ----------

    def test_top_level_has_the_departments_in_rows_beside_the_mayors_office(self):
        from hr.orgchart import build_office_tree

        chart = build_office_tree()
        head = chart.nodes[0]
        self.assertEqual(head.title, "Municipal Government of San Pascual")
        offices = [e.child for e in chart.edges if e.parent is head]
        self.assertEqual(len(offices), Office.objects.filter(parent=None).count())
        self.assertEqual(offices[0].title, "Office of the Municipal Mayor")  # leads its branch
        self.assertIn("Office of the Municipal Treasurer", [o.title for o in offices])
        self.assertEqual(offices[0].more, "+ 6 divisions, 1 coterminous")
        trunk = head.x + head.w / 2
        last_row_y = max(o.y for o in offices)
        for o in offices:  # the trunk only runs past rows through the gap between boxes
            if o.y < last_row_y:
                self.assertFalse(o.x < trunk < o.x + o.w, o.title)
        self._assert_no_overlap(chart)

    def test_clicking_the_mayors_office_shows_its_own_structure(self):
        self.client.force_login(self.hrmo_user)
        mayor = self.office_a.pk
        top = self.client.get("/offices/chart/")
        self.assertContains(top, f'href="/offices/chart/?office={mayor}"')  # the box opens it
        resp = self.client.get(f"/offices/chart/?office={mayor}")
        self.assertEqual(resp.context["heading"], "Office of the Municipal Mayor")
        for text in ("Municipal Mayor I", "Executive Assistant II", "Local Special Bodies",
                     "Cooperatives Development Division (CDD)", "Office of the Municipal Administrator"):
            self.assertContains(resp, text)
        self.assertNotContains(resp, "Local Government Unit")  # the municipality box is gone
        self.assertNotContains(resp, "Office of the Municipal Treasurer")  # departments stay on top
        self.assertContains(resp, f'href="/offices/{mayor}/chart/"')  # the head opens the staff chart
        self.assertEqual(resp.context["legend"][1][0], "Proposed for creation")
        self.assertIn(f"?office={mayor}", resp.context["export_urls"]["pptx"])
        self.assertEqual([name for name, _url in resp.context["trail"]], ["Municipal Government of San Pascual", "Office of the Municipal Mayor"])
        self.assertContains(self.client.get("/offices/chart/?office=nonsense"), "Local Government Unit")

    def test_mayors_structure_follows_the_chart_and_its_legend(self):
        from hr.orgchart import FILLS, PLACEHOLDER, TEXTS, build_office_tree

        chart = build_office_tree(focus=self.office_a.pk)
        head = chart.nodes[0]
        self.assertEqual((head.title, head.subtitle, head.style), ("Municipal Mayor I", "SG 27", StructurePosition.Status.ELECTED))
        # The Mayor's own staff: a row on one line just below the box, joined to the trunk.
        line_to = {id(e.child): e for e in chart.edges}
        self.assertEqual(
            [(n.title, line_to[id(n)].kind, line_to[id(n)].parent.title, line_to[id(n)].dashed) for n in head.side],
            [
                ("Security Officer II", "side-left", "Senior Administrative Assistant III (Private Secretary II)", False),
                ("Senior Administrative Assistant III (Private Secretary II)", "left", "Municipal Mayor I", False),
                ("Executive Assistant II", "right", "Municipal Mayor I", False),
                ("Local Special Bodies", "side-right", "Executive Assistant II", True),
                ("Office of the Municipal Administrator", "side-right", "Local Special Bodies", True),
            ],
        )
        self.assertTrue(all(n.y >= head.y + head.h for n in head.side))  # lowered, as on HRMO's chart
        self.assertEqual(len({round(n.y + n.h / 2, 1) for n in head.side}), 1)
        divisions = [e.child for e in chart.edges if e.parent is head and e.kind == "below"]
        self.assertEqual([d.title for d in divisions][::5], ["Cooperatives Development Division (CDD)", "Sports Development and Regulation Division (SDRD)"])
        self.assertEqual(len(divisions), 6)
        # The whole chart wraps its divisions into rows to fit a screen; the trunk runs down between
        # the divisions of every row but the last. A slide keeps its divisions in one row, as on the chart.
        rows = sorted({d.y for d in divisions})
        self.assertEqual(len(rows), 2)
        trunk = head.x + head.w / 2
        self.assertFalse(any(d.x < trunk < d.x + d.w for d in divisions if d.y != rows[-1]))
        slide = build_office_tree(focus=self.office_a.pk, part=1)
        self.assertEqual(len({e.child.y for e in slide.edges if e.parent is slide.nodes[0] and e.kind == "below"}), 1)
        self.assertTrue(all(d.y > max(n.y + n.h for n in head.side) for d in divisions))

        # Each division: its chief, the sections under the chief, a column of positions under each.
        gsd = next(d for d in divisions if d.title.startswith("General Services"))
        (chief,) = gsd.children
        self.assertEqual(chief.title, "General Services Officer / Administrative Officer V")
        self.assertEqual([s.title for s in chief.children][:2], ["Records & Archival Management Section", "Procurement and Supplies Section"])
        under = self._under(gsd)
        self.assertAlmostEqual(gsd.w, max(n.x + n.w for n in under) - min(n.x for n in under))  # spans it all
        (records_head,) = chief.children[0].children
        self.assertEqual([c.title for c in records_head.children], ["Administrative Officer I (Records Officer I)"] * 2)
        (bookbinder,) = records_head.joined  # under both Records Officers, joined to each
        self.assertEqual([e.parent for e in chart.edges if e.child is bookbinder], records_head.children)
        self.assertTrue(bookbinder.y > max(c.y + c.h for c in records_head.children))
        self.assertEqual((FILLS[records_head.style], TEXTS[records_head.style]), ("#e00000", "#ffff00"))  # red: vacant, proposed
        self.assertEqual(FILLS[bookbinder.style], "#a9d18e")  # green: existing filled-up
        sdrd = divisions[-1]
        self.assertEqual([s.children[0].style for s in sdrd.children], [PLACEHOLDER, PLACEHOLDER])
        self._assert_no_overlap(chart)

    def test_structure_lines_never_loop(self):
        from hr.models import OrgUnit
        from hr.orgchart import build_office_tree

        guard = StructurePosition.objects.get(office=self.office_a, title="Security Guard II")
        driver = StructurePosition.objects.get(office=self.office_a, title="Driver I")
        driver.reports_to = guard
        driver.save()
        guard.reports_to = driver
        guard.save()
        titles = [n.title for n in build_office_tree(focus=self.office_a.pk).nodes]
        self.assertIn("Security Guard II", titles)
        self.assertIn("Driver I", titles)
        elsewhere = StructurePosition(
            office=self.office_a, unit=OrgUnit.objects.get(acronym="CDD"), title="X",
            status=StructurePosition.Status.PROPOSED, reports_to=guard,
        )
        with self.assertRaises(ValidationError):
            elsewhere.clean()

    def test_wide_structure_exports_as_slides_like_hrmos(self):
        from pptx import Presentation
        from hr.orgchart import build_office_tree

        mayor = self.office_a.pk
        self.assertEqual(build_office_tree(focus=mayor).parts, 2)
        one, two = (build_office_tree(focus=mayor, part=part) for part in (1, 2))
        self.assertEqual((one.part_label, two.part_label), ("Part 1 of 2 · Divisions 1–3", "Part 2 of 2 · Divisions 4–6"))
        self.assertEqual(
            [n.title for n in one.nodes[0].children],
            ["Cooperatives Development Division (CDD)", "Community Affairs and Public Relations Division (CAPRD)",
             "General Services Division (GSD)", "Continued on next slide"],
        )
        back = two.nodes[0].children[0]
        self.assertEqual((back.title, back.subtitle, back.dashed), ("Continued from previous slide", "CDD · CAPRD · GSD", True))
        self.assertEqual(len(two.nodes[0].side), 5)  # every slide keeps the Mayor and the row beside the box
        self.client.force_login(self.hrmo_user)
        prs = Presentation(BytesIO(self.client.get(f"/offices/chart/export/pptx/?office={mayor}").content))
        self.assertEqual(len(prs.slides), 2)
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)
        self.assertTrue(self.client.get(f"/offices/chart/export/png/?office={mayor}").content.startswith(b"\x89PNG"))

    def _under(self, node):
        found = []
        for child in node.children + node.joined:
            found += [child] + self._under(child)
        return found

    def test_structure_seed_keeps_admin_changes(self):
        from hr.models import OrgUnit

        self.assertEqual(OrgUnit.objects.filter(office=self.office_a, parent=None, beside_head=False).count(), 6)
        self.assertEqual(StructurePosition.objects.filter(office=self.office_a).count(), 79)
        StructurePosition.objects.filter(title="Messenger").update(status=StructurePosition.Status.PROPOSED)
        call_command("seed_structure")
        self.assertEqual(StructurePosition.objects.get(title="Messenger").status, StructurePosition.Status.PROPOSED)
        call_command("seed_structure", "--replace")
        self.assertEqual(StructurePosition.objects.get(title="Messenger").status, StructurePosition.Status.FILLED)

    def test_slide_exports(self):
        from pptx import Presentation
        from hr.orgchart import SITE_LEFT, SITE_RIGHT, build_office_tree, render_png, render_pptx

        chart = build_office_tree(focus=self.office_a.pk)
        png = Image.open(BytesIO(render_png(chart, "Test")))
        self.assertAlmostEqual(png.width / png.height, 16 / 9, places=2)
        prs = Presentation(BytesIO(render_pptx(chart, "Test")))
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)  # fits a normal slide
        shapes = list(prs.slides[0].shapes)
        names = {s.shape_id: s.name for s in shapes}
        self.assertIn("Legend: Proposed for creation", names.values())
        dashed = [s for s in shapes if s._element.tag.endswith("}cxnSp") and "prstDash" in s._element.xml]
        self.assertEqual(len(dashed), 2)
        ends = {names[d._element.nvCxnSpPr.cNvCxnSpPr.endCxn.id]: d._element.nvCxnSpPr.cNvCxnSpPr for d in dashed}
        glue = ends["Office of the Municipal Administrator"]
        self.assertEqual((names[glue.stCxn.id], glue.stCxn.idx), ("Local Special Bodies", SITE_RIGHT))
        self.assertEqual(glue.endCxn.idx, SITE_LEFT)
        self.client.force_login(self.hrmo_user)
        resp = self.client.get(f"/offices/chart/export/pptx/?office={self.office_a.pk}")
        self.assertIn("office-tree-mo-", resp["Content-Disposition"])

    def _assert_no_overlap(self, chart):
        boxes = [(n.x, n.y, n.x + n.w, n.y + n.h, n.title) for n in chart.nodes]
        for i, a in enumerate(boxes):
            for b in boxes[i + 1:]:
                self.assertFalse(a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3], (a, b))

    # --- Exports (O4) --------------------------------------------------------------------

    def test_exports_return_files_and_are_audit_logged(self):
        self.client.force_login(self.staff_a)
        magic = {"png": b"\x89PNG", "pdf": b"%PDF", "pptx": b"PK"}
        for fmt, start in magic.items():
            resp = self.client.get(f"/offices/{self.office_a.pk}/chart/export/{fmt}/")
            self.assertEqual(resp.status_code, 200, fmt)
            self.assertTrue(resp.content.startswith(start), fmt)
            self.assertIn(f'.{fmt}"', resp["Content-Disposition"])
        self.assertEqual(AuditLog.objects.filter(action=AuditLog.Action.EXPORT, model_name="Office").count(), 3)
        self.assertEqual(self.client.get(f"/offices/{self.office_a.pk}/chart/export/svg/").status_code, 404)
        self.client.force_login(self.hrmo_user)
        for fmt in magic:
            self.assertEqual(self.client.get(f"/offices/chart/export/{fmt}/").status_code, 200)

    def test_pptx_boxes_are_editable_and_lines_are_glued(self):
        from pptx import Presentation
        from hr.orgchart import build_staff_chart, render_pptx

        prs = Presentation(BytesIO(render_pptx(build_staff_chart(self.office_a).chart, "Test")))
        shapes = list(prs.slides[0].shapes)
        boxes = [s for s in shapes if s.shape_type == 1]  # autoshapes
        self.assertEqual(
            sorted(b.text_frame.paragraphs[0].text for b in boxes),
            ["Francisco Dagohoy", "Gabriela Silang", "Jose Rizal"],
        )
        connectors = [s for s in shapes if s._element.tag.endswith("}cxnSp")]
        self.assertEqual(len(connectors), 2)
        box_ids = {b.shape_id for b in boxes}
        for cxn in connectors:
            glue = cxn._element.nvCxnSpPr.cNvCxnSpPr
            self.assertIn(glue.stCxn.id, box_ids)
            self.assertIn(glue.endCxn.id, box_ids)

    def test_connector_frame_maps_start_and_end(self):
        from hr.orgchart import connector_frame

        for start, end in (((100, 50), (40, 120)), ((100, 50), (180, 120)), ((10, 0), (30, 40))):
            rot, _flip_h, _flip_v, (x, y, w, h) = connector_frame(start, end)
            self.assertIn(rot, (90, 270))  # first leg runs down
            self.assertAlmostEqual(x + w / 2, (start[0] + end[0]) / 2)
            self.assertAlmostEqual(y + h / 2, (start[1] + end[1]) / 2)

    # --- Editing reporting lines (O2) -----------------------------------------------------

    def _post_lines(self, user, data):
        self.client.force_login(user)
        return self.client.post(f"/offices/{self.office_a.pk}/chart/edit/", data)

    def test_hrmo_edits_lines_with_version_check_and_audit(self):
        resp = self._post_lines(self.hrmo_user, {
            f"reports_to_{self.clerk.pk}": str(self.boss.pk), f"version_{self.clerk.pk}": str(self.clerk.version),
        })
        self.assertRedirects(resp, f"/offices/{self.office_a.pk}/chart/", fetch_redirect_response=False)
        self.clerk.refresh_from_db()
        self.assertEqual(self.clerk.reports_to, self.boss)
        self.assertEqual(self.clerk.version, 2)
        self.eng_a.refresh_from_db()
        self.assertEqual(self.eng_a.reports_to, self.boss)  # rows not posted stay as they were
        self.assertTrue(AuditLog.objects.filter(
            action=AuditLog.Action.UPDATE, model_name="Engagement", object_id=str(self.clerk.pk)
        ).exists())

        stale = self._post_lines(self.hrmo_user, {f"reports_to_{self.clerk.pk}": "", f"version_{self.clerk.pk}": "1"})
        self.assertRedirects(stale, f"/offices/{self.office_a.pk}/chart/edit/", fetch_redirect_response=False)
        self.clerk.refresh_from_db()
        self.assertEqual(self.clerk.reports_to, self.boss)

    def test_cycles_other_offices_and_job_order_supervisors_are_refused(self):
        bad = [
            {f"reports_to_{self.boss.pk}": str(self.clerk.pk), f"version_{self.boss.pk}": "1"},  # loop
            {f"reports_to_{self.clerk.pk}": str(self.eng_b.pk), f"version_{self.clerk.pk}": "1"},  # other office
            {f"reports_to_{self.clerk.pk}": str(self.jo.pk), f"version_{self.clerk.pk}": "1"},  # JO as supervisor
            {f"reports_to_{self.clerk.pk}": str(self.clerk.pk), f"version_{self.clerk.pk}": "1"},  # self
        ]
        for data in bad:
            resp = self._post_lines(self.hrmo_user, data)
            self.assertEqual(resp.status_code, 200, data)
            self.assertContains(resp, "Nothing was saved")
        self.boss.refresh_from_db()
        self.clerk.refresh_from_db()
        self.assertIsNone(self.boss.reports_to)
        self.assertEqual(self.clerk.reports_to, self.eng_a)

    def test_only_hrmo_edits_lines(self):
        for user in (self.head_a, self.staff_a, self.emp_a_user, self.auditor_user):
            resp = self._post_lines(user, {f"reports_to_{self.clerk.pk}": "", f"version_{self.clerk.pk}": "1"})
            self.assertEqual(resp.status_code, 403)
        self.clerk.refresh_from_db()
        self.assertEqual(self.clerk.reports_to, self.eng_a)

    def test_engagement_form_checks_reports_to(self):
        form = EngagementForm(instance=self.clerk, user=self.hrmo_user, data={
            "office": self.office_a.pk, "type": Engagement.Type.CASUAL, "position_title": "Administrative Aide IV",
            "start_date": "2025-01-01", "status": Engagement.Status.ACTIVE, "reports_to": self.eng_b.pk,
            "version": self.clerk.version,
        })
        self.assertFalse(form.is_valid())
        self.assertIn("same office", form.errors["reports_to"][0])

    def test_ending_an_appointment_hands_reports_to_the_renewal(self):
        from hr.orgchart import hand_over_direct_reports

        renewal = Engagement.objects.create(
            person=self.person_a, office=self.office_a, type=Engagement.Type.PLANTILLA,
            position_title="Executive Assistant V", salary_grade=23, step=1, item_number="SP-MO-002",
            start_date=date(2026, 1, 1), status=Engagement.Status.ACTIVE, reports_to=self.boss,
        )
        self.eng_a.status = Engagement.Status.ENDED
        self.eng_a.save()
        self.assertEqual(hand_over_direct_reports(self.eng_a), 2)
        self.clerk.refresh_from_db()
        self.jo.refresh_from_db()
        self.assertEqual((self.clerk.reports_to, self.jo.reports_to), (renewal, renewal))

    @override_settings(DEBUG=True)
    def test_seed_reporting_lines_fills_only_empty_lines(self):
        from io import StringIO

        self.boss.position_title = "Municipal Mayor"
        self.clerk.reports_to = None
        self.clerk.save()
        call_command("seed_reporting_lines", stdout=StringIO())
        self.clerk.refresh_from_db()
        self.eng_a.refresh_from_db()
        self.assertIsNotNone(self.clerk.reports_to)
        self.assertEqual(self.eng_a.reports_to, self.boss)  # existing line kept


@override_settings(SALARY_TRACKING_START="2020-01-01", SALARY_LOOKAHEAD_DAYS=90)
class SalaryActionsTestCase(TestCase):
    """SA1-SA6: the daily check, the HRMO workflow and who sees salary actions."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp
    _seq = 700

    def setUp(self):
        self.setUp_roles()

    def _person(self, last="Mabini"):
        SalaryActionsTestCase._seq += 1
        return Person.objects.create(
            employee_no=f"SP-0{SalaryActionsTestCase._seq}", last_name=last, first_name="Apolinario",
            birthdate=date(1985, 1, 1),
        )

    def _eng(self, person, start, eng_type=Engagement.Type.PLANTILLA, item="SP-MTO-010", step=1, end=None,
             status=Engagement.Status.ACTIVE, office=None):
        plantilla = eng_type == Engagement.Type.PLANTILLA
        return Engagement.objects.create(
            person=person, office=office or self.office_b, type=eng_type, position_title="Revenue Officer I",
            salary_grade=11 if plantilla else None, step=step if plantilla else None,
            item_number=item if plantilla else "", start_date=start, end_date=end, status=status,
        )

    def _cases(self, person, kind=None):
        qs = SalaryAction.objects.filter(person=person)
        return qs.filter(kind=kind) if kind else qs

    # --- The check (SA1, SA2) ---------------------------------------------------------

    def test_step_increment_due_after_three_years_and_check_is_idempotent(self):
        from hr.salary_actions import run_check

        person = self._person()
        self._eng(person, date(2023, 3, 1))
        run_check(today=date(2026, 3, 1))
        case = self._cases(person, SalaryAction.Kind.STEP_INCREMENT).get()
        self.assertEqual((case.status, case.due_date, case.from_step, case.to_step),
                         (SalaryAction.Status.DUE, date(2026, 3, 1), 1, 2))
        self.assertEqual(run_check(today=date(2026, 3, 2)), (0, 0))
        self.assertEqual(SalaryCheckRun.objects.count(), 2)

    def test_case_opens_only_within_the_lookahead(self):
        from hr.salary_actions import run_check

        person = self._person()
        self._eng(person, date(2023, 3, 1))
        run_check(today=date(2025, 11, 1))  # 120 days early
        self.assertFalse(self._cases(person).exists())
        run_check(today=date(2025, 12, 15))
        self.assertEqual(self._cases(person).get().status, SalaryAction.Status.UPCOMING)
        run_check(today=date(2026, 3, 1))
        self.assertEqual(self._cases(person).get().status, SalaryAction.Status.DUE)

    @override_settings(SALARY_TRACKING_START="2026-01-01")
    def test_milestones_before_go_live_are_assumed_done_with_a_warning(self):
        from hr.salary_actions import run_check

        person = self._person()
        self._eng(person, date(2017, 2, 1), step=3)
        run_check(today=date(2026, 1, 15))
        case = self._cases(person).get()
        self.assertEqual(case.due_date, date(2026, 2, 1))  # 2020 and 2023 assumed granted on paper
        self.assertIn("assumed granted on paper", case.warnings[0])

    def test_renewals_of_the_same_item_count_together_and_promotion_restarts(self):
        from hr.salary_actions import position_start

        person = self._person()
        first = self._eng(person, date(2020, 1, 1), end=date(2021, 12, 31), status=Engagement.Status.ENDED)
        renewal = self._eng(person, date(2022, 1, 1))
        promoted = self._eng(person, date(2024, 6, 1), item="SP-MTO-020")
        engs = [first, renewal, promoted]
        self.assertEqual(position_start(renewal, engs), date(2020, 1, 1))
        self.assertEqual(position_start(promoted, engs), date(2024, 6, 1))

    def test_leave_without_pay_moves_the_due_date(self):
        from hr.salary_actions import run_check

        person = self._person()
        self._eng(person, date(2023, 3, 1))
        ServiceInterruption.objects.create(person=person, kind=ServiceInterruption.Kind.LWOP,
                                           start_date=date(2024, 1, 10), end_date=date(2024, 1, 19))
        run_check(today=date(2026, 3, 1))
        case = self._cases(person).get()
        self.assertEqual((case.due_date, case.extension_days), (date(2026, 3, 11), 10))
        self.assertEqual(case.status, SalaryAction.Status.UPCOMING)

    def test_step_eight_and_job_orders_get_no_step_increment(self):
        from hr.salary_actions import run_check

        top = self._person("Aguinaldo")
        self._eng(top, date(2023, 3, 1), step=8)
        jo = self._person("Luna")
        self._eng(jo, date(2023, 3, 1), eng_type=Engagement.Type.JOB_ORDER)
        run_check(today=date(2026, 3, 1))
        self.assertFalse(self._cases(top, SalaryAction.Kind.STEP_INCREMENT).exists())
        self.assertFalse(self._cases(jo).exists())

    def test_loyalty_counts_unbroken_appointive_service_but_not_job_orders(self):
        from hr.salary_actions import run_check

        person = self._person()
        self._eng(person, date(2014, 1, 1), eng_type=Engagement.Type.JOB_ORDER, end=date(2015, 12, 31),
                  status=Engagement.Status.ENDED)
        self._eng(person, date(2016, 1, 1), eng_type=Engagement.Type.CASUAL, end=date(2019, 12, 31),
                  status=Engagement.Status.ENDED)
        self._eng(person, date(2020, 1, 1), step=2)
        run_check(today=date(2025, 12, 1))
        case = self._cases(person, SalaryAction.Kind.LOYALTY).get()
        self.assertEqual((case.years, case.service_from, case.due_date), (10, date(2016, 1, 1), date(2026, 1, 1)))
        self.assertTrue(any("not counted" in w for w in case.warnings))

    def test_a_gap_restarts_loyalty_counting(self):
        from hr.salary_actions import run_check

        person = self._person()
        self._eng(person, date(2016, 1, 1), eng_type=Engagement.Type.CASUAL, end=date(2019, 6, 30),
                  status=Engagement.Status.ENDED)
        self._eng(person, date(2020, 1, 1))
        run_check(today=date(2025, 12, 1))
        self.assertFalse(self._cases(person, SalaryAction.Kind.LOYALTY).exists())

    def test_cases_that_no_longer_apply_are_closed_and_deferred_ones_come_back(self):
        from hr.salary_actions import run_check

        person = self._person()
        eng = self._eng(person, date(2023, 3, 1))
        run_check(today=date(2026, 2, 1))
        eng.status = Engagement.Status.RESIGNED
        eng.save()
        run_check(today=date(2026, 2, 2))
        case = self._cases(person).get()
        self.assertEqual(case.status, SalaryAction.Status.NOT_ELIGIBLE)
        self.assertIn("daily check", case.decision_note)

        other = self._person("Jacinto")
        self._eng(other, date(2023, 3, 1))
        run_check(today=date(2026, 3, 1))
        deferred = self._cases(other).get()
        deferred.status, deferred.deferred_until = SalaryAction.Status.DEFERRED, date(2026, 4, 1)
        deferred.save()
        run_check(today=date(2026, 4, 1))
        deferred.refresh_from_db()
        self.assertEqual(deferred.status, SalaryAction.Status.DUE)

    # --- Workflow (SA3) -----------------------------------------------------------------

    def _due_case(self):
        from django.utils import timezone
        from hr.salary_actions import add_years, run_check

        person = self._person()
        eng = self._eng(person, add_years(timezone.localdate(), -3))
        run_check()
        return self._cases(person).get(), eng

    def _act(self, case, action, user=None, **data):
        self.client.force_login(user or self.hrmo_user)
        case.refresh_from_db()
        payload = {"action": action, "version": str(case.version), **data}
        return self.client.post(f"/salary-actions/{case.pk}/", payload, follow=True)

    def test_hrmo_review_signature_and_issue_raises_the_step(self):
        case, eng = self._due_case()
        self._act(case, "start_review")
        resp = self._act(case, "signature", check=["rating"])
        self.assertContains(resp, "Tick every checklist item first")
        case.refresh_from_db()
        self.assertEqual(case.status, SalaryAction.Status.REVIEW)
        self._act(case, "signature", check=["rating", "no_suspension", "lwop", "funds"])
        case.refresh_from_db()
        self.assertEqual(case.status, SalaryAction.Status.SIGNATURE)
        self._act(case, "issue", nosi_number="NOSI-2026-001", signed_on="2026-01-05", effective_date=str(case.due_date))
        case.refresh_from_db()
        eng.refresh_from_db()
        self.assertEqual(case.status, SalaryAction.Status.ISSUED)
        self.assertEqual((eng.step, eng.version), (2, 2))
        self.assertEqual(case.issued_by, self.hrmo_user)
        self.assertTrue(AuditLog.objects.filter(model_name="Engagement", object_id=str(eng.pk),
                                                action=AuditLog.Action.UPDATE).exists())
        self.assertTrue(AuditLog.objects.filter(model_name="SalaryAction", object_id=str(case.pk)).exists())

        from hr.salary_actions import next_milestones
        from hr.salary_actions import add_years
        self.assertEqual(next_milestones(case.person)["step"]["due_date"], add_years(case.effective_date, 3))

    def test_issue_is_refused_when_the_step_changed_meanwhile(self):
        case, eng = self._due_case()
        case.status = SalaryAction.Status.SIGNATURE
        case.save()
        eng.step = 3
        eng.save()
        resp = self._act(case, "issue", nosi_number="N-1", signed_on="2026-01-05", effective_date="2026-01-05")
        self.assertContains(resp, "is now at step 3")
        case.refresh_from_db()
        self.assertEqual(case.status, SalaryAction.Status.SIGNATURE)

    def test_stale_version_defer_and_not_eligible_need_reasons(self):
        case, _eng = self._due_case()
        self.client.force_login(self.hrmo_user)
        resp = self.client.post(f"/salary-actions/{case.pk}/", {"action": "start_review", "version": "99"}, follow=True)
        self.assertContains(resp, "Someone changed this case")
        resp = self._act(case, "defer", deferred_until="2099-01-01")
        self.assertContains(resp, "Give a reason")
        resp = self._act(case, "defer", deferred_until="2099-01-01", note="Pending IPCR")
        case.refresh_from_db()
        self.assertEqual((case.status, case.deferred_until), (SalaryAction.Status.DEFERRED, date(2099, 1, 1)))
        self._act(case, "not_eligible", note="Unsatisfactory rating")
        case.refresh_from_db()
        self.assertEqual(case.status, SalaryAction.Status.NOT_ELIGIBLE)

    def test_recording_leave_without_pay_re_dates_open_cases(self):
        case, _eng = self._due_case()
        self.client.force_login(self.hrmo_user)
        start = case.service_from + __import__("datetime").timedelta(days=30)
        self.client.post(f"/salary-actions/{case.pk}/interruptions/", {
            "kind": "lwop", "start_date": start.isoformat(),
            "end_date": (start + __import__("datetime").timedelta(days=4)).isoformat(), "notes": "",
        })
        case.refresh_from_db()
        self.assertEqual(case.extension_days, 5)
        self.assertEqual(case.status, SalaryAction.Status.UPCOMING)

    # --- Who sees what (SA4) ------------------------------------------------------------

    def test_only_hrmo_works_cases_and_auditor_reads_them(self):
        case, _eng = self._due_case()
        for user in (self.head_a, self.staff_a, self.mayor, self.emp_a_user, self.sysadmin):
            self.client.force_login(user)
            self.assertEqual(self.client.get("/salary-actions/").status_code, 403, user)
            self.assertEqual(self.client.get(f"/salary-actions/{case.pk}/").status_code, 403, user)
        self.client.force_login(self.auditor_user)
        self.assertContains(self.client.get("/salary-actions/"), case.person.full_name)
        self.assertNotContains(self.client.get(f"/salary-actions/{case.pk}/"), "Start review")
        self.assertEqual(self._act(case, "start_review", user=self.auditor_user).status_code, 403)
        self.assertEqual(self.client.post("/salary-actions/check/").status_code, 403)

    def test_tracker_runs_the_check_when_it_has_not_run_today(self):
        from django.utils import timezone
        from hr.salary_actions import add_years

        person = self._person()
        self._eng(person, add_years(timezone.localdate(), -3))
        self.client.force_login(self.hrmo_user)
        resp = self.client.get("/salary-actions/")
        self.assertContains(resp, person.full_name)
        self.assertEqual(SalaryCheckRun.objects.get().source, "page")
        self.client.get("/salary-actions/")
        self.assertEqual(SalaryCheckRun.objects.count(), 1)

    def test_calendar_and_person_page_show_milestones_to_hrmo_only(self):
        case, _eng = self._due_case()
        month = f"?month={case.due_date:%Y-%m}"
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get("/calendar/" + month), "step 1 → 2")
        self.assertContains(self.client.get(f"/people/{case.person_id}/"), "Next salary increase (step)")
        self.client.force_login(self.head_a)
        self.assertNotContains(self.client.get("/calendar/" + month), "step 1 → 2")


@override_settings(RENEWAL_LOOKAHEAD_DAYS=45)
class RemindersTestCase(TestCase):
    """RM1-RM5: contract renewals, retirement and year-end bonus lists."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp
    _seq = 900
    TODAY = date(2026, 10, 17)

    def setUp(self):
        self.setUp_roles()

    def _person(self, last="Del Pilar", birthdate=date(1985, 1, 1)):
        RemindersTestCase._seq += 1
        return Person.objects.create(employee_no=f"SP-0{RemindersTestCase._seq}", last_name=last,
                                     first_name="Gregorio", birthdate=birthdate)

    def _eng(self, person, eng_type, start, end=None, status=Engagement.Status.ACTIVE, item="SP-MTO-050"):
        plantilla = eng_type == Engagement.Type.PLANTILLA
        return Engagement.objects.create(
            person=person, office=self.office_b, type=eng_type, position_title="Clerk",
            salary_grade=8 if plantilla else None, step=1 if plantilla else None,
            item_number=item if plantilla else "", start_date=start, end_date=end, status=status,
        )

    # --- Renewals (RM1) -------------------------------------------------------------------

    def test_renewals_list_lapsed_and_soon_and_drop_renewed(self):
        from hr.reminders import renewal_rows

        lapsed = self._eng(self._person("Lapsed"), Engagement.Type.JOB_ORDER, date(2026, 7, 1), date(2026, 10, 7))
        soon = self._eng(self._person("Soon"), Engagement.Type.CONTRACT_OF_SERVICE, date(2026, 7, 1), date(2026, 11, 20))
        self._eng(self._person("Later"), Engagement.Type.JOB_ORDER, date(2026, 7, 1), date(2026, 12, 31))
        renewed_person = self._person("Renewed")
        renewed = self._eng(renewed_person, Engagement.Type.JOB_ORDER, date(2026, 7, 1), date(2026, 11, 1))
        self._eng(renewed_person, Engagement.Type.JOB_ORDER, date(2026, 11, 2), date(2027, 4, 30))
        rows = {r["eng"].pk: r for r in renewal_rows(self.TODAY)}
        self.assertEqual(set(rows), {lapsed.pk, soon.pk})
        self.assertTrue(rows[lapsed.pk]["lapsed"])
        self.assertEqual(rows[lapsed.pk]["ago"], 10)
        self.assertNotIn(renewed.pk, rows)

    # --- Retirement (RM2) -----------------------------------------------------------------

    def test_retirement_groups(self):
        from hr.reminders import retirement_rows

        past = self._person("Past", date(1961, 9, 1))
        self._eng(past, Engagement.Type.PLANTILLA, date(2010, 1, 1))
        soon = self._person("Soon", date(1962, 3, 1))
        self._eng(soon, Engagement.Type.CASUAL, date(2020, 1, 1))
        optional = self._person("Optional", date(1965, 5, 1))
        self._eng(optional, Engagement.Type.CASUAL, date(2008, 1, 1), date(2011, 12, 31), Engagement.Status.ENDED)
        self._eng(optional, Engagement.Type.PLANTILLA, date(2012, 1, 1), item="SP-MTO-051")
        young_service = self._person("ShortService", date(1964, 5, 1))
        self._eng(young_service, Engagement.Type.PLANTILLA, date(2020, 1, 1), item="SP-MTO-052")
        jo = self._person("OldJO", date(1950, 1, 1))
        self._eng(jo, Engagement.Type.JOB_ORDER, date(2026, 1, 1), date(2026, 12, 31))
        groups = retirement_rows(self.TODAY)
        names = {k: [r["person"].last_name for r in v] for k, v in groups.items()}
        self.assertEqual(names, {"past": ["Past"], "soon": ["Soon"], "optional": ["Optional"]})
        self.assertEqual(groups["optional"][0]["service_years"], 18)

    # --- Year-end bonus (RM3) -------------------------------------------------------------

    def test_year_end_bonus_eligibility(self):
        from hr.reminders import ELIGIBLE, NOT_IN_SERVICE, PRORATED, yeb_rows

        full = self._person("Full")
        self._eng(full, Engagement.Type.CASUAL, date(2026, 1, 1), date(2026, 3, 31), Engagement.Status.ENDED)
        self._eng(full, Engagement.Type.CONTRACTUAL, date(2026, 9, 1), date(2026, 12, 31))  # 90 + 61 days
        short = self._person("Short")
        self._eng(short, Engagement.Type.PLANTILLA, date(2026, 8, 1))  # 92 days by Oct 31
        left = self._person("Left")
        self._eng(left, Engagement.Type.PLANTILLA, date(2020, 1, 1), date(2026, 6, 30), Engagement.Status.RESIGNED)
        jo = self._person("JobOrder")
        self._eng(jo, Engagement.Type.JOB_ORDER, date(2026, 1, 1), date(2026, 12, 31))
        rows = {r["person"].last_name: r for r in yeb_rows(2026, self.TODAY)}
        self.assertEqual(rows["Full"]["result"], ELIGIBLE)
        self.assertEqual(rows["Full"]["days"], 151)
        self.assertEqual(rows["Short"]["result"], PRORATED)
        self.assertEqual(rows["Left"]["result"], NOT_IN_SERVICE)
        self.assertNotIn("JobOrder", rows)

    # --- Pages, exports, calendar (RM4) ---------------------------------------------------

    def test_only_hrmo_and_auditor_see_the_lists(self):
        pages = ["/tracker/renewals/", "/tracker/retirement/", "/tracker/year-end-bonus/"]
        for user in (self.head_a, self.staff_a, self.mayor, self.emp_a_user, self.sysadmin):
            self.client.force_login(user)
            for page in pages:
                self.assertEqual(self.client.get(page).status_code, 403, (user, page))
        for user in (self.hrmo_user, self.auditor_user):
            self.client.force_login(user)
            for page in pages:
                self.assertEqual(self.client.get(page).status_code, 200, (user, page))

    def test_excel_exports_are_audited_and_formula_safe(self):
        from openpyxl import load_workbook

        person = self._person("=HYPERLINK(1)")
        self._eng(person, Engagement.Type.JOB_ORDER, date(2025, 1, 1), timezone_today_plus(5))
        self.client.force_login(self.hrmo_user)
        for page in ("/tracker/renewals/", "/tracker/retirement/", "/tracker/year-end-bonus/"):
            resp = self.client.get(page + "?export=xlsx")
            self.assertEqual(resp.status_code, 200, page)
            self.assertTrue(resp.content.startswith(b"PK"))
        sheet = load_workbook(BytesIO(self.client.get("/tracker/renewals/?export=xlsx").content)).active
        names = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertIn("'=HYPERLINK(1), Gregorio", names)
        self.assertEqual(AuditLog.objects.filter(action=AuditLog.Action.EXPORT, model_name="Engagement").count(), 4)

    def test_calendar_shows_65th_birthdays_to_hrmo_only(self):
        from hr.salary_actions import add_years

        turning = timezone_today_plus(10)
        person = self._person("Retiree", add_years(turning, -65))
        self._eng(person, Engagement.Type.PLANTILLA, date(2010, 1, 1))
        month = f"?month={turning:%Y-%m}"
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get("/calendar/" + month), "Turns 65 (compulsory retirement)")
        self.client.force_login(self.head_a)
        self.assertNotContains(self.client.get("/calendar/" + month), "Turns 65")

    def test_menu_dot_when_a_contract_lapsed(self):
        self.client.force_login(self.hrmo_user)
        self.assertNotContains(self.client.get("/people/"), "needs attention")
        self._eng(self._person(), Engagement.Type.JOB_ORDER, date(2025, 1, 1), timezone_today_plus(-3))
        self.assertContains(self.client.get("/people/"), "needs attention")


def timezone_today_plus(days):
    from datetime import timedelta
    from django.utils import timezone

    return timezone.localdate() + timedelta(days=days)


class RequiredMeetingsTestCase(TestCase):
    """MT1-MT5: required meeting series, periods, attendance and who sees them."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()

    def _series(self, **kwargs):
        defaults = {"name": "HR committee", "starts_on": date(2026, 1, 15)}
        defaults.update(kwargs)
        members = defaults.pop("members", [])
        series = MeetingSeries.objects.create(**defaults)
        series.members.set(members)
        return series

    def test_periods_are_calendar_quarters_and_creation_is_idempotent(self):
        from hr.meetings import ensure_meetings, period_of

        self.assertEqual(period_of(date(2026, 11, 3), 3), (date(2026, 10, 1), date(2026, 12, 31)))
        self.assertEqual(period_of(date(2026, 8, 3), 6), (date(2026, 7, 1), date(2026, 12, 31)))
        self.assertEqual(period_of(date(2026, 2, 3), 1), (date(2026, 2, 1), date(2026, 2, 28)))
        series = self._series()
        ensure_meetings(date(2026, 10, 17))
        ensure_meetings(date(2026, 10, 18))
        self.assertEqual([m.period_label for m in series.meetings.order_by("period_start")],
                         ["Q1 2026", "Q2 2026", "Q3 2026", "Q4 2026"])

    def test_status_follows_the_dates(self):
        from hr.meetings import ensure_meetings

        series = self._series()
        ensure_meetings(date(2026, 10, 17))
        q3, q4 = series.meetings.get(period_label_start(7)), series.meetings.get(period_label_start(10))
        today = date(2026, 10, 17)
        self.assertEqual(q3.status_on(today), Meeting.Status.MISSED)
        self.assertEqual(q4.status_on(today), Meeting.Status.DUE)
        q4.scheduled_on = date(2026, 11, 20)
        self.assertEqual(q4.status_on(today), Meeting.Status.SCHEDULED)
        q3.not_held_reason = "Typhoon"
        self.assertEqual(q3.status_on(today), Meeting.Status.NOT_HELD)

    def test_expected_people_by_audience(self):
        from hr.meetings import expected_people

        everyone = self._series(name="All staff")
        office = self._series(name="Mayor's office", audience=MeetingSeries.Audience.OFFICE, office=self.office_a)
        chosen = self._series(name="Committee", audience=MeetingSeries.Audience.CHOSEN, members=[self.person_b])
        self.assertEqual(set(expected_people(everyone)), {self.person_a, self.person_b})
        self.assertEqual(list(expected_people(office)), [self.person_a])
        self.assertEqual(list(expected_people(chosen)), [self.person_b])

    def test_hrmo_sets_up_a_series_and_records_attendance(self):
        self.client.force_login(self.hrmo_user)
        resp = self.client.post("/tracker/meetings/new/", {
            "name": "Selection board", "legal_basis": "", "frequency": "quarterly", "audience": "chosen",
            "office": self.office_b.pk, "members": [self.person_a.pk, self.person_b.pk],
            "starts_on": "2026-01-01", "is_active": "on", "notes": "",
        })
        self.assertRedirects(resp, "/tracker/meetings/", fetch_redirect_response=False)
        series = MeetingSeries.objects.get(name="Selection board")
        self.assertIsNone(series.office)  # office only kept for one-office meetings
        from django.utils import timezone

        current = series.meetings.get(period_start__lte=timezone.localdate(), period_end__gte=timezone.localdate())
        self.client.post(f"/tracker/meetings/{current.pk}/", {
            "action": "held", "version": "1", "held_on": timezone.localdate().isoformat(),
            "attended": [str(self.person_a.pk), "99999"], "minutes_location": "Cabinet B",
        })
        current.refresh_from_db()
        self.assertEqual(current.status_on(timezone.localdate()), Meeting.Status.HELD)
        self.assertEqual(list(current.attendees.all()), [self.person_a])  # unknown ids ignored
        self.assertEqual(current.attendance_count, 1)
        self.assertTrue(AuditLog.objects.filter(model_name="Meeting", object_id=str(current.pk)).exists())
        stale = self.client.post(f"/tracker/meetings/{current.pk}/", {"action": "reopen", "version": "1"}, follow=True)
        self.assertContains(stale, "Someone changed this meeting")

    def test_series_form_requires_office_or_members(self):
        self.client.force_login(self.hrmo_user)
        base = {"name": "X", "frequency": "quarterly", "starts_on": "2026-01-01", "is_active": "on"}
        resp = self.client.post("/tracker/meetings/new/", {**base, "audience": "office"})
        self.assertContains(resp, "Choose the office whose staff attend.")
        resp = self.client.post("/tracker/meetings/new/", {**base, "audience": "chosen"})
        self.assertContains(resp, "Choose at least one person.")
        self.assertFalse(MeetingSeries.objects.exists())

    def test_held_needs_a_past_date_and_not_held_needs_a_reason(self):
        from datetime import timedelta
        from django.utils import timezone

        self._series(starts_on=timezone.localdate())
        self.client.force_login(self.hrmo_user)
        self.client.get("/tracker/meetings/")
        meeting = Meeting.objects.get()
        future = (timezone.localdate() + timedelta(days=3)).isoformat()
        resp = self.client.post(f"/tracker/meetings/{meeting.pk}/", {"action": "held", "version": "1", "held_on": future}, follow=True)
        self.assertContains(resp, "today or earlier")
        resp = self.client.post(f"/tracker/meetings/{meeting.pk}/", {"action": "not_held", "version": "1", "reason": ""}, follow=True)
        self.assertContains(resp, "Give the reason")

    def test_only_hrmo_and_auditor_open_the_tracker_and_only_hrmo_edits(self):
        series = self._series()
        self.client.force_login(self.hrmo_user)
        self.client.get("/tracker/meetings/")
        meeting = series.meetings.first()
        for user in (self.head_a, self.staff_a, self.mayor, self.emp_a_user, self.sysadmin):
            self.client.force_login(user)
            self.assertEqual(self.client.get("/tracker/meetings/").status_code, 403, user)
            self.assertEqual(self.client.get(f"/tracker/meetings/{meeting.pk}/").status_code, 403, user)
            self.assertEqual(self.client.get("/tracker/meetings/new/").status_code, 403, user)
        self.client.force_login(self.auditor_user)
        self.assertEqual(self.client.get(f"/tracker/meetings/{meeting.pk}/").status_code, 200)
        resp = self.client.post(f"/tracker/meetings/{meeting.pk}/", {"action": "reopen", "version": "1"})
        self.assertEqual(resp.status_code, 403)

    def test_calendar_shows_meetings_to_their_audience(self):
        from django.utils import timezone

        today = timezone.localdate()
        everyone = self._series(name="All-staff assembly", starts_on=today)
        chosen = self._series(name="Selection board", starts_on=today, audience=MeetingSeries.Audience.CHOSEN,
                              members=[self.person_b])
        from hr.meetings import ensure_meetings
        ensure_meetings(today)
        for series in (everyone, chosen):
            m = series.meetings.get()
            m.scheduled_on = today
            m.save()
        month = f"?month={today:%Y-%m}"
        self.client.force_login(self.emp_a_user)  # person A: not on the board
        resp = self.client.get("/calendar/" + month)
        self.assertContains(resp, "All-staff assembly")
        self.assertNotContains(resp, "Selection board")
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get("/calendar/" + month), "Selection board")

    def test_menu_dot_near_the_end_of_an_unheld_period(self):
        from hr.meetings import needs_attention

        self._series(starts_on=date(2026, 1, 1))
        self.assertFalse(needs_attention(date(2026, 11, 1)))  # Q4 ends in 60 days
        self.assertTrue(needs_attention(date(2026, 12, 15)))
        m = Meeting.objects.get(period_start=date(2026, 10, 1))
        m.held_on = date(2026, 12, 10)
        m.save()
        self.assertFalse(needs_attention(date(2026, 12, 15)))


def period_label_start(month):
    from django.db.models import Q

    return Q(period_start=date(2026, month, 1))


class SalaryCheckSchedulingTestCase(TestCase):
    """SA5: the check runs after each deploy and when the tracker or calendar finds it stale."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()

    def test_calendar_runs_a_stale_check_once_for_hrmo_only(self):
        self.client.force_login(self.head_a)
        self.client.get("/calendar/")
        self.assertFalse(SalaryCheckRun.objects.exists())  # office roles never trigger it
        self.client.force_login(self.hrmo_user)
        self.client.get("/calendar/")
        self.client.get("/calendar/")
        self.assertEqual(list(SalaryCheckRun.objects.values_list("source", flat=True)), ["calendar"])

    def test_deploy_runs_the_check_and_never_fails_because_of_it(self):
        from io import StringIO
        from unittest.mock import patch
        from hr.management.commands.prepare_database import Command

        out, err = StringIO(), StringIO()
        command = Command(stdout=out, stderr=err)
        command._salary_check()
        self.assertEqual(SalaryCheckRun.objects.get().source, "deploy")
        self.assertIn("Salary actions check:", out.getvalue())
        with patch("hr.salary_actions.run_check", side_effect=RuntimeError("boom")):
            command._salary_check()
        self.assertIn("skipped: boom", err.getvalue())


class LeaveWorkflowTestCase(TestCase):
    """L1-L9: filing CS Form No. 6 online, the approval chain, the ledger and who sees what."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()
        self.eng_a.leave_eligible = True
        self.eng_a.save()
        LeaveCredit.objects.create(person=self.person_a, leave_type=LeaveType.VACATION, days=Decimal("10"),
                                   as_of=date(2026, 1, 1), note="Beginning balance")
        LeaveCredit.objects.create(person=self.person_a, leave_type=LeaveType.SICK, days=Decimal("5"),
                                   as_of=date(2026, 1, 1), note="Beginning balance")

    def _file(self, user, **data):
        payload = {"leave_type": "vl", "detail": "within_ph", "detail_text": "", "start_date": "2026-11-02",
                   "end_date": "2026-11-04", "days": "", "reason": "", "leave_type_other": ""}
        payload.update(data)
        self.client.force_login(user)
        return self.client.post("/leave/new/" + ("?paper=1" if "person" in data else ""), payload)

    def _act(self, user, req, action="done", **extra):
        self.client.force_login(user)
        req.refresh_from_db()
        return self.client.post(f"/leave/{req.pk}/", {"action": action, "version": req.version, **extra})

    def test_working_days_skip_weekends_and_holidays(self):
        from hr.leave import working_days

        self.assertEqual(working_days(date(2026, 11, 2), date(2026, 11, 6)), 5)
        self.assertEqual(working_days(date(2026, 10, 31), date(2026, 11, 1)), 0)
        CalendarEvent.objects.create(title="Holiday", kind=CalendarEvent.Kind.HOLIDAY, start_date=date(2026, 11, 4))
        self.assertEqual(working_days(date(2026, 11, 2), date(2026, 11, 6)), 4)
        self.client.force_login(self.emp_a_user)
        data = self.client.get("/leave/days/?start=2026-11-02&end=2026-11-06").json()
        self.assertEqual(data, {"days": 4, "holidays": ["2026-11-04"]})

    def test_employee_files_and_the_chain_runs_to_a_ledger_deduction(self):
        from hr.leave import balance

        resp = self._file(self.emp_a_user)
        req = LeaveRequest.objects.get(person=self.person_a)
        self.assertRedirects(resp, f"/leave/{req.pk}/", fetch_redirect_response=False)
        self.assertEqual((req.days, req.office, req.position_title), (Decimal("3"), self.office_a, "Executive Assistant IV"))
        self.assertEqual([s.kind for s in req.steps.all()], ["recommend", "certify", "approve"])
        self.assertTrue(AuditLog.objects.filter(model_name="LeaveApprovalStep").exists())

        # Nobody signs their own leave, and the Mayor cannot skip ahead of the office head.
        self._act(self.emp_a_user, req)
        self._act(self.mayor, req)
        self.assertEqual(req.current_step.kind, "recommend")

        self.client.force_login(self.head_a)
        self.assertContains(self.client.get("/leave/inbox/"), self.person_a.full_name)
        self._act(self.head_a, req)
        self._act(self.hrmo_user, req)
        req.refresh_from_db()
        self.assertEqual((req.days_with_pay, req.days_without_pay), (Decimal("3"), Decimal("0")))
        self.assertEqual(req.certified_balances["vl"], "10.000")
        self.assertFalse(LeaveCredit.objects.filter(leave_request=req).exists())  # nothing until approval
        self._act(self.mayor, req)
        req.refresh_from_db()
        self.assertEqual(req.status, LeaveRequest.Status.APPROVED)
        self.assertEqual(balance(self.person_a, LeaveType.VACATION), Decimal("7"))

    def test_only_the_office_head_of_the_filing_office_recommends(self):
        from hr.leave import LeaveError, act, can_act, file_request

        head_b = get_user_model().objects.create_user("head_b", "hb@test.com", "x")
        head_b.groups.add(Group.objects.get(name=ROLE_OFFICE_HEAD))
        UserProfile.objects.create(user=head_b, office=self.office_b)
        req = file_request(LeaveRequest(person=self.person_a, leave_type="vl", detail="within_ph",
                                        start_date=date(2026, 11, 2), end_date=date(2026, 11, 2), days=1), None)
        self.assertEqual(can_act(head_b, req.current_step), (False, False))
        with self.assertRaises(LeaveError):
            act(req, head_b, "done")
        self.assertEqual(can_act(self.head_a, req.current_step), (True, False))
        # HRMO may only record the paper form for another signatory (L6).
        self.assertEqual(can_act(self.hrmo_user, req.current_step), (True, True))

    def test_disapproval_needs_a_reason_and_closes_without_charging(self):
        self._file(self.emp_a_user)
        req = LeaveRequest.objects.get(person=self.person_a)
        self._act(self.head_a, req, "disapproved")
        req.refresh_from_db()
        self.assertEqual(req.status, LeaveRequest.Status.PENDING)
        self._act(self.head_a, req, "disapproved", note="Year-end audit week")
        req.refresh_from_db()
        self.assertEqual((req.status, req.decision_note), (LeaveRequest.Status.REJECTED, "Year-end audit week"))
        self.assertFalse(LeaveCredit.objects.filter(leave_request=req).exists())

    def test_sick_leave_beyond_its_balance_is_charged_to_vacation_then_without_pay(self):
        from hr.leave import act, balance, file_request

        req = file_request(LeaveRequest(person=self.person_a, leave_type="sl", detail="out_patient", detail_text="Flu",
                                        start_date=date(2026, 11, 2), end_date=date(2026, 11, 24), days=Decimal("17")),
                           None)
        for user in (self.head_a, self.hrmo_user, self.mayor):
            act(req, user, "done")
            req.save()
        self.assertEqual((req.days_with_pay, req.days_without_pay), (Decimal("15"), Decimal("2")))
        self.assertEqual(balance(self.person_a, LeaveType.SICK), Decimal("0"))
        self.assertEqual(balance(self.person_a, LeaveType.VACATION), Decimal("0"))

    def test_hrmo_cancels_approved_leave_and_the_days_come_back(self):
        from hr.leave import LeaveError, act, balance, cancel, file_request

        req = file_request(LeaveRequest(person=self.person_a, leave_type="vl", detail="within_ph",
                                        start_date=date(2026, 11, 2), end_date=date(2026, 11, 3), days=2), None)
        for user in (self.head_a, self.hrmo_user, self.mayor):
            act(req, user, "done")
            req.save()
        self.assertEqual(balance(self.person_a, LeaveType.VACATION), Decimal("8"))
        with self.assertRaises(LeaveError):
            cancel(req, self.emp_a_user)
        self._act(self.hrmo_user, req, "cancel")  # no reason: refused
        req.refresh_from_db()
        self.assertEqual(req.status, LeaveRequest.Status.APPROVED)
        self._act(self.hrmo_user, req, "cancel", note="Recalled to work")
        req.refresh_from_db()
        self.assertEqual(req.status, LeaveRequest.Status.CANCELLED)
        self.assertEqual(balance(self.person_a, LeaveType.VACATION), Decimal("10"))

    def test_stale_version_saves_nothing(self):
        self._file(self.emp_a_user)
        req = LeaveRequest.objects.get(person=self.person_a)
        self.client.force_login(self.head_a)
        self.client.post(f"/leave/{req.pk}/", {"action": "done", "version": req.version + 5})
        self.assertEqual(req.current_step.kind, "recommend")

    def test_who_can_see_an_application(self):
        self._file(self.emp_a_user)
        req = LeaveRequest.objects.get(person=self.person_a)
        for user, code in ((self.emp_a_user, 200), (self.head_a, 200), (self.mayor, 200), (self.hrmo_user, 200),
                           (self.auditor_user, 200), (self.staff_a, 403), (self.sysadmin, 403)):
            self.client.force_login(user)
            self.assertEqual(self.client.get(f"/leave/{req.pk}/").status_code, code, user.username)
        self.client.force_login(self.staff_a)
        self.assertEqual(self.client.get(f"/leave/{req.pk}/print/").status_code, 403)
        self.client.force_login(self.head_a)
        for path in ("/leave/register/", "/leave/balances/"):
            self.assertEqual(self.client.get(path).status_code, 403, path)

    def test_printed_form_is_filled_in_and_logged(self):
        self._file(self.emp_a_user, commutation_requested="on")
        req = LeaveRequest.objects.get(person=self.person_a)
        self._act(self.head_a, req)
        self._act(self.hrmo_user, req)
        self.client.force_login(self.emp_a_user)
        resp = self.client.get(f"/leave/{req.pk}/print/")
        self.assertContains(resp, "APPLICATION FOR LEAVE")
        self.assertContains(resp, "Executive Assistant IV")
        self.assertContains(resp, "10.000")  # 7.A total earned
        self.assertContains(resp, "7.000")   # 7.A balance after this application
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.Action.EXPORT, model_name="LeaveRequest").exists())

    def test_no_credit_leave_for_appointments_without_leave(self):
        user_b = get_user_model().objects.create_user("emp_b", "b@test.com", "x")
        UserProfile.objects.create(user=user_b, person=self.person_b)
        self.assertContains(self._file(user_b), "does not earn leave")
        self._file(user_b, leave_type="paternity", detail="")
        self.assertEqual(LeaveRequest.objects.get(person=self.person_b).leave_type, "paternity")

    def test_hrmo_enters_a_paper_form_and_records_paper_signatures(self):
        self._file(self.hrmo_user, person=self.person_a.pk)
        req = LeaveRequest.objects.get(person=self.person_a)
        self.assertEqual(req.filed_by, self.hrmo_user)
        self._act(self.hrmo_user, req)
        step = req.steps.get(kind="recommend")
        self.assertEqual((step.status, step.from_paper), ("done", True))
        self.client.force_login(self.emp_a_user)
        self.assertNotContains(self.client.get("/leave/new/?paper=1"), 'name="person"')  # employees file their own

    def test_an_office_heads_own_leave_skips_the_recommendation(self):
        from hr.leave import file_request

        self.emp_a_user.profile.delete()  # the head's own account is the one linked to person A
        self.head_a.profile.person = self.person_a
        self.head_a.profile.save()
        req = file_request(LeaveRequest(person=self.person_a, leave_type="vl", detail="within_ph",
                                        start_date=date(2026, 11, 2), end_date=date(2026, 11, 2), days=1), None)
        self.assertEqual(req.steps.get(kind="recommend").status, "skipped")
        self.assertEqual(req.current_step.kind, "certify")

    def test_monthly_credits_are_posted_once(self):
        from hr.leave import post_monthly_credits

        self.assertEqual(len(post_monthly_credits(2026, 9)), 1)
        self.assertEqual(post_monthly_credits(2026, 9), [])
        self.assertEqual(LeaveCredit.objects.filter(note="Earned September 2026").count(), 2)
        self.client.force_login(self.staff_a)
        self.assertEqual(self.client.post("/leave/credits/monthly/", {"month": "2026-09"}).status_code, 403)

    def test_balances_list_prints_and_exports(self):
        self.client.force_login(self.hrmo_user)
        resp = self.client.get("/leave/balances/?print=1")
        self.assertContains(resp, self.person_a.full_name)
        self.assertContains(resp, "15.00")
        self.assertEqual(self.client.get("/leave/balances/?export=xlsx").status_code, 200)
        self.assertNotContains(self.client.get("/leave/balances/"), self.person_b.full_name)  # earns no leave

    def test_my_hr_home_page(self):
        self.client.force_login(self.emp_a_user)
        resp = self.client.get("/")
        self.assertContains(resp, "My leave")
        self.assertContains(resp, "10.000")
        self.assertContains(resp, "My appointment")
        self.assertNotContains(resp, "Workforce today")
        self.client.force_login(self.hrmo_user)
        resp = self.client.get("/")
        self.assertContains(resp, "HRMO desk")
        self.assertContains(resp, "Workforce today")
        # Each role gets its own desk; office roles never see the municipality-wide counts.
        self.client.force_login(self.head_a)
        resp = self.client.get("/")
        self.assertContains(resp, "My office")
        self.assertNotContains(resp, "Workforce today")
        self.client.force_login(self.mayor)
        resp = self.client.get("/")
        self.assertContains(resp, "Workforce today")
        self.assertNotContains(resp, "HRMO desk")
        self.client.force_login(self.auditor_user)
        self.assertContains(self.client.get("/"), "Audit desk")
        self.client.force_login(self.sysadmin)
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_file_and_print_opens_the_filled_in_form(self):
        resp = self._file(self.emp_a_user, then="print")
        req = LeaveRequest.objects.get(person=self.person_a)
        self.assertRedirects(resp, f"/leave/{req.pk}/print/?autoprint=1", fetch_redirect_response=False)
        page = self.client.get(f"/leave/{req.pk}/print/?autoprint=1")
        self.assertContains(page, "window.print()")
        self.assertContains(page, "Sec. 51, Rule XVI, Omnibus Rules Implementing E.O. No. 292")  # 6.A legal basis
        self.assertContains(page, "INSTRUCTIONS AND REQUIREMENTS")

    @override_settings(LEAVE_FORM_MAYOR="Juana Dela Cruz")
    def test_blank_form_has_the_employees_details_only(self):
        self.client.force_login(self.emp_a_user)
        page = self.client.get("/leave/print/blank/")
        self.assertContains(page, "Rizal")
        self.assertContains(page, "Executive Assistant IV")
        self.assertContains(page, "Blank form")
        self.assertContains(page, "Juana Dela Cruz")
        self.assertNotContains(page, 'class="box on"')  # nothing ticked
        # HRMO prints one for an employee; others cannot pick a person.
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get(f"/leave/print/blank/?person={self.person_b.pk}"), "Bonifacio")
        self.client.force_login(self.head_a)
        self.assertEqual(self.client.get(f"/leave/print/blank/?person={self.person_b.pk}").status_code, 403)

    def test_register_lists_leave_recorded_without_an_office(self):
        # Leave entered before Step 8 (e.g. by seed_fake) has no office or steps.
        LeaveRequest.objects.create(person=self.person_a, leave_type="vl", start_date=date(2026, 4, 10),
                                    end_date=date(2026, 4, 10), days=1, status=LeaveRequest.Status.APPROVED)
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get("/leave/register/"), self.person_a.full_name)

    def test_menu_has_at_most_six_items(self):
        for user in (self.superuser, self.hrmo_user, self.auditor_user, self.mayor, self.head_a, self.emp_a_user):
            self.client.force_login(user)
            html = self.client.get("/calendar/").content.decode()
            menu = html.split('<ul class="navbar-nav me-auto">', 1)[1].split('<ul class="navbar-nav">', 1)[0]
            self.assertLessEqual(menu.count('<li class="nav-item'), 6, user.username)  # dropdown entries are not nav-items


class PlantillaVacancyMemoTestCase(TestCase):
    """RP1, V1-V5 and M1-M4: the Plantilla report, vacancies and the HRMPSB, memos with acknowledgment."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        from django.utils import timezone

        self.setUp_roles()
        self.today = timezone.localdate()
        self.vacancy = Vacancy.objects.create(
            office=self.office_a, position_title="Clerk I", item_number="SP-MO-099", salary_grade=3,
            monthly_salary=Decimal("14000"), education="Two years of college", eligibility="Subprofessional",
            status=Vacancy.Status.OPEN, published_on=self.today, deadline=self.today + timedelta(days=5),
        )

    def _employee_b(self):
        user_b = get_user_model().objects.create_user("emp_b", "b@test.com", "x")
        UserProfile.objects.create(user=user_b, person=self.person_b)
        return user_b

    def test_plantilla_lists_filled_and_vacant_items_for_hrmo_and_auditor_only(self):
        self.eng_a.appointment_status = Engagement.AppointmentStatus.PERMANENT
        self.eng_a.save()
        self.client.force_login(self.hrmo_user)
        resp = self.client.get("/reports/plantilla/")
        self.assertContains(resp, "SP-MO-001")
        self.assertContains(resp, "SP-MO-099")
        self.assertContains(resp, "(Vacant)")
        self.assertContains(resp, "600,000.00")  # 50,000 a month
        self.assertContains(resp, "01/01/2025")  # original appointment
        self.assertNotContains(resp, "Revenue Collector I")  # casual, not plantilla
        self.assertEqual(self.client.get("/reports/plantilla/?export=xlsx").status_code, 200)
        self.assertEqual(self.client.get("/reports/plantilla/?print=1").status_code, 200)
        for user in (self.mayor, self.head_a, self.emp_a_user):
            self.client.force_login(user)
            self.assertEqual(self.client.get("/reports/plantilla/").status_code, 403, user.username)

    def test_vacancies_are_public_notices_and_employees_apply_online(self):
        draft = Vacancy.objects.create(office=self.office_a, position_title="Driver", item_number="SP-MO-100",
                                       salary_grade=3, monthly_salary=Decimal("14000"), education="HS", eligibility="-")
        self.client.force_login(self.emp_a_user)
        self.assertContains(self.client.get("/vacancies/"), "Clerk I")
        self.assertEqual(self.client.get(f"/vacancies/{draft.pk}/").status_code, 404)
        self.client.post(f"/vacancies/{self.vacancy.pk}/", {"action": "apply"})
        self.client.post(f"/vacancies/{self.vacancy.pk}/", {"action": "apply"})
        applicant = Applicant.objects.get(vacancy=self.vacancy)
        self.assertEqual((applicant.person, applicant.applied_online), (self.person_a, True))
        resp = self.client.get(f"/vacancies/{self.vacancy.pk}/")
        self.assertContains(resp, "You signified interest")
        self.assertNotContains(resp, "Applicants (")
        self.assertEqual(self.client.post(f"/vacancies/{self.vacancy.pk}/", {"action": "add_applicant"}).status_code, 403)
        self.client.force_login(self.head_a)
        self.assertNotContains(self.client.get(f"/vacancies/{self.vacancy.pk}/"), "Applicants (")
        self.client.force_login(self.mayor)
        self.assertContains(self.client.get(f"/vacancies/{self.vacancy.pk}/"), "Applicants (1)")

    def test_applications_close_after_the_deadline(self):
        self.vacancy.published_on = self.today - timedelta(days=10)
        self.vacancy.deadline = self.today - timedelta(days=1)
        self.vacancy.save()
        self.client.force_login(self.emp_a_user)
        self.client.post(f"/vacancies/{self.vacancy.pk}/", {"action": "apply"})
        self.assertFalse(Applicant.objects.exists())

    def test_resolution_draft_names_the_qualified_applicants(self):
        for last, ok in (("Reyes", True), ("Santos", True), ("Cruz", False)):
            Applicant.objects.create(vacancy=self.vacancy, last_name=last, first_name="Ana", signified_on=self.today,
                                     meets_qs=ok)
        self.client.force_login(self.hrmo_user)
        resp = self.client.get(f"/vacancies/{self.vacancy.pk}/print/resolution/")
        self.assertContains(resp, "Ana Reyes and Ana Santos")
        self.assertContains(resp, "respectively")
        self.assertContains(self.client.get(f"/vacancies/{self.vacancy.pk}/print/publication/"), "SP-MO-099")
        self.client.force_login(self.emp_a_user)
        self.assertEqual(self.client.get(f"/vacancies/{self.vacancy.pk}/print/resolution/").status_code, 403)

    def test_memo_reaches_its_audience_and_is_acknowledged_once(self):
        memo = Memo.objects.create(number="Memo 1", dated=self.today, subject="Fiesta schedule", addressed_to="MO staff",
                                   audience=Memo.Audience.OFFICE, office=self.office_a, body="Text",
                                   comply_by=self.today + timedelta(days=2))
        self.client.force_login(self._employee_b())
        self.assertEqual(self.client.get(f"/memos/{memo.pk}/").status_code, 403)
        self.assertNotContains(self.client.get("/memos/"), "Fiesta schedule")

        self.client.force_login(self.emp_a_user)
        self.assertContains(self.client.get("/"), "Fiesta schedule")  # on the home page until read
        self.client.post(f"/memos/{memo.pk}/", {"action": "ack"})
        self.client.post(f"/memos/{memo.pk}/", {"action": "ack"})
        self.assertEqual(memo.acknowledgments.count(), 1)
        self.assertNotContains(self.client.get("/"), "Fiesta schedule")
        self.assertEqual(self.client.get(f"/memos/{memo.pk}/acknowledgments/").status_code, 403)

        self.client.force_login(self.head_a)
        self.assertContains(self.client.get(f"/memos/{memo.pk}/acknowledgments/"), "1 of 1 acknowledged")

    def test_hrmo_records_paper_signatures_only_for_addressees(self):
        memo = Memo.objects.create(number="Memo 2", dated=self.today, subject="ICAS", addressed_to="All", body="Text")
        self.client.force_login(self.hrmo_user)
        self.client.post(f"/memos/{memo.pk}/acknowledgments/", {"person": [self.person_b.pk, 999999]})
        self.assertEqual(list(memo.acknowledgments.values_list("person_id", "from_paper")), [(self.person_b.pk, True)])
        self.client.force_login(self.head_a)
        resp = self.client.post(f"/memos/{memo.pk}/acknowledgments/", {"person": [self.person_a.pk]})
        self.assertEqual(resp.status_code, 403)

    def test_memo_scan_must_be_a_real_file(self):
        self.client.force_login(self.hrmo_user)
        resp = self.client.post("/memos/new/", {
            "number": "Memo 3", "dated": self.today.isoformat(), "subject": "S", "addressed_to": "All",
            "audience": "everyone", "body": "B", "requires_ack": "on", "is_published": "on",
            "scan": SimpleUploadedFile("memo.pdf", b"not a pdf", content_type="application/pdf"),
        })
        self.assertContains(resp, "not a valid PDF")
        self.assertFalse(Memo.objects.exists())

    def test_calendar_shows_vacancy_and_memo_deadlines(self):
        Memo.objects.create(number="Memo 4", dated=self.today, subject="Submit ICAS forms", addressed_to="MO",
                            audience=Memo.Audience.OFFICE, office=self.office_a, body="B", comply_by=self.vacancy.deadline)
        url = f"/calendar/day/{self.vacancy.deadline.isoformat()}/"
        self.client.force_login(self.emp_a_user)
        resp = self.client.get(url)
        self.assertContains(resp, "Last day to apply: Clerk I")
        self.assertContains(resp, "Submit ICAS forms")
        self.client.force_login(self._employee_b())
        self.assertNotContains(self.client.get(url), "Submit ICAS forms")

    def test_seed_demo_board_adds_once(self):
        from io import StringIO

        with override_settings(DEBUG=True):
            call_command("seed_demo_board", stdout=StringIO())
            out = StringIO()
            call_command("seed_demo_board", stdout=out)
        self.assertIn("already added", out.getvalue())
        self.assertEqual(Vacancy.objects.filter(notes__startswith="[demo board]").count(), 2)
        self.assertTrue(Memo.objects.filter(number__endswith="(sample)").exists())


class CertificateRequestTestCase(TestCase):
    """CR1-CR6: service records and certificates of employment, requested online or at the counter."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()
        User = get_user_model()
        self.emp_a = self.emp_a_user
        # A former employee who kept an account: resigned, no active appointment.
        self.former = Person.objects.create(employee_no="SP-00303", last_name="Mabini", first_name="Apolinario",
                                            sex="M", birthdate=date(1964, 7, 23))
        Engagement.objects.create(person=self.former, office=self.office_b, type=Engagement.Type.CASUAL,
                                  position_title="Clerk I", start_date=date(2001, 1, 2), end_date=date(2004, 12, 31),
                                  status=Engagement.Status.ENDED)
        Engagement.objects.create(person=self.former, office=self.office_b, type=Engagement.Type.PLANTILLA,
                                  position_title="Clerk II", salary_grade=6, step=2, item_number="SP-MTO-090",
                                  rate=Decimal("20000.00"), rate_basis=Engagement.RateBasis.MONTHLY,
                                  start_date=date(2005, 1, 3), end_date=date(2020, 6, 30),
                                  status=Engagement.Status.RESIGNED)
        self.former_user = User.objects.create_user("former", "former@test.com", "formerpass123")
        UserProfile.objects.create(user=self.former_user, person=self.former)

    def _file(self, user, **data):
        self.client.force_login(user)
        return self.client.post("/certificates/new/", {"kind": "service_record", "purpose": "GSIS claim",
                                                       "copies": 1, "contact": "", **data})

    # --- Service record rows (CR1) ----------------------------------------------------------

    def test_service_rows_split_at_step_increments_and_mark_separation(self):
        from hr.certificates import service_rows

        eng = self.former.engagements.get(type=Engagement.Type.PLANTILLA)
        SalaryAction.objects.create(kind=SalaryAction.Kind.STEP_INCREMENT, person=self.former, engagement=eng,
                                    milestone="Step 2", status=SalaryAction.Status.ISSUED, service_from=date(2005, 1, 3),
                                    due_date=date(2008, 1, 3), from_step=1, to_step=2, effective_date=date(2008, 1, 3))
        ServiceInterruption.objects.create(person=self.former, kind=ServiceInterruption.Kind.LWOP,
                                           start_date=date(2010, 3, 1), end_date=date(2010, 3, 10))
        Engagement.objects.create(person=self.former, office=self.office_b, type=Engagement.Type.JOB_ORDER,
                                  position_title="Encoder", start_date=date(2021, 1, 4), end_date=date(2021, 6, 30),
                                  status=Engagement.Status.ENDED)
        rows = service_rows(self.former)
        self.assertEqual([(r["designation"], r["start"], r["end"]) for r in rows], [
            ("Clerk I", date(2001, 1, 2), date(2004, 12, 31)),
            ("Clerk II", date(2005, 1, 3), date(2008, 1, 2)),
            ("Clerk II", date(2008, 1, 3), date(2020, 6, 30)),
        ])  # the job order is not government service
        self.assertEqual(rows[0]["cause"], "")  # followed by a new appointment, not a separation
        self.assertEqual([r["grade"] for r in rows[1:]], ["SG 6, step 1", "SG 6, step 2"])
        self.assertEqual(rows[1]["salary"], "")
        self.assertTrue(rows[1]["salary_missing"])
        self.assertEqual(rows[2]["salary"], "₱240,000.00 p.a.")
        self.assertEqual((rows[2]["lwop"], rows[2]["cause"], rows[2]["separated_on"]), (10, "Resigned", date(2020, 6, 30)))

    def test_last_ended_contract_is_end_of_contract(self):
        from hr.certificates import service_rows

        self.former.engagements.filter(type=Engagement.Type.PLANTILLA).delete()
        self.assertEqual(service_rows(self.former)[-1]["cause"], "End of contract")

    # --- Requests (CR2) ---------------------------------------------------------------------

    def test_employee_requests_online_with_control_number(self):
        resp = self._file(self.emp_a, contact="0917-000-0000")
        req = CertificateRequest.objects.get()
        self.assertRedirects(resp, f"/certificates/{req.pk}/")
        self.assertEqual((req.person, req.received_via, req.requester, req.status, req.filed_by),
                         (self.person_a, "online", "self", "pending", self.emp_a))
        self.assertRegex(req.control_no, r"^CR-\d{4}-0001$")
        self._file(self.emp_a, kind="coe")
        self.assertTrue(CertificateRequest.objects.filter(control_no__endswith="-0002").exists())

    def test_former_employee_can_request_and_sees_past_tense(self):
        self._file(self.former_user, kind="coe_pay", purpose="Employment abroad")
        req = CertificateRequest.objects.get(person=self.former)
        self.client.force_login(self.hrmo_user)
        resp = self.client.get(f"/certificates/{req.pk}/print/")
        self.assertContains(resp, "was employed")
        self.assertContains(resp, "MR. APOLINARIO MABINI")
        self.assertContains(resp, "₱20,000.00 a month")
        self.assertContains(resp, req.control_no)
        self.assertNotContains(resp, "PREVIEW")

    def test_job_order_only_cannot_ask_for_service_record(self):
        self.person_a.engagements.update(type=Engagement.Type.JOB_ORDER, salary_grade=None, step=None, item_number="")
        resp = self._file(self.emp_a)
        self.assertContains(resp, "not government service")
        self.assertFalse(CertificateRequest.objects.exists())

    def test_open_request_limit(self):
        for _ in range(5):
            self._file(self.emp_a)
        resp = self._file(self.emp_a)
        self.assertContains(resp, "already have 5 requests")
        self.assertEqual(CertificateRequest.objects.count(), 5)

    def test_hrmo_records_request_from_heir_of_former_employee(self):
        self.client.force_login(self.hrmo_user)
        url = f"/people/{self.former.pk}/certificates/new/"
        resp = self.client.post(url, {"kind": "service_record", "purpose": "GSIS survivorship claim", "copies": 2,
                                      "received_via": "letter", "requester": "heir", "requester_name": ""})
        self.assertContains(resp, "Give the name of the representative or heir")
        self.client.post(url, {"kind": "service_record", "purpose": "GSIS survivorship claim", "copies": 2,
                               "received_via": "letter", "requester": "heir", "requester_name": "Maria Mabini"})
        req = CertificateRequest.objects.get()
        self.assertEqual((req.requester_name, req.received_via, req.copies), ("Maria Mabini", "letter", 2))
        resp = self.client.get(f"/certificates/{req.pk}/print/")
        self.assertContains(resp, "SERVICE RECORD")
        self.assertContains(resp, "Maria Mabini (heir or next of kin)")
        # HRMO's form cannot claim the request was filed online by the person.
        self.client.post(url, {"kind": "coe", "purpose": "x", "copies": 1, "received_via": "online",
                               "requester": "self"})
        self.assertEqual(CertificateRequest.objects.count(), 1)

    # --- Who sees what (CR5) ----------------------------------------------------------------

    def test_print_is_hrmo_and_auditor_only_and_audited(self):
        self._file(self.emp_a)
        req = CertificateRequest.objects.get()
        self.assertEqual(self.client.get(f"/certificates/{req.pk}/").status_code, 200)
        self.assertEqual(self.client.get(f"/certificates/{req.pk}/print/").status_code, 403)  # pay data (R6)
        for user in (self.head_a, self.mayor, self.staff_a, self.sysadmin, self.former_user):
            self.client.force_login(user)
            self.assertEqual(self.client.get(f"/certificates/{req.pk}/").status_code, 403, user)
            self.assertEqual(self.client.get("/certificates/").status_code, 403, user)
        self.client.force_login(self.auditor_user)
        self.assertEqual(self.client.get(f"/certificates/{req.pk}/print/").status_code, 200)
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.Action.EXPORT, object_repr__contains=req.control_no).exists())
        self.assertEqual(self.client.get(f"/people/{self.person_a.pk}/certificates/new/").status_code, 403)
        self.client.force_login(self.hrmo_user)
        self.assertEqual(self.client.get(f"/people/{self.person_a.pk}/certificates/new/").status_code, 200)

    def test_preview_from_person_record(self):
        self.client.force_login(self.hrmo_user)
        resp = self.client.get(f"/people/{self.person_a.pk}/preview/service_record/")
        self.assertContains(resp, "NOT FOR RELEASE")
        self.assertContains(resp, "Executive Assistant IV")
        self.assertEqual(self.client.get(f"/people/{self.person_a.pk}/preview/bogus/").status_code, 403)
        self.client.force_login(self.emp_a)
        self.assertEqual(self.client.get(f"/people/{self.person_a.pk}/preview/coe/").status_code, 403)

    # --- Workflow (CR3) ---------------------------------------------------------------------

    def _act(self, user, req, action, **data):
        self.client.force_login(user)
        req.refresh_from_db()
        return self.client.post(f"/certificates/{req.pk}/", {"action": action, "version": req.version, **data})

    def test_prepare_release_and_home_page(self):
        self._file(self.emp_a)
        req = CertificateRequest.objects.get()
        self.assertEqual(self._act(self.emp_a, req, "ready").status_code, 403)
        self._act(self.hrmo_user, req, "ready")
        req.refresh_from_db()
        self.assertEqual((req.status, req.prepared_by), ("ready", self.hrmo_user))
        self.client.force_login(self.emp_a)
        self.assertContains(self.client.get("/"), "document ready for release at HRMO")
        self._act(self.hrmo_user, req, "release", released_to="")
        req.refresh_from_db()
        self.assertEqual(req.status, "ready")
        self._act(self.hrmo_user, req, "release", released_to="Jose Rizal", note="UMID")
        req.refresh_from_db()
        self.assertEqual((req.status, req.released_to, req.note), ("released", "Jose Rizal", "UMID"))
        self.assertIsNotNone(req.released_at)

    def test_decline_needs_reason_and_stale_version_refused(self):
        self._file(self.emp_a)
        req = CertificateRequest.objects.get()
        self._act(self.hrmo_user, req, "decline")
        req.refresh_from_db()
        self.assertEqual(req.status, "pending")
        self.client.post(f"/certificates/{req.pk}/", {"action": "ready", "version": req.version - 1})
        req.refresh_from_db()
        self.assertEqual(req.status, "pending")
        self._act(self.hrmo_user, req, "decline", note="Bring an authorization letter.")
        req.refresh_from_db()
        self.assertEqual(req.status, "declined")

    def test_only_the_filer_cancels_while_pending(self):
        self._file(self.emp_a)
        req = CertificateRequest.objects.get()
        self.assertEqual(self._act(self.hrmo_user, req, "cancel").status_code, 403)
        self._act(self.emp_a, req, "cancel")
        req.refresh_from_db()
        self.assertEqual(req.status, "cancelled")
        self.client.force_login(self.hrmo_user)
        self.client.post(f"/people/{self.person_a.pk}/certificates/new/", {
            "kind": "coe", "purpose": "Loan", "copies": 1, "received_via": "walk_in", "requester": "self"})
        recorded = CertificateRequest.objects.get(received_via="walk_in")
        self.assertEqual(self._act(self.emp_a, recorded, "cancel").status_code, 403)

    def test_db_requires_release_details(self):
        self._file(self.emp_a)
        req = CertificateRequest.objects.get()
        req.status = CertificateRequest.Status.RELEASED
        with self.assertRaises(IntegrityError), transaction.atomic():
            req.save()


class PersonalDataSheetTestCase(TestCase):
    """PD1-PD7: the PDS filled in online, verified by HRMO, and printed as CS Form No. 212."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()
        self.emp = self.emp_a_user
        self.url = f"/people/{self.person_a.pk}/pds/"

    def _complete(self, person):
        """Fill every required section through the models."""
        from hr.models import PdsEducation, PdsProfile, PdsReference

        answers = {q: False for q in ["q34a", "q34b", "q35a", "q35b", "q36", "q37", "q38a", "q38b", "q39",
                                       "q40a", "q40b", "q40c"]}
        PdsProfile.objects.update_or_create(person=person, defaults=dict(
            place_of_birth="San Pascual, Masbate", civil_status="single", res_barangay="Poblacion",
            res_city="San Pascual", res_province="Masbate", father_surname="Rizal", father_first_name="Francisco",
            mother_surname="Alonso", mother_first_name="Teodora", gov_id_type="UMID", gov_id_number="0011-2233",
            **answers))
        PdsEducation.objects.create(person=person, level="college", school="Masbate State College", degree="BSBA",
                                    year_from=1998, year_to=2002, year_graduated=2002)
        for n in range(3):
            PdsReference.objects.create(person=person, name=f"Reference {n}", address="San Pascual")

    def _rows(self, prefix, rows, initial=0):
        data = {f"{prefix}-TOTAL_FORMS": len(rows), f"{prefix}-INITIAL_FORMS": initial,
                f"{prefix}-MIN_NUM_FORMS": 0, f"{prefix}-MAX_NUM_FORMS": 1000}
        for i, row in enumerate(rows):
            data.update({f"{prefix}-{i}-{k}": v for k, v in row.items()})
        return data

    # --- Filling in (PD2) -------------------------------------------------------------------

    def test_employee_fills_sections_then_sends(self):
        from hr.models import PdsSubmission

        self.client.force_login(self.emp)
        self.assertContains(self.client.get("/"), "sections of your Personal Data Sheet to fill in")
        self.client.post(f"{self.url}submit/")
        self.assertFalse(PdsSubmission.objects.exists())  # required sections still empty
        resp = self.client.post(f"{self.url}personal/", {
            "version": 1, "place_of_birth": "San Pascual", "civil_status": "married", "citizenship": "filipino",
            "res_barangay": "Poblacion", "res_city": "San Pascual", "perm_same": "on", "height_m": "1.65",
        })
        self.assertRedirects(resp, self.url)
        resp = self.client.post(f"{self.url}family/", {
            "version": 2, "father_surname": "Rizal", "mother_surname": "Alonso",
            **self._rows("family", [{"full_name": "Maria Rizal", "birthdate": "2010-05-01"}]),
        })
        self.assertRedirects(resp, self.url)
        self.assertContains(self.client.get(self.url), "Fill in every required section")  # married, no spouse yet
        self.client.post(f"{self.url}family/", {"version": 3, "father_surname": "Rizal", "mother_surname": "Alonso",
                                                "spouse_surname": "Bracken", **self._rows("family", [])})
        self.client.post(f"{self.url}education/", self._rows("education", [
            {"level": "college", "school": "Masbate State College", "year_from": 1998, "year_to": 2002}]))
        self.client.post(f"{self.url}questions/", {"version": 4, **{q: "False" for q in [
            "q34a", "q34b", "q35a", "q35b", "q36", "q37", "q38a", "q38b", "q39", "q40a", "q40b", "q40c"]}})
        self.client.post(f"{self.url}references/", {
            "version": 5, "gov_id_type": "UMID", "gov_id_number": "00-1234",
            **self._rows("references", [{"name": f"Ref {n}"} for n in range(3)]),
        })
        resp = self.client.post(f"{self.url}submit/")
        sub = PdsSubmission.objects.get()
        self.assertEqual((sub.status, sub.submitted_by), ("submitted", self.emp))
        self.assertEqual(sub.snapshot["children"][0]["full_name"], "Maria Rizal")
        self.assertEqual(sub.snapshot["profile"]["spouse_surname"], "Bracken")
        self.assertNotContains(self.client.get("/"), "Personal Data Sheet to fill in")

    def test_yes_needs_details_and_fake_ids_only(self):
        self.client.force_login(self.emp)
        answers = {q: "False" for q in ["q34a", "q34b", "q35a", "q35b", "q36", "q37", "q38a", "q38b", "q39",
                                         "q40a", "q40b", "q40c"]}
        resp = self.client.post(f"{self.url}questions/", {"version": 1, **answers, "q34a": "True", "q40b": "True"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.content.decode().count("You answered Yes: give the details."), 2)
        resp = self.client.post(f"{self.url}references/", {"version": 1, "gov_id_type": "UMID",
                                                           "gov_id_number": "1234-5678", **self._rows("references", [])})
        self.assertContains(resp, "fake ID numbers only")
        resp = self.client.post(f"{self.url}references/", {"version": 1, **self._rows(
            "references", [{"name": f"Ref {n}"} for n in range(4)])})
        self.assertContains(resp, "Please submit at most 3 forms")

    # --- Verification (PD4) -----------------------------------------------------------------

    def test_verify_then_changes_are_marked_and_resent(self):
        from hr import pds
        from hr.models import PdsEducation, PdsSubmission

        self._complete(self.person_a)
        self.client.force_login(self.emp)
        self.client.post(f"{self.url}submit/")
        self.client.post(f"{self.url}submit/")  # a second send replaces the waiting copy
        sub = PdsSubmission.objects.get()
        self.client.force_login(self.hrmo_user)
        self.client.post(f"/pds/review/{sub.pk}/", {"action": "return", "version": sub.version})
        sub.refresh_from_db()
        self.assertEqual(sub.status, "submitted")  # a reason is required
        self.client.post(f"/pds/review/{sub.pk}/", {"action": "verify", "version": sub.version})
        sub.refresh_from_db()
        self.assertEqual((sub.status, sub.reviewed_by), ("verified", self.hrmo_user))
        self.assertEqual(pds.state(self.person_a)["unsent"], set())
        # HRMO correcting the mobile number is not a change the employee must resend.
        Person.objects.filter(pk=self.person_a.pk).update(mobile="0917-000-0000")
        self.assertEqual(pds.state(self.person_a)["unsent"], set())
        PdsEducation.objects.create(person=self.person_a, level="graduate", school="Bicol University")
        self.assertEqual(pds.state(self.person_a)["unsent"], {"education"})
        self.client.force_login(self.emp)
        self.assertContains(self.client.get(self.url), "You changed 1 section")
        self.client.post(f"{self.url}submit/")
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get("/pds/queue/"), "<td>III. Educational background</td>", html=False)
        waiting = PdsSubmission.objects.get(status="submitted")
        self.client.post(f"/pds/review/{waiting.pk}/", {"action": "return", "version": waiting.version,
                                                        "note": "Attach the diploma."})
        self.client.force_login(self.emp)
        self.assertContains(self.client.get("/"), "Personal Data Sheet returned by HRMO")
        self.assertContains(self.client.get(self.url), "Attach the diploma.")

    def test_hrmo_encodes_paper_pds_and_verifies(self):
        from hr.models import PdsSubmission

        self._complete(self.person_b)
        self.client.force_login(self.hrmo_user)
        self.assertEqual(self.client.get(f"/people/{self.person_b.pk}/pds/personal/").status_code, 200)
        self.client.post(f"/people/{self.person_b.pk}/pds/submit/", {"verify": "1"})
        sub = PdsSubmission.objects.get(person=self.person_b)
        self.assertEqual((sub.status, sub.reviewed_by), ("verified", self.hrmo_user))
        # An employee cannot verify their own.
        self._complete(self.person_a)
        self.client.force_login(self.emp)
        self.client.post(f"{self.url}submit/", {"verify": "1"})
        self.assertEqual(PdsSubmission.objects.get(person=self.person_a).status, "submitted")

    # --- Who sees what (PD5) ----------------------------------------------------------------

    def test_only_self_hrmo_and_auditor(self):
        from hr.models import PdsSubmission

        self._complete(self.person_a)
        sub = PdsSubmission.objects.create(person=self.person_a, snapshot={})
        for user in (self.head_a, self.mayor, self.staff_a, self.sysadmin):
            self.client.force_login(user)
            for url in (self.url, f"{self.url}print/", f"{self.url}personal/", f"/pds/review/{sub.pk}/"):
                self.assertEqual(self.client.get(url).status_code, 403, (user, url))
            self.assertEqual(self.client.get("/pds/queue/").status_code, 403, user)
        self.client.force_login(self.auditor_user)
        self.assertEqual(self.client.get(self.url).status_code, 200)
        self.assertEqual(self.client.get(f"{self.url}personal/").status_code, 403)
        self.assertEqual(self.client.post(f"/pds/review/{sub.pk}/", {"action": "verify", "version": 1}).status_code, 403)
        self.assertEqual(self.client.get(f"/people/{self.person_b.pk}/pds/").status_code, 200)
        self.client.force_login(self.emp)
        self.assertEqual(self.client.get(f"/people/{self.person_b.pk}/pds/").status_code, 403)

    def test_print_fills_lgu_work_and_hides_pay_from_employee(self):
        self._complete(self.person_a)
        self.client.force_login(self.emp)
        resp = self.client.get(f"{self.url}print/")
        self.assertContains(resp, "PERSONAL DATA SHEET")
        self.assertContains(resp, "Executive Assistant IV")
        self.assertContains(resp, "Not yet verified by HRMO")
        self.assertNotContains(resp, "50,000.00")
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.Action.EXPORT, object_repr__startswith="PDS print").exists())
        self.client.force_login(self.hrmo_user)
        resp = self.client.get(f"{self.url}print/")
        self.assertContains(resp, "50,000.00")
        self.assertContains(resp, "22-1")

    def test_ids_and_answers_are_redacted_in_audit(self):
        from hr.models import PdsProfile

        self._complete(self.person_a)
        profile = PdsProfile.objects.get(person=self.person_a)
        profile.gov_id_number = "00-9999"
        profile.q36_details = "Sample case"
        profile.save()
        entry = AuditLog.objects.filter(model_name="PdsProfile").order_by("-pk").first()
        self.assertEqual(entry.changes["gov_id_number"], "[REDACTED]")
        self.assertEqual(entry.changes["q36_details"], "[REDACTED]")

    def test_opening_a_section_changes_nothing(self):
        from hr.models import PdsProfile

        self.client.force_login(self.emp)
        for section in ("personal", "family", "questions", "references"):
            self.assertEqual(self.client.get(f"{self.url}{section}/").status_code, 200)
        self.assertFalse(PdsProfile.objects.filter(person=self.person_a).exists())
        resp = self.client.post(f"{self.url}other/", {"version": 1, "skills": "Driving"})
        self.assertRedirects(resp, self.url)
        self.assertEqual(PdsProfile.objects.get(person=self.person_a).skills, "Driving")


class TalentTrackersTestCase(TestCase):
    """EL1-EL3, LD1-LD3, SP1-SP5, SN1-SN4, CT1-CT5: licences, training hours, ratings, SALN, COC/CTO."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()
        self.today = date(2026, 10, 20)

    # --- Licences and eligibility (EL1-EL3) -------------------------------------------------

    def test_licence_rows_and_reminders(self):
        from hr import qualifications
        from hr.models import PdsEligibility

        PdsEligibility.objects.create(person=self.person_a, name="Registered Nurse (RA 1080)",
                                      license_number="0012345", license_valid_until=date(2026, 9, 1))
        PdsEligibility.objects.create(person=self.person_b, name="Career Service Professional",
                                      license_valid_until=date(2026, 12, 1))
        PdsEligibility.objects.create(person=self.person_b, name="Driver's license",
                                      license_valid_until=date(2029, 1, 1))
        rows = qualifications.licence_rows(self.today)
        self.assertEqual([(r["person"], r["state"]) for r in rows],
                         [(self.person_a, "expired"), (self.person_b, "soon")])
        self.assertEqual(rows[0]["ago"], 49)
        found = qualifications.licence_rows(self.today, query="driver")
        self.assertEqual([r["state"] for r in found], ["valid"])
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get("/tracker/licences/"), "Registered Nurse")
        self.client.force_login(self.emp_a_user)
        self.assertContains(self.client.get("/"), "licence on your PDS expired or expires soon")
        self.client.force_login(self.head_a)
        self.assertEqual(self.client.get("/tracker/licences/").status_code, 403)

    def test_eligibility_match_is_a_hint(self):
        from hr.qualifications import eligibility_match

        wanted = "Career Service (Professional) / Second Level Eligibility"
        self.assertEqual(eligibility_match(wanted, ["Career Service Professional", "Career Service Sub-Professional",
                                                    "Driver's license"]), ["Career Service Professional"])

    def test_applying_prefills_requirements_and_hrmo_sees_evidence(self):
        from hr.models import Applicant, PdsEducation, PdsEligibility, PdsSubmission, PerformanceRating

        vacancy = Vacancy.objects.create(office=self.office_a, position_title="Clerk I", item_number="SP-MO-099",
                                         salary_grade=3, monthly_salary=Decimal("15000"), education="College",
                                         eligibility="Career Service (Professional)", status=Vacancy.Status.OPEN,
                                         published_on=date.today(), deadline=date.today() + timedelta(days=10))
        PdsEligibility.objects.create(person=self.person_a, name="Career Service Professional")
        PdsEducation.objects.create(person=self.person_a, level="college", school="Masbate State College",
                                    degree="BSBA", year_graduated=2002)
        PdsSubmission.objects.create(person=self.person_a, snapshot={}, status="verified")
        PerformanceRating.objects.create(person=self.person_a, year=2026, semester=1, rating=Decimal("4.6"))
        self.client.force_login(self.emp_a_user)
        self.client.post(f"/vacancies/{vacancy.pk}/", {"action": "apply"})
        applicant = Applicant.objects.get()
        self.assertEqual((applicant.has_pds, applicant.has_rating, applicant.has_eligibility, applicant.has_tor),
                         (True, True, True, False))
        self.client.force_login(self.hrmo_user)
        page = self.client.get(f"/vacancies/{vacancy.pk}/")
        self.assertContains(page, "BSBA")
        self.assertContains(page, "4.600 Outstanding")
        self.assertContains(page, 'title="Matches the eligibility asked for; HRMO decides"')

    # --- Learning and development (LD1-LD3) -------------------------------------------------

    def test_record_training_for_attendees_without_reverification(self):
        from hr import pds
        from hr.models import PdsSubmission, PdsTraining

        before = pds.snapshot(self.person_a)
        PdsSubmission.objects.create(person=self.person_a, snapshot=before, status="verified")
        self.client.force_login(self.hrmo_user)
        resp = self.client.post("/tracker/learning/record/", {
            "title": "Records Management Seminar", "date_from": "2026-03-02", "date_to": "2026-03-03",
            "hours": "16", "kind": "technical", "sponsor": "CSC RO V", "person": [self.person_a.pk, self.person_b.pk],
        })
        self.assertRedirects(resp, "/tracker/learning/")
        self.assertEqual(PdsTraining.objects.filter(by_hrmo=True).count(), 2)
        self.assertEqual(pds.state(self.person_a)["unsent"], set())  # not a change the employee must send
        page = self.client.get("/tracker/learning/?year=2026")
        self.assertEqual(page.context["total"], 32)
        self.client.force_login(self.emp_a_user)
        page = self.client.get(f"/people/{self.person_a.pk}/pds/training/")
        self.assertContains(page, "Recorded by HRMO")
        self.assertEqual(page.context["formset"].queryset.count(), 0)  # read-only to the employee
        self.client.force_login(self.auditor_user)
        self.assertEqual(self.client.get("/tracker/learning/record/").status_code, 403)

    # --- SPMS ratings (SP1-SP5) -------------------------------------------------------------

    def test_office_sheet_saves_rejects_and_removes(self):
        from hr.models import PerformanceRating

        url = f"/tracker/ratings/{self.office_a.pk}/?year=2026&semester=1"
        self.client.force_login(self.hrmo_user)
        resp = self.client.post(url, {f"rating-{self.person_a.pk}": "6"})
        self.assertContains(resp, "Nothing was saved")
        self.client.post(url, {f"rating-{self.person_a.pk}": "3.75", f"note-{self.person_a.pk}": "IPCR on file"})
        rating = PerformanceRating.objects.get(person=self.person_a)
        self.assertEqual((rating.rating, rating.adjectival, rating.office), (Decimal("3.750"), "Very Satisfactory", self.office_a))
        self.client.post(url, {f"rating-{self.person_a.pk}": ""})
        self.assertFalse(PerformanceRating.objects.exists())
        self.client.force_login(self.auditor_user)
        self.assertEqual(self.client.post(url, {f"rating-{self.person_a.pk}": "4"}).status_code, 403)

    def test_ratings_tick_the_step_increment_checklist(self):
        from hr import spms
        from hr.models import PerformanceRating

        case = SalaryAction.objects.create(kind=SalaryAction.Kind.STEP_INCREMENT, person=self.person_a,
                                           engagement=self.eng_a, milestone="Step 2", service_from=date(2025, 1, 1),
                                           due_date=date(2026, 6, 30), from_step=1, to_step=2,
                                           status=SalaryAction.Status.DUE)
        for year, sem, value in [(2025, 1, "3.2"), (2025, 2, "2.6")]:
            PerformanceRating.objects.create(person=self.person_a, year=year, semester=sem, rating=Decimal(value))
        self.assertFalse(spms.prefill_checklist(case, self.today))  # Jan-Jun 2026 missing
        PerformanceRating.objects.create(person=self.person_a, year=2026, semester=1, rating=Decimal("4"))
        self.client.force_login(self.hrmo_user)
        self.client.post(f"/salary-actions/{case.pk}/", {"action": "start_review", "version": case.version})
        case.refresh_from_db()
        self.assertTrue(case.checklist.get("rating"))
        PerformanceRating.objects.filter(year=2025, semester=2).update(rating=Decimal("2.4"))
        self.assertFalse(spms.case_coverage(case, self.today)["all_satisfactory"])

    def test_who_sees_ratings(self):
        from hr.models import PerformanceRating
        from hr.permissions import can_view_ratings

        PerformanceRating.objects.create(person=self.person_a, year=2026, semester=1, rating=Decimal("4"))
        self.assertTrue(can_view_ratings(self.head_a, self.person_a))
        self.assertFalse(can_view_ratings(self.head_a, self.person_b))
        self.assertFalse(can_view_ratings(self.staff_a, self.person_a))
        self.assertTrue(can_view_ratings(self.emp_a_user, self.person_a))
        self.client.force_login(self.head_a)
        self.assertContains(self.client.get(f"/people/{self.person_a.pk}/"), "Performance ratings (IPCR)")

    # --- SALN (SN1-SN4) -----------------------------------------------------------------------

    @override_settings(SALN_FIRST_YEAR=2020)
    def test_saln_reported_then_received(self):
        from django.utils import timezone

        from hr import saln
        from hr.models import SalnFiling

        today = timezone.localdate()
        year = saln.filing_year(today)
        self.assertTrue(saln.is_covered(self.person_a))  # plantilla
        self.assertFalse(saln.is_covered(self.person_b))  # casual
        self.client.force_login(self.emp_a_user)
        if (saln.deadline(year) - today).days <= saln.REMIND_DAYS:
            self.assertContains(self.client.get("/"), f"SALN for {year}")
        self.client.post("/saln/")
        filing = SalnFiling.objects.get(person=self.person_a, year=year)
        self.assertIsNotNone(filing.reported_at)
        self.client.get("/saln/")  # shows the "Recorded" message
        self.assertNotContains(self.client.get("/"), f"SALN for {year}")
        self.client.force_login(self._b_user())
        self.assertEqual(self.client.post("/saln/").status_code, 403)
        self.client.force_login(self.hrmo_user)
        self.client.post(f"/tracker/saln/?year={year}", {"person": [self.person_a.pk, self.person_b.pk],
                                                          "received_on": today.isoformat()})
        filing.refresh_from_db()
        self.assertEqual((filing.received_on, filing.received_by), (today, self.hrmo_user))
        self.assertFalse(SalnFiling.objects.filter(person=self.person_b).exists())  # not covered
        self.client.force_login(self.head_a)
        page = self.client.get("/tracker/saln/")
        self.assertContains(page, "Rizal")
        self.assertEqual(self.client.post("/tracker/saln/", {"person": [self.person_a.pk]}).status_code, 403)
        self.client.force_login(self.staff_a)
        self.assertEqual(self.client.get("/tracker/saln/").status_code, 403)

    def test_saln_not_before_first_year_or_year_end(self):
        from hr import saln

        with override_settings(SALN_FIRST_YEAR=2026):
            self.assertIsNone(saln.my_reminder(self.person_a, date(2026, 4, 1)))  # 2025 was filed on paper
            self.assertEqual(saln.filing_year(date(2026, 10, 1)), 2026)
            self.assertFalse(saln.is_open(2026, date(2026, 10, 1)))
            self.assertIsNotNone(saln.my_reminder(self.person_a, date(2027, 4, 1)))

    def _b_user(self):
        user = get_user_model().objects.create_user("emp_b", "emp_b@test.com", "emppass123")
        UserProfile.objects.create(user=user, person=self.person_b)
        return user

    # --- Overtime credits and CTO (CT1-CT5) ---------------------------------------------------

    def test_coc_rates_caps_and_expiry(self):
        from hr import coc

        entry, credited, short = coc.credit(self.person_a, date(2026, 9, 7), Decimal("3"), "workday", "OO 1", None, self.today)
        self.assertEqual((credited, short, entry.expires_on), (Decimal("3.00"), Decimal("0"), date(2027, 9, 7)))
        _, credited, _ = coc.credit(self.person_a, date(2026, 9, 12), Decimal("4"), "restday", "OO 2", None, self.today)
        self.assertEqual(credited, Decimal("6.00"))
        _, credited, short = coc.credit(self.person_a, date(2026, 9, 13), Decimal("24"), "restday", "OO 3", None, self.today)
        self.assertEqual((credited, short), (Decimal("31.00"), Decimal("5.00")))  # 40 a month
        self.assertEqual(coc.balance(self.person_a, self.today), Decimal("40"))
        self.assertEqual(coc.balance(self.person_a, date(2027, 9, 8)), Decimal("37"))  # the first lot expired

    def test_cto_charged_on_approval_and_restored_on_cancel(self):
        from hr import coc
        from hr.models import CocEntry, LeaveRequest

        coc.credit(self.person_a, date(2026, 9, 7), Decimal("8"), "workday", "OO 1", None, self.today)
        self.client.force_login(self.emp_a_user)
        self.client.post("/leave/new/", {"leave_type": "cto", "detail": "", "detail_text": "", "start_date": "2026-11-02",
                                         "end_date": "2026-11-02", "days": "", "reason": "", "leave_type_other": ""})
        req = LeaveRequest.objects.get(leave_type="cto")
        self.assertEqual(coc.held_hours(self.person_a), Decimal("8"))
        for user in (self.head_a, self.hrmo_user, self.mayor):
            self.client.force_login(user)
            req.refresh_from_db()
            self.client.post(f"/leave/{req.pk}/", {"action": "done", "version": req.version})
        req.refresh_from_db()
        self.assertEqual(req.status, "approved")
        self.assertEqual(coc.balance(self.person_a, date(2026, 11, 3)), Decimal("0"))
        self.assertEqual(CocEntry.objects.get(kind="used").hours, Decimal("-8"))
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get(f"/leave/{req.pk}/print/"), "Compensatory time-off (CTO)")
        req.refresh_from_db()
        self.client.post(f"/leave/{req.pk}/", {"action": "cancel", "version": req.version, "note": "Recalled to work"})
        self.assertEqual(coc.balance(self.person_a, date(2026, 11, 3)), Decimal("8"))

    def test_job_order_earns_no_coc(self):
        from hr import coc
        from hr.models import CocEntry

        self.eng_b.type = Engagement.Type.JOB_ORDER
        self.eng_b.save()
        self.assertFalse(coc.earns_coc(self.person_b))
        self.client.force_login(self.hrmo_user)
        resp = self.client.post("/leave/coc/record/", {"worked_on": "2026-09-07", "hours": "2", "day_kind": "workday",
                                                       "reference": "OO 4", "person": [self.person_a.pk, self.person_b.pk]})
        self.assertRedirects(resp, "/leave/coc/record/")
        self.assertEqual(list(CocEntry.objects.values_list("person_id", flat=True)), [self.person_a.pk])
        self.client.force_login(self.emp_a_user)
        self.assertEqual(self.client.get(f"/leave/coc/{self.person_a.pk}/").status_code, 200)
        self.assertEqual(self.client.get(f"/leave/coc/{self.person_b.pk}/").status_code, 403)


class AppVersionTestCase(TestCase):
    """H9: the release number in the footer, with the exact build on hover."""

    @override_settings(APP_VERSION="1.1.0", APP_BUILD="8d520fe, 2026-10-25")
    def test_version_once_on_login_page(self):
        page = self.client.get("/accounts/login/")
        self.assertContains(page, "v1.1.0", count=1)
        self.assertContains(page, 'title="Build 8d520fe, 2026-10-25"')

    @override_settings(APP_VERSION="development", APP_BUILD="")
    def test_development_without_build(self):
        page = self.client.get("/accounts/login/")
        self.assertContains(page, "development version")
        self.assertNotContains(page, "title=\"Build")

    def test_committed_version_is_semantic(self):
        from django.conf import settings as django_settings

        version = (django_settings.BASE_DIR / "VERSION").read_text(encoding="utf-8").strip()
        self.assertRegex(version, r"^\d+\.\d+\.\d+$")
        changelog = (django_settings.BASE_DIR / "CHANGELOG.md").read_text(encoding="utf-8")
        self.assertIn(f"## {version} ", changelog)  # every release is described

    def test_version_follows_the_commit_messages(self):
        import importlib.util
        from django.conf import settings as django_settings

        spec = importlib.util.spec_from_file_location("version_script", django_settings.BASE_DIR / "scripts" / "version.py")
        script = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(script)
        bump = script.next_version

        self.assertEqual(bump("1.1.1", []), "1.1.1")
        self.assertEqual(bump("1.1.1", ["docs(deploy): notes", "chore: tidy", "test: more"]), "1.1.1")
        self.assertEqual(bump("1.1.1", ["fix(leave): crash"]), "1.1.2")
        self.assertEqual(bump("1.1.1", ["Tidy the menu"]), "1.1.2")  # no type: assume users see a change
        self.assertEqual(bump("1.1.1", ["fix: a", "feat(accounts): b", "fix: c"]), "1.2.0")  # one bump per release
        self.assertEqual(bump("1.1.1", ["feat(api)!: new login"]), "2.0.0")
        self.assertEqual(bump("1.9.3", ["fix: x\n\nBREAKING CHANGE: the PDS is stored differently"]), "2.0.0")

    def test_deploy_releases_the_unreleased_changelog(self):
        import importlib.util
        from django.conf import settings as django_settings

        spec = importlib.util.spec_from_file_location("deploy_script", django_settings.BASE_DIR / "scripts" / "deploy.py")
        script = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(script)

        committed = "# Changelog\n\n## Unreleased (next minor)\n\nAdded:\n- Org charts zoom.\n\nFixed:\n- Paging.\n\n## 1.2.0 (October 26, 2026)\n\n- Old.\n"
        released, notes = script.release_changelog(committed, "1.3.0", "October 27, 2026")
        self.assertEqual(notes, "Added:\n- Org charts zoom.\n\nFixed:\n- Paging.")
        self.assertIn("## Unreleased\n\n## 1.3.0 (October 27, 2026)\n\nAdded:\n- Org charts zoom.", released)
        self.assertTrue(released.endswith("- Paging.\n\n## 1.2.0 (October 26, 2026)\n\n- Old.\n"))
        # No entries: the commit subjects are the notes.
        empty = "## Unreleased\n\n## 1.2.0 (x)\n"
        self.assertEqual(script.release_changelog(empty, "1.2.1", "y", ["- fix: a"])[1], "- fix: a")

        # An entry not committed yet stays under Unreleased, under its label; the release keeps the rest.
        working = committed.replace("- Paging.\n", "- Paging.\n- Two-step login clock message.\n")
        after = script.carry_uncommitted(committed, working, released)
        self.assertIn("## Unreleased\n\nFixed:\n- Two-step login clock message.\n\n## 1.3.0", after)
        self.assertEqual(after.count("Two-step login"), 1)
        self.assertEqual(script.carry_uncommitted(committed, committed, released), released)
        with self.assertRaises(script.Stop):  # an uncommitted edit to a released line
            script.carry_uncommitted(committed, committed.replace("- Old.", "- Older."), released)


class HrmoOwnRecordTestCase(TestCase):
    """R10: HRMO staff never change their own HR record; another HRMO staff member does."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp
    _complete = PersonalDataSheetTestCase._complete

    def setUp(self):
        self.setUp_roles()
        self.eng_a.leave_eligible = True
        self.eng_a.save()
        # The HRMO user has their own record: a plantilla appointment in Office A.
        self.person_h = Person.objects.create(employee_no="SP-00303", last_name="Mabini", first_name="Apolinario",
                                              birthdate=date(1985, 7, 23))
        self.eng_h = Engagement.objects.create(
            person=self.person_h, office=self.office_a, type=Engagement.Type.PLANTILLA,
            position_title="Human Resource Management Officer I", salary_grade=11, step=1, item_number="SP-MO-099",
            rate=Decimal("27000.00"), rate_basis=Engagement.RateBasis.MONTHLY, leave_eligible=True,
            start_date=date(2024, 1, 2), status=Engagement.Status.ACTIVE,
        )
        UserProfile.objects.create(user=self.hrmo_user, person=self.person_h)
        self.hrmo2 = get_user_model().objects.create_user("hrmo2", "hrmo2@test.com", "hrmo2pass123")
        self.hrmo2.groups.add(Group.objects.get(name=ROLE_HRMO))
        self.client.force_login(self.hrmo_user)

    def test_own_person_record_appointments_and_documents_are_read_only(self):
        doc = Document.objects.create(person=self.person_h, doc_type=Document.Type.APPOINTMENT)
        for url in (f"/people/{self.person_h.pk}/edit/", f"/people/{self.person_h.pk}/engagements/new/",
                    f"/engagements/{self.eng_h.pk}/edit/", f"/engagements/{self.eng_h.pk}/end/",
                    f"/engagements/{self.eng_h.pk}/renew/", f"/people/{self.person_h.pk}/documents/new/",
                    f"/engagements/{self.eng_h.pk}/documents/new/", f"/documents/{doc.pk}/edit/",
                    f"/documents/{doc.pk}/replace/"):
            self.assertEqual(self.client.get(url).status_code, 403, url)
        resp = self.client.post(f"/engagements/{self.eng_h.pk}/edit/", {"step": 8, "version": self.eng_h.version})
        self.assertEqual(resp.status_code, 403)
        self.eng_h.refresh_from_db()
        self.assertEqual(self.eng_h.step, 1)
        # No edit links on one's own record; they are still there on someone else's.
        self.assertNotContains(self.client.get(f"/people/{self.person_h.pk}/"), f"/people/{self.person_h.pk}/edit/")
        self.assertContains(self.client.get(f"/people/{self.person_a.pk}/"), f"/people/{self.person_a.pk}/edit/")
        self.assertEqual(self.client.get(f"/engagements/{self.eng_a.pk}/edit/").status_code, 200)
        self.client.force_login(self.hrmo2)
        self.assertEqual(self.client.get(f"/engagements/{self.eng_h.pk}/edit/").status_code, 200)

    def test_no_leave_credits_on_own_card(self):
        credit = {"leave_type": LeaveType.VACATION, "days": "30", "as_of": "2026-01-01", "note": "Beginning balance"}
        self.assertEqual(self.client.post(f"/leave/card/{self.person_h.pk}/", credit).status_code, 403)
        self.assertFalse(LeaveCredit.objects.filter(person=self.person_h).exists())
        self.client.post(f"/leave/card/{self.person_a.pk}/", credit)
        self.assertTrue(LeaveCredit.objects.filter(person=self.person_a, days=Decimal("30")).exists())
        self.client.force_login(self.hrmo2)
        self.client.post(f"/leave/card/{self.person_h.pk}/", credit)
        self.assertTrue(LeaveCredit.objects.filter(person=self.person_h).exists())

    def test_no_overtime_credit_for_oneself(self):
        from hr.models import CocEntry

        self.assertNotContains(self.client.get("/leave/coc/record/"), f'value="{self.person_h.pk}"')
        self.client.post("/leave/coc/record/", {
            "worked_on": (date.today() - timedelta(days=1)).isoformat(), "hours": "4",
            "day_kind": CocEntry.DayKind.WORKDAY, "reference": "OO 12", "person": [self.person_a.pk, self.person_h.pk],
        })
        self.assertTrue(CocEntry.objects.filter(person=self.person_a).exists())
        self.assertFalse(CocEntry.objects.filter(person=self.person_h).exists())

    def test_no_training_rating_or_saln_receipt_for_oneself(self):
        from hr import spms
        from hr.models import PdsTraining, PerformanceRating, SalnFiling

        self.client.post("/tracker/learning/record/", {"title": "Records management", "date_from": "2026-03-02",
                                                       "hours": "8", "person": [self.person_a.pk, self.person_h.pk]})
        self.assertTrue(PdsTraining.objects.filter(person=self.person_a).exists())
        self.assertFalse(PdsTraining.objects.filter(person=self.person_h).exists())

        year, semester = spms.last_ended(date.today())
        self.client.post(f"/tracker/ratings/{self.office_a.pk}/?year={year}&semester={semester}",
                         {f"rating-{self.person_a.pk}": "4.2", f"rating-{self.person_h.pk}": "5.0"})
        self.assertTrue(PerformanceRating.objects.filter(person=self.person_a).exists())
        self.assertFalse(PerformanceRating.objects.filter(person=self.person_h).exists())

        self.client.post("/tracker/saln/", {"person": [self.person_a.pk, self.person_h.pk], "received_on": "2026-04-15"})
        self.assertTrue(SalnFiling.objects.filter(person=self.person_a, received_on__isnull=False).exists())
        self.assertFalse(SalnFiling.objects.filter(person=self.person_h, received_on__isnull=False).exists())

    def test_no_work_on_own_salary_action(self):
        case = SalaryAction.objects.create(kind=SalaryAction.Kind.STEP_INCREMENT, person=self.person_h,
                                           engagement=self.eng_h, milestone="Step 2", service_from=date(2024, 1, 2),
                                           due_date=date(2027, 1, 2), from_step=1, to_step=2,
                                           status=SalaryAction.Status.DUE)
        resp = self.client.post(f"/salary-actions/{case.pk}/", {"action": "start_review", "version": case.version})
        self.assertEqual(resp.status_code, 403)
        resp = self.client.post(f"/salary-actions/{case.pk}/interruptions/", {
            "kind": ServiceInterruption.Kind.LWOP, "start_date": "2025-03-02", "end_date": "2025-03-03"})
        self.assertEqual(resp.status_code, 403)
        case.refresh_from_db()
        self.assertEqual(case.status, SalaryAction.Status.DUE)
        self.assertFalse(ServiceInterruption.objects.filter(person=self.person_h).exists())
        self.assertEqual(self.client.get(f"/salary-actions/{case.pk}/").status_code, 200)  # still visible

    def test_no_cancelling_own_approved_leave(self):
        from hr.leave import LeaveError, cancel, file_request

        req = file_request(LeaveRequest(person=self.person_h, leave_type="vl", detail="within_ph",
                                        start_date=date(2026, 11, 2), end_date=date(2026, 11, 2), days=1), None)
        req.status = LeaveRequest.Status.APPROVED
        req.save()
        with self.assertRaises(LeaveError):
            cancel(req, self.hrmo_user, note="Changed my mind")
        cancel(req, self.hrmo2, note="Recalled to duty")
        self.assertEqual(req.status, LeaveRequest.Status.CANCELLED)

    def test_no_verifying_own_pds(self):
        from hr.models import PdsSubmission

        self._complete(self.person_h)
        self.client.post(f"/people/{self.person_h.pk}/pds/submit/", {"verify": "1"})
        sub = PdsSubmission.objects.get(person=self.person_h)
        self.assertEqual(sub.status, "submitted")
        resp = self.client.post(f"/pds/review/{sub.pk}/", {"action": "verify", "version": sub.version})
        self.assertEqual(resp.status_code, 403)
        sub.refresh_from_db()
        self.assertEqual(sub.status, "submitted")
        self.client.force_login(self.hrmo2)
        self.client.post(f"/pds/review/{sub.pk}/", {"action": "verify", "version": sub.version})
        sub.refresh_from_db()
        self.assertEqual((sub.status, sub.reviewed_by), ("verified", self.hrmo2))

    def test_no_setting_own_supervisor(self):
        url = f"/offices/{self.office_a.pk}/chart/edit/"
        self.assertNotContains(self.client.get(url), f'name="reports_to_{self.eng_h.pk}"')
        self.client.post(url, {
            f"reports_to_{self.eng_h.pk}": str(self.eng_a.pk), f"version_{self.eng_h.pk}": str(self.eng_h.version),
            f"reports_to_{self.eng_a.pk}": str(self.eng_h.pk), f"version_{self.eng_a.pk}": str(self.eng_a.version),
        })
        self.eng_h.refresh_from_db()
        self.eng_a.refresh_from_db()
        self.assertIsNone(self.eng_h.reports_to)
        self.assertEqual(self.eng_a.reports_to, self.eng_h)  # other rows on the same sheet still save
        self.client.force_login(self.hrmo2)
        self.assertContains(self.client.get(url), f'name="reports_to_{self.eng_h.pk}"')
        self.client.post(url, {
            f"reports_to_{self.eng_h.pk}": str(self.eng_a.pk), f"version_{self.eng_h.pk}": str(self.eng_h.version),
            f"reports_to_{self.eng_a.pk}": "", f"version_{self.eng_a.pk}": str(self.eng_a.version),
        })
        self.eng_h.refresh_from_db()
        self.assertEqual(self.eng_h.reports_to, self.eng_a)


class AuditGroupedViewTestCase(TestCase):
    """/audit/?view=grouped groups the filtered records by who acted; the flat list stays the default."""

    def setUp(self):
        call_command("seed_offices")
        call_command("setup_roles")
        User = get_user_model()
        self.auditor = User.objects.create_user("aud_grp", password="Password123!")
        self.auditor.groups.add(Group.objects.get(name=ROLE_AUDITOR))
        self.hrmo = User.objects.create_user("hrmo_grp", password="Password123!")
        self.hrmo.groups.add(Group.objects.get(name=ROLE_HRMO))
        self.office = Office.objects.get(name="Office of the Municipal Mayor")

    def _log(self, username, n, action=AuditLog.Action.VIEW):
        return log_audit_event(action=action, model_name="Person", object_repr=f"GRP-{n}",
                               username=username, office=self.office if username else None)

    def _interleave(self, count):
        # Two people working at once, plus an unattributed system event.
        for n in range(count):
            self._log("zeta", n)
            self._log("alpha", n)
        self._log(None, "sys")

    def test_flat_mode_is_default_and_chronological(self):
        self._interleave(3)
        self.client.force_login(self.auditor)
        resp = self.client.get("/audit/?q=GRP-")
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(resp.context["grouped"])
        self.assertEqual(resp.context["groups"], [])
        rows = list(resp.context["page_obj"])
        self.assertEqual(len(rows), 7)
        self.assertEqual([r.id for r in rows], sorted((r.id for r in rows), reverse=True))
        self.assertContains(resp, "?view=grouped&q=GRP-")

    def test_grouped_mode_groups_by_user_with_system_last(self):
        self._interleave(3)
        self.client.force_login(self.auditor)
        resp = self.client.get("/audit/?view=grouped&q=GRP-")
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.context["grouped"])
        groups = resp.context["groups"]
        self.assertEqual([g["username"] for g in groups], ["alpha", "zeta", ""])
        self.assertEqual([g["total"] for g in groups], [3, 3, 1])
        alpha = groups[0]
        self.assertEqual([r.object_repr for r in alpha["rows"]], ["GRP-2", "GRP-1", "GRP-0"])
        self.assertEqual(alpha["office_name"], self.office.name)
        self.assertLessEqual(alpha["earliest"], alpha["latest"])
        self.assertContains(resp, "System / Anonymous")
        self.assertContains(resp, 'data-bs-toggle="collapse"')
        self.assertContains(resp, 'id="audit-group-3"')
        self.assertContains(resp, '<input type="hidden" name="view" value="grouped">', html=True)

    def test_grouped_mode_keeps_filters(self):
        self._interleave(2)
        self._log("alpha", "upd", action=AuditLog.Action.UPDATE)
        self.client.force_login(self.auditor)
        resp = self.client.get("/audit/?view=grouped&q=GRP-&action=update")
        groups = resp.context["groups"]
        self.assertEqual([(g["username"], g["total"]) for g in groups], [("alpha", 1)])
        resp = self.client.get("/audit/?view=grouped&q=GRP-&user=zet")
        self.assertEqual([(g["username"], g["total"]) for g in resp.context["groups"]], [("zeta", 2)])

    def test_grouped_pagination_keeps_mode_and_filters_and_totals(self):
        self._interleave(30)  # 61 rows: alpha 30, zeta 30, system 1
        self.client.force_login(self.auditor)
        url = "/audit/?view=grouped&q=GRP-&start_date=2000-01-01"
        page1 = self.client.get(url)
        groups = page1.context["groups"]
        self.assertEqual([(g["username"], len(g["rows"]), g["total"]) for g in groups],
                         [("alpha", 30, 30), ("zeta", 20, 30)])
        self.assertContains(page1, "20 on this page")
        query = page1.context["page_query"]
        for part in ("view=grouped", "q=GRP-", "start_date=2000-01-01"):
            self.assertIn(part, query)
        self.assertNotIn("page=", query)

        page2 = self.client.get(f"/audit/?page=2&{query}")
        groups = page2.context["groups"]
        self.assertEqual([(g["username"], len(g["rows"]), g["total"]) for g in groups],
                         [("zeta", 10, 30), ("", 1, 1)])

        # The toggle back to the flat list keeps the filters but starts at page 1.
        page2_toggle = page2.context["filter_query"]
        self.assertNotIn("view=", page2_toggle)
        self.assertNotIn("page=", page2_toggle)
        flat = self.client.get("/audit/?page=2&q=GRP-&start_date=2000-01-01")
        self.assertFalse(flat.context["grouped"])
        self.assertEqual(len(list(flat.context["page_obj"])), 11)

    def test_grouped_view_is_auditor_only(self):
        resp = self.client.get("/audit/?view=grouped")
        self.assertEqual(resp.status_code, 302)
        self.client.force_login(self.hrmo)
        self.assertEqual(self.client.get("/audit/?view=grouped").status_code, 403)
        self.assertEqual(self.client.get("/audit/export/csv/?view=grouped").status_code, 403)

    def test_csv_export_stays_chronological(self):
        self._interleave(3)
        self.client.force_login(self.auditor)
        import csv

        content = self.client.get("/audit/export/csv/?view=grouped").content.decode()
        reprs = [row[7] for row in csv.reader(content.splitlines()[1:]) if row[7].startswith("GRP-")]
        self.assertEqual(reprs, ["GRP-sys", "GRP-2", "GRP-2", "GRP-1", "GRP-1", "GRP-0", "GRP-0"])


class ForgotPasswordTestCase(TestCase):
    """A2: "Forgot password?" asks HRMO, who checks an ID and issues a one-time password."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()
        # The fixture's employee account for person A, renamed to the employee number (D9).
        self.employee = self.emp_a_user
        self.employee.username = "SP-00101"
        self.employee.set_password("Forgotten-Pass-2026")
        self.employee.save()

    def _ask(self, username, contact="", ip="203.0.113.7"):
        return self.client.post("/accounts/forgot/", {"username": username, "contact": contact}, REMOTE_ADDR=ip)

    def _issue_from_request(self, req, **data):
        payload = {"action": "issue", "id_checked": "LGU ID", "confirm": "on"}
        payload.update(data)
        return self.client.post(f"/accounts/reset-requests/{req.pk}/", payload)

    def test_login_page_links_to_forgot_password(self):
        self.assertContains(self.client.get("/accounts/login/"), 'href="/accounts/forgot/"')
        self.assertEqual(self.client.get("/accounts/forgot/").status_code, 200)

    def test_known_and_unknown_usernames_get_the_same_answer(self):
        from hr.models import PasswordResetRequest

        known = self._ask("sp-00101", "0917 000 0000")  # case does not matter
        unknown = self._ask("SP-99999")
        self.assertRedirects(known, "/accounts/forgot/sent/")
        self.assertRedirects(unknown, "/accounts/forgot/sent/")
        stored = {r.username_entered: r for r in PasswordResetRequest.objects.all()}
        self.assertEqual(stored["sp-00101"].user, self.employee)
        self.assertEqual(stored["sp-00101"].contact, "0917 000 0000")
        self.assertIsNone(stored["SP-99999"].user)
        self.assertEqual(stored["sp-00101"].status, PasswordResetRequest.Status.PENDING)
        self.assertTrue(self.employee.check_password("Forgotten-Pass-2026"))  # nothing changes until HRMO acts

    def test_duplicates_and_floods_are_not_stored(self):
        from hr.models import PasswordResetRequest

        self._ask("SP-00101")
        self._ask("SP-00101")
        self.assertEqual(PasswordResetRequest.objects.filter(user=self.employee).count(), 1)
        for i in range(10):
            self.assertRedirects(self._ask(f"nobody-{i}", ip="198.51.100.9"), "/accounts/forgot/sent/")
        self.assertEqual(PasswordResetRequest.objects.filter(ip_address="198.51.100.9").count(), 5)

    def test_logged_in_user_is_sent_to_change_password(self):
        self.client.force_login(self.employee)
        self.assertRedirects(self.client.get("/accounts/forgot/"), "/accounts/password/")

    def test_only_hrmo_sees_the_queue_and_issues(self):
        from hr.models import PasswordResetRequest

        self._ask("SP-00101")
        req = PasswordResetRequest.objects.get()
        urls = ["/accounts/reset-requests/", f"/accounts/reset-requests/{req.pk}/", f"/people/{self.person_a.pk}/password/"]
        self.assertRedirects(self.client.get(urls[0]), "/accounts/login/?next=/accounts/reset-requests/")
        for user in (self.head_a, self.staff_a, self.auditor_user, self.mayor, self.sysadmin, self.employee):
            self.client.force_login(user)
            for url in urls:
                self.assertEqual(self.client.get(url).status_code, 403, (user, url))
            self.assertEqual(self._issue_from_request(req).status_code, 403, user)
        self.assertTrue(get_user_model().objects.get(pk=self.employee.pk).check_password("Forgotten-Pass-2026"))
        self.client.force_login(self.hrmo_user)
        page = self.client.get(urls[0])
        self.assertContains(page, "SP-00101")
        self.assertContains(page, self.person_a.full_name)
        self.assertContains(self.client.get("/"), 'href="/accounts/reset-requests/"')

    def test_hrmo_issues_a_one_time_password_after_checking_an_id(self):
        from django.utils import timezone
        from hr.models import PasswordResetRequest

        self._ask("SP-00101")
        req = PasswordResetRequest.objects.get()
        self.client.force_login(self.hrmo_user)
        # Both the ID and the tick are required.
        self.assertEqual(self._issue_from_request(req, id_checked="").status_code, 200)
        self.assertEqual(self._issue_from_request(req, confirm="").status_code, 200)
        req.refresh_from_db()
        self.assertEqual(req.status, PasswordResetRequest.Status.PENDING)

        resp = self._issue_from_request(req, id_checked="UMID")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Cache-Control"], "no-store")
        password = resp.context["password"]
        self.assertContains(resp, password)
        req.refresh_from_db()
        self.assertEqual((req.status, req.handled_by, req.id_checked),
                         (PasswordResetRequest.Status.ISSUED, self.hrmo_user, "UMID"))
        profile = UserProfile.objects.get(user=self.employee)
        self.assertTrue(profile.must_change_password)
        self.assertLess(profile.password_expires_at - timezone.now(), timedelta(hours=25))

        # A second POST (refresh) does not issue another password.
        self._issue_from_request(req)
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password(password))
        self.assertFalse(self.employee.check_password("Forgotten-Pass-2026"))

        resets = AuditLog.objects.filter(action=AuditLog.Action.PASSWORD_RESET, object_id=str(self.employee.pk),
                                         username="hrmo")
        self.assertEqual(resets.count(), 1)
        for row in AuditLog.objects.all():
            self.assertNotIn(password, f"{row.object_repr} {row.changes}")
        self.assertTrue(verify_chain().is_valid)

        # The person logs in with it and must set their own password first.
        self.client.logout()
        resp = self.client.post("/accounts/login/", {"username": "SP-00101", "password": password})
        self.assertEqual(resp.status_code, 302)
        self.assertRedirects(self.client.get(f"/people/{self.person_a.pk}/"), "/accounts/password/")

    def test_issuing_lifts_a_login_lockout(self):
        from axes.models import AccessAttempt
        from hr.models import PasswordResetRequest

        for _ in range(5):
            self.client.post("/accounts/login/", {"username": "SP-00101", "password": "wrong-guess"})
        self.assertTrue(AccessAttempt.objects.filter(username="SP-00101").exists())
        self._ask("SP-00101")
        self.client.force_login(self.hrmo_user)
        password = self._issue_from_request(PasswordResetRequest.objects.get()).context["password"]
        self.assertFalse(AccessAttempt.objects.filter(username="SP-00101").exists())
        self.client.logout()
        self.client.post("/accounts/login/", {"username": "SP-00101", "password": password})
        self.assertEqual(int(self.client.session["_auth_user_id"]), self.employee.pk)

    def test_accounts_with_more_access_than_hrmo_stay_with_the_superuser(self):
        from hr.models import PasswordResetRequest

        # An Office Head with their own personnel record, a delegate, and HRMO's own account.
        profile = UserProfile.objects.get(user=self.head_a)
        profile.person = self.person_b
        profile.save()
        UserProfile.objects.create(user=self.delegate_user, person=Person.objects.create(
            employee_no="SP-00909", last_name="Silang", first_name="Gabriela", birthdate=date(1981, 3, 19)))
        Delegate.objects.create(principal=self.mayor, delegate=self.delegate_user,
                                valid_from=date.today(), valid_to=date.today() + timedelta(days=7))
        UserProfile.objects.create(user=self.hrmo_user, person=Person.objects.create(
            employee_no="SP-00910", last_name="Mabini", first_name="Apolinario", birthdate=date(1985, 7, 23)))
        for user in (self.head_a, self.delegate_user, self.hrmo_user):
            old_hash = user.password
            self.client.logout()
            self._ask(user.username)
            self.client.force_login(self.hrmo_user)
            req = PasswordResetRequest.objects.get(user=user)
            self.assertNotContains(self.client.get(f"/accounts/reset-requests/{req.pk}/"), 'value="issue"')
            self.assertEqual(self._issue_from_request(req).status_code, 200)  # page again, nothing issued
            user.refresh_from_db()
            self.assertEqual(user.password, old_hash, user)
            req.refresh_from_db()
            self.assertEqual(req.status, PasswordResetRequest.Status.PENDING)

    def test_walk_in_from_the_person_record_closes_the_online_request(self):
        from hr.models import PasswordResetRequest

        self._ask("SP-00101")
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.get(f"/people/{self.person_a.pk}/"), f"/people/{self.person_a.pk}/password/")
        self.assertNotContains(self.client.get(f"/people/{self.person_b.pk}/"), f"/people/{self.person_b.pk}/password/")
        resp = self.client.post(f"/people/{self.person_a.pk}/password/",
                                {"action": "issue", "id_checked": "LGU ID", "confirm": "on"})
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password(resp.context["password"]))
        statuses = sorted(PasswordResetRequest.objects.values_list("received_via", "status"))
        self.assertEqual(statuses, [("counter", "issued"), ("online", "issued")])

    def test_dismiss_needs_a_reason(self):
        from hr.models import PasswordResetRequest

        self._ask("SP-99999")
        req = PasswordResetRequest.objects.get()
        self.client.force_login(self.hrmo_user)
        url = f"/accounts/reset-requests/{req.pk}/"
        self.assertEqual(self.client.post(url, {"action": "dismiss", "note": ""}).status_code, 200)
        self.assertRedirects(self.client.post(url, {"action": "dismiss", "note": "No such employee"}),
                             "/accounts/reset-requests/")
        req.refresh_from_db()
        self.assertEqual((req.status, req.note, req.handled_by), ("dismissed", "No such employee", self.hrmo_user))


@override_settings(MFA_REQUIRED=True)
class MfaTestCase(TestCase):
    """A5: second login step with an authenticator app (TOTP) and one-use recovery codes."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()
        from hr import mfa

        self.mfa = mfa

    def _enroll(self, user):
        from hr.models import MfaDevice

        secret = self.mfa.new_secret()
        MfaDevice.objects.create(user=user, secret=secret)
        return secret

    def _code(self, secret, offset=0):
        return self.mfa.code_at(secret, self.mfa.current_step() + offset)

    def test_codes_match_rfc_6238(self):
        # RFC 6238 appendix B, SHA-1 key "12345678901234567890"; the RFC shows 8 digits, we use the last 6.
        secret = "GEZDGNBVGY3TQOJQGEZDGNBVGY3TQOJQ"
        self.assertEqual(self.mfa.code_at(secret, 59 // 30), "287082")
        self.assertEqual(self.mfa.code_at(secret, 1111111109 // 30), "081804")
        self.assertEqual(self.mfa.matching_step(secret, "081 804", now=1111111109), 1111111109 // 30)
        self.assertIsNone(self.mfa.matching_step(secret, "081804", now=1111111109 + 120))
        # Pasted from a phone, a code can carry a non-breaking or thin space instead of a plain one.
        self.assertEqual(self.mfa.matching_step(secret, "081\u00a0804", now=1111111109), 1111111109 // 30)
        self.assertEqual(self.mfa.matching_step(secret, " 081\u2009804\n", now=1111111109), 1111111109 // 30)

    def test_required_roles_must_set_up_before_anything_else(self):
        for user in (self.hrmo_user, self.auditor_user, self.sysadmin, self.superuser):
            self.client.force_login(user)
            self.assertRedirects(self.client.get("/people/"), "/accounts/two-step/setup/", fetch_redirect_response=False)
            self.assertRedirects(self.client.get("/admin/"), "/accounts/two-step/setup/", fetch_redirect_response=False)
        # Plain employees and Office Heads are not forced; they may opt in.
        for user in (self.emp_a_user, self.head_a):
            self.client.force_login(user)
            self.assertEqual(self.client.get("/").status_code, 200)

    def test_menus_are_hidden_before_the_second_step(self):
        self.client.force_login(self.hrmo_user)
        page = self.client.get("/accounts/two-step/setup/")
        self.assertEqual(page.status_code, 200)
        self.assertNotContains(page, "Forgotten passwords")
        self.assertNotContains(page, "Personnel")

    def test_setup_stores_the_app_and_shows_recovery_codes_once(self):
        from hr.models import AuditLog, MfaDevice, MfaRecoveryCode

        self.client.force_login(self.hrmo_user)
        page = self.client.get("/accounts/two-step/setup/")
        self.assertContains(page, "data:image/svg+xml")
        self.assertEqual(page["Cache-Control"], "no-store")
        secret = self.client.session[self.mfa.SESSION_PENDING_SECRET]
        self.assertFalse(MfaDevice.objects.filter(user=self.hrmo_user).exists())  # nothing stored before a code

        self.assertContains(self.client.post("/accounts/two-step/setup/", {"code": "000000"}), "not right")
        done = self.client.post("/accounts/two-step/setup/", {"code": self._code(secret)})
        self.assertContains(done, "only time they are shown")
        self.assertEqual(done["Cache-Control"], "no-store")
        self.assertEqual(MfaDevice.objects.get(user=self.hrmo_user).secret, secret)
        self.assertEqual(MfaRecoveryCode.objects.filter(user=self.hrmo_user).count(), 10)
        self.assertEqual(self.client.get("/people/").status_code, 200)

        self.assertTrue(AuditLog.objects.filter(action=AuditLog.Action.MFA_SETUP, user_id=self.hrmo_user.pk).exists())
        self.assertFalse(any(secret in str(e.changes) for e in AuditLog.objects.all()))

    def test_login_with_an_app_asks_for_a_code_and_keeps_the_page(self):
        secret = self._enroll(self.hrmo_user)
        response = self.client.post("/accounts/login/", {"username": "hrmo", "password": "hrmopass123"})
        self.assertEqual(response.status_code, 302)
        self.assertRedirects(self.client.get("/people/"), "/accounts/two-step/code/?next=%2Fpeople%2F",
                             fetch_redirect_response=False)
        wrong = self.client.post("/accounts/two-step/code/", {"code": "000000", "next": "/people/"})
        self.assertContains(wrong, "not right")
        ok = self.client.post("/accounts/two-step/code/", {"code": self._code(secret), "next": "/people/"})
        self.assertRedirects(ok, "/people/", fetch_redirect_response=False)
        self.assertEqual(self.client.get("/people/").status_code, 200)

    def test_an_app_is_needed_even_when_not_required(self):
        with override_settings(MFA_REQUIRED=False):
            self.client.force_login(self.hrmo_user)
            self.assertEqual(self.client.get("/people/").status_code, 200)
            self._enroll(self.emp_a_user)
            self.client.force_login(self.emp_a_user)
            self.assertEqual(self.client.get("/").status_code, 302)

    def test_a_code_works_only_once(self):
        secret = self._enroll(self.hrmo_user)
        code = self._code(secret)
        self.client.force_login(self.hrmo_user)
        self.assertEqual(self.client.post("/accounts/two-step/code/", {"code": code}).status_code, 302)
        other = self.client_class()
        other.force_login(self.hrmo_user)
        self.assertContains(other.post("/accounts/two-step/code/", {"code": code}), "just used")
        self.assertEqual(other.get("/people/").status_code, 302)

    def test_full_flow_set_up_log_out_then_a_code_then_a_recovery_code(self):
        """The first thing anyone does after setup: log out and straight back in with the code still showing."""
        from hr.models import MfaDevice

        login = {"username": "hrmo", "password": "hrmopass123"}
        self.client.post("/accounts/login/", login)
        self.assertRedirects(self.client.get("/people/"), "/accounts/two-step/setup/", fetch_redirect_response=False)
        self.client.get("/accounts/two-step/setup/")
        secret = self.client.session[self.mfa.SESSION_PENDING_SECRET]
        setup_code = self._code(secret)
        codes = self.client.post("/accounts/two-step/setup/", {"code": setup_code}).context["codes"]
        self.assertEqual(self.client.get("/people/").status_code, 200)

        self.client.post("/accounts/logout/")
        self.client.post("/accounts/login/", login)
        # The app still shows the setup code: refused, but said plainly and not counted as a wrong guess.
        for _ in range(self.mfa.MAX_FAILURES):
            again = self.client.post("/accounts/two-step/code/", {"code": setup_code, "next": "/people/"})
            self.assertContains(again, "just used")
            self.assertNotContains(again, "not right")
        self.assertEqual(MfaDevice.objects.get(user=self.hrmo_user).failures, 0)
        ok = self.client.post("/accounts/two-step/code/", {"code": self._code(secret, 1), "next": "/people/"})
        self.assertRedirects(ok, "/people/", fetch_redirect_response=False)
        self.assertEqual(self.client.get("/people/").status_code, 200)

        self.client.post("/accounts/logout/")
        self.client.post("/accounts/login/", login)
        self.assertEqual(self.client.get("/people/").status_code, 302)
        ok = self.client.post("/accounts/two-step/code/", {"code": codes[0], "next": "/people/"})
        self.assertRedirects(ok, "/people/", fetch_redirect_response=False)
        self.assertEqual(self.mfa.recovery_codes_left(self.hrmo_user), 9)

    def test_a_code_from_a_clock_that_is_off_says_so(self):
        """A server clock 93 s slow (Windows Time stopped) refused every code from the phone as "not right"."""
        from hr.models import AuditLog, MfaDevice

        self.client.force_login(self.hrmo_user)
        self.client.get("/accounts/two-step/setup/")
        secret = self.client.session[self.mfa.SESSION_PENDING_SECRET]
        page = self.client.post("/accounts/two-step/setup/", {"code": self._code(secret, 3)})
        self.assertContains(page, "for a time 90 seconds ahead of this server")
        self.assertNotContains(page, "not right")
        self.assertContains(self.client.post("/accounts/two-step/setup/", {"code": self._code(secret, -20)}),
                            "10 minutes behind this server")
        # Further off than 10 minutes it is just a wrong code.
        self.assertContains(self.client.post("/accounts/two-step/setup/", {"code": self._code(secret, 25)}), "not right")
        self.assertFalse(MfaDevice.objects.filter(user=self.hrmo_user).exists())
        self.assertEqual(self.mfa.clock_offset(secret, self.mfa.code_at(secret, 1003), now=1000 * 30), 90)
        self.assertIsNone(self.mfa.clock_offset(secret, self.mfa.code_at(secret, 1001), now=1000 * 30))

        # At login too: refused, said plainly, and not counted towards the lock.
        self.client.post("/accounts/two-step/setup/", {"code": self._code(secret)})
        other = self.client_class()
        other.force_login(self.hrmo_user)
        for _ in range(self.mfa.MAX_FAILURES):
            self.assertContains(other.post("/accounts/two-step/code/", {"code": self._code(secret, 3)}), "ahead of")
        self.assertEqual(MfaDevice.objects.get(user=self.hrmo_user).failures, 0)
        self.assertTrue(AuditLog.objects.filter(
            action=AuditLog.Action.MFA_FAILED, changes__reason="clocks differ by 90 s").exists())

    def test_the_login_code_typed_again_on_two_step_login_says_so(self):
        from hr.models import AuditLog, MfaDevice

        secret = self._enroll(self.hrmo_user)
        self.mfa.new_recovery_codes(self.hrmo_user)
        old = set(self.hrmo_user.mfa_recovery_codes.values_list("code_hash", flat=True))
        self.client.force_login(self.hrmo_user)
        code = self._code(secret)
        self.client.post("/accounts/two-step/code/", {"code": code})
        page = self.client.post("/accounts/two-step/", {"action": "codes", "code": code})
        self.assertContains(page, "just used")
        self.assertEqual(set(self.hrmo_user.mfa_recovery_codes.values_list("code_hash", flat=True)), old)
        self.assertEqual(MfaDevice.objects.get(user=self.hrmo_user).failures, 0)
        self.assertTrue(AuditLog.objects.filter(
            action=AuditLog.Action.MFA_FAILED, changes__reason="code already used").exists())
        # A wrong code still counts.
        self.assertContains(self.client.post("/accounts/two-step/", {"action": "codes", "code": "000000"}), "not right")
        self.assertEqual(MfaDevice.objects.get(user=self.hrmo_user).failures, 1)

    def test_logging_in_again_needs_the_second_step_again(self):
        secret = self._enroll(self.hrmo_user)
        self.client.force_login(self.hrmo_user)
        self.client.post("/accounts/two-step/code/", {"code": self._code(secret)})
        self.assertEqual(self.client.get("/people/").status_code, 200)
        self.client.force_login(self.hrmo_user)
        self.assertEqual(self.client.get("/people/").status_code, 302)

    def test_five_wrong_codes_lock_the_account_not_just_the_session(self):
        from hr.models import AuditLog

        secret = self._enroll(self.hrmo_user)
        self.client.force_login(self.hrmo_user)
        for _ in range(4):
            self.assertEqual(self.client.post("/accounts/two-step/code/", {"code": "000000"}).status_code, 200)
        self.assertRedirects(self.client.post("/accounts/two-step/code/", {"code": "000000"}), "/accounts/login/",
                             fetch_redirect_response=False)
        self.assertNotIn("_auth_user_id", self.client.session)
        # A fresh password login does not reset the count, and even the right code is refused.
        self.client.force_login(self.hrmo_user)
        self.assertRedirects(self.client.post("/accounts/two-step/code/", {"code": self._code(secret)}),
                             "/accounts/login/", fetch_redirect_response=False)
        self.assertEqual(AuditLog.objects.filter(action=AuditLog.Action.MFA_FAILED).count(), 5)

    def test_recovery_code_works_once(self):
        self._enroll(self.hrmo_user)
        codes = self.mfa.new_recovery_codes(self.hrmo_user)
        self.client.force_login(self.hrmo_user)
        self.assertEqual(self.client.post("/accounts/two-step/code/", {"code": codes[0].upper()}).status_code, 302)
        self.assertEqual(self.mfa.recovery_codes_left(self.hrmo_user), 9)
        self.client.force_login(self.hrmo_user)
        self.assertContains(self.client.post("/accounts/two-step/code/", {"code": codes[0]}), "not right")

    def test_next_never_leaves_the_site(self):
        secret = self._enroll(self.hrmo_user)
        self.client.force_login(self.hrmo_user)
        response = self.client.post("/accounts/two-step/code/", {"code": self._code(secret), "next": "https://evil.example/"})
        self.assertRedirects(response, "/", fetch_redirect_response=False)

    def test_one_time_password_then_setup_then_password_change(self):
        from hr.accounts import require_password_change

        require_password_change(self.hrmo_user)
        self.client.force_login(self.hrmo_user)
        self.assertRedirects(self.client.get("/people/"), "/accounts/two-step/setup/", fetch_redirect_response=False)
        self.assertEqual(self.client.get("/accounts/two-step/setup/").status_code, 200)  # no redirect loop
        secret = self.client.session[self.mfa.SESSION_PENDING_SECRET]
        self.client.post("/accounts/two-step/setup/", {"code": self._code(secret)})
        self.assertRedirects(self.client.get("/people/"), "/accounts/password/", fetch_redirect_response=False)

    def test_required_roles_cannot_turn_it_off_others_can(self):
        secret = self._enroll(self.hrmo_user)
        self.client.force_login(self.hrmo_user)
        self.client.post("/accounts/two-step/code/", {"code": self._code(secret, -1)})
        self.assertNotContains(self.client.get("/accounts/two-step/"), 'value="off"')
        self.client.post("/accounts/two-step/", {"action": "off", "code": self._code(secret)})
        self.assertTrue(self.mfa.has_device(self.hrmo_user))

        secret = self._enroll(self.emp_a_user)
        self.client.force_login(self.emp_a_user)
        self.client.post("/accounts/two-step/code/", {"code": self._code(secret, -1)})
        self.client.post("/accounts/two-step/", {"action": "off", "code": self._code(secret)})
        self.assertFalse(self.mfa.has_device(self.emp_a_user))

    def test_reset_command_and_admin_action(self):
        from django.contrib.admin.sites import site
        from hr.models import AuditLog

        self._enroll(self.hrmo_user)
        self.mfa.new_recovery_codes(self.hrmo_user)
        with self.assertRaises(CommandError):
            call_command("reset_mfa", "hrmo")  # a reason is required
        call_command("reset_mfa", "HRMO", reason="Phone lost, LGU ID checked")
        self.assertFalse(self.mfa.has_device(self.hrmo_user))
        self.assertEqual(self.mfa.recovery_codes_left(self.hrmo_user), 0)
        entry = AuditLog.objects.get(action=AuditLog.Action.MFA_RESET)
        self.assertEqual((entry.username, entry.changes["reason"]), ("manage.py reset_mfa", "Phone lost, LGU ID checked"))

        # In /admin/ only the emergency superuser gets the action: System Admin must not unlock HR accounts.
        user_admin = site._registry[get_user_model()]
        request = RequestFactory().get("/admin/auth/user/")
        request.user = self.sysadmin
        self.assertNotIn("reset_two_step_login", user_admin.get_actions(request))
        request.user = self.superuser
        self.assertIn("reset_two_step_login", user_admin.get_actions(request))


@override_settings(MFA_REQUIRED=True, MFA_REMEMBER_DAYS=30)
class MfaRememberBrowserTestCase(TestCase):
    """A5: "Remember this browser" skips the code until it expires, the password changes, or it is forgotten."""

    setUp_roles = Step4CSecurityAndRolesTestCase.setUp

    def setUp(self):
        self.setUp_roles()
        from hr import mfa
        from hr.models import MfaDevice

        self.mfa = mfa
        self.secret = mfa.new_secret()
        MfaDevice.objects.create(user=self.hrmo_user, secret=self.secret)
        self.cookie = mfa.remember_cookie_name(self.hrmo_user)

    def _pass(self, remember=True, offset=0):
        self.client.force_login(self.hrmo_user)
        data = {"code": self.mfa.code_at(self.secret, self.mfa.current_step() + offset)}
        if remember:
            data["remember"] = "1"
        return self.client.post("/accounts/two-step/code/", data)

    def _login_again(self):
        self.client.force_login(self.hrmo_user)
        return self.client.get("/people/").status_code

    def test_remembered_browser_skips_the_code_after_the_next_login(self):
        from hr.models import AuditLog

        response = self._pass()
        self.assertEqual(response.cookies[self.cookie]["max-age"], 30 * 86400)
        self.assertTrue(response.cookies[self.cookie]["httponly"])
        self.assertEqual(self._login_again(), 200)
        self.assertTrue(AuditLog.objects.filter(
            action=AuditLog.Action.MFA_PASSED, changes__method="remembered browser").exists())
        # Another browser still needs the code.
        other = self.client_class()
        other.force_login(self.hrmo_user)
        self.assertEqual(other.get("/people/").status_code, 302)

    def test_not_ticked_means_not_remembered(self):
        self.assertNotIn(self.cookie, self._pass(remember=False).cookies)
        self.assertEqual(self._login_again(), 302)

    def test_option_off_when_days_is_zero(self):
        with override_settings(MFA_REMEMBER_DAYS=0):
            self.client.force_login(self.hrmo_user)
            self.assertNotContains(self.client.get("/accounts/two-step/code/"), 'name="remember"')
            self.assertNotIn(self.cookie, self._pass().cookies)

    def test_forget_password_change_and_reset_each_end_it(self):
        self._pass(offset=-1)
        self.client.post("/accounts/two-step/", {"action": "forget"})
        self.assertEqual(self._login_again(), 302)

        self._pass(offset=0)
        self.hrmo_user.set_password("Another-Pass-2026")
        self.hrmo_user.save()
        self.assertEqual(self._login_again(), 302)

        self._pass(offset=1)
        self.mfa.reset(self.hrmo_user, "Phone lost")
        from hr.models import MfaDevice

        MfaDevice.objects.create(user=self.hrmo_user, secret=self.secret)
        self.assertEqual(self._login_again(), 302)

    def test_expired_or_borrowed_cookies_do_not_count(self):
        from unittest import mock
        from django.http import HttpResponse

        old = HttpResponse()
        with mock.patch("django.core.signing.time.time", return_value=time.time() - 31 * 86400):
            self.mfa.remember_browser(old, self.hrmo_user)
        self.client.cookies[self.cookie] = old.cookies[self.cookie].value
        self.assertEqual(self._login_again(), 302)

        # The HRMO's valid cookie, copied under another account's name, does not let that account in.
        from hr.models import MfaDevice

        MfaDevice.objects.create(user=self.emp_a_user, secret=self.mfa.new_secret())
        fresh = HttpResponse()
        self.mfa.remember_browser(fresh, self.hrmo_user)
        self.client.cookies[self.mfa.remember_cookie_name(self.emp_a_user)] = fresh.cookies[self.cookie].value
        self.client.force_login(self.emp_a_user)
        self.assertEqual(self.client.get("/").status_code, 302)
